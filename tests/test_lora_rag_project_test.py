"""Synthetic-only guards for the追加固定test100; no actual test/holdout reads."""
import copy
import json
from collections import Counter

import pytest

from src.generation import lora_rag_project_test as experiment
from tests.test_lora_rag_integration import synthetic_sample, populate_rag_tasks, configuration
from tests.test_generation import PromptPreflightTests


def config():
    return experiment.read_json(experiment.resolve(experiment.CONFIG))


@pytest.mark.parametrize('change', ['path', 'guard', 'resume', 'cohort', 'protocol'])
def test_rejects_forbidden_inputs_before_hashing(monkeypatch, change):
    cfg = config()
    if change == 'path':
        cfg['test_inputs']['path'] = 'forbidden/holdout.json'
    elif change == 'guard':
        cfg['guards']['output_based_adjustment'] = True
    elif change == 'resume':
        cfg['generation']['resume'] = True
    elif change == 'cohort':
        cfg['review']['products'] = 32
    else:
        cfg['development_protocol']['sha256'] = 'wrong'
    monkeypatch.setattr(experiment, 'read_json', lambda _: cfg)
    monkeypatch.setattr(experiment, 'digest', lambda _: pytest.fail('Must reject before reading input'))
    with pytest.raises(ValueError):
        experiment.checked_config()


def hundred():
    return [synthetic_sample(str(i), f'category{i % 8}') for i in range(100)]


def test_mapping_deterministic_global_and_category_balanced_before_outputs():
    rows = hundred()
    mapping = experiment.mapping_for(rows)
    assert mapping == experiment.mapping_for(list(reversed(rows)))
    assert len(mapping) == 100
    assert len({p['product_id'] for p in mapping}) == 100
    assert Counter(p['A'] for p in mapping)[experiment.VARIANTS[0]] == 50
    cats = {s['product_id']: s['category_l2'] for s in rows}
    for c in set(cats.values()):
        assignments = Counter(p['A'] for p in mapping if cats[p['product_id']] == c)
        assert abs(assignments[experiment.VARIANTS[0]] - assignments[experiment.VARIANTS[1]]) <= 1
    assert all(p['A'] != p['B'] for p in mapping)


def test_attempt_consumed_even_if_interrupted(tmp_path):
    path = experiment.claim_run(tmp_path)
    assert json.loads((path / 'attempt_started.json').read_text())['single_attempt_consumed']
    with pytest.raises(FileExistsError):
        experiment.claim_run(tmp_path)


def test_source_audit_scope_copy_only_no_status_or_policy_changes(monkeypatch):
    rows = [{'product_id': str(i), 'split': 'test'} for i in range(100)]
    audit = {str(i): {'product_id': str(i), 'scope': 'test100', 'status': 'REVIEW',
                     'blocked_fields': 'capacity'} for i in range(100)}
    original = copy.deepcopy(audit)
    seen = []
    def builder(record, row, policy, **kwargs):
        seen.append(copy.deepcopy(row))
        assert row['scope'] == record['split'] == 'test'
        assert policy == {'frozen': True}
        assert kwargs['source_sha256'] == 'fixed'
        return [{'product_id': record['product_id']}]
    monkeypatch.setattr(experiment, 'build_fact_units', builder)
    facts = experiment.test_facts(rows, audit, {'frozen': True}, {'path': 'synthetic', 'sha256': 'fixed'})
    assert len(facts) == len(seen) == 100 and audit == original
    assert all(r['status'] == 'REVIEW' and r['blocked_fields'] == 'capacity' for r in seen)


def test_identity_block_empty_context_keeps_id_no_fact_restoration(monkeypatch):
    s = synthetic_sample()
    s['source_quality']['identity_block'] = True
    s['used_attributes'] = {}
    f = {'product_id': '1', 'fact_id': 'blocked1', 'canonical_field': 'capacity',
         'quality_status': 'blocked_conflict'}
    context = {**s['rag_context'], 'identity_facts': [], 'mandatory_core_facts': [],
               'mandatory_injection_expected_fact_ids': [], 'identity_mandatory_coverage_fact_ids': []}
    monkeypatch.setattr(experiment.treatment, 'select_rag_v2_context', lambda *a, **k: copy.deepcopy(context))
    c, withheld = experiment.gated_context(s, [f], {})
    assert c['product_id'] == '1' and len(withheld) == 1
    assert s['core_attribute_count'] == 2
    f['quality_status'] = 'eligible'
    with pytest.raises(ValueError, match='blocked'):
        experiment.gated_context(s, [f], {})


def test_reused_frozen_prompt_seed_tokenizer_no_new_content_candidate():
    cfg = configuration()
    s = populate_rag_tasks(synthetic_sample(), cfg)
    for t in experiment.TASK_TYPES:
        assert s['rag_tasks'][t]['seed'] == s['tasks'][t]['seed']
        assert s['rag_tasks'][t]['messages'] == experiment.treatment.integration_messages(s, t, s['rag_context'], cfg)
    _, meta = experiment.treatment.preflight(PromptPreflightTests.FakeTokenizer(), s, 'title', cfg)
    assert meta['mandatory_core_expected'] == meta['mandatory_core_verified'] == 1
    with pytest.raises(ValueError):
        experiment.treatment.preflight(PromptPreflightTests.FakeTokenizer(drop_mandatory_when_tokenized=True), s, 'title', cfg)


@pytest.fixture
def workbook(tmp_path, monkeypatch):
    cfg = config()
    dev_cfg = configuration()
    samples = {s['product_id']: s for s in hundred()}
    versions = {v: {} for v in experiment.VARIANTS}
    for s in samples.values():
        populate_rag_tasks(s, dev_cfg)
        for variant in experiment.VARIANTS:
            tasks = {}
            for t in experiment.TASK_TYPES:
                raw = '["品牌：甲牌", "容量：500mL", "品类：保温杯"]' if t == 'selling_points' else '甲牌500mL保温杯'
                tasks[t] = {'raw_text': raw, 'parsed': experiment.parse_task(t, raw), 'latency_seconds': 1,
                    'memory': {'peak_allocated_mib': 1, 'peak_reserved_mib': 2}, 'hit_max_new_tokens': False,
                    'preflight': {'input_token_count': 100, 'mandatory_core_expected': 1, 'mandatory_core_verified': 1}
                    if variant == experiment.VARIANTS[1] else {'input_token_count': 50}}
            p = experiment.assemble(s, tasks, variant)
            p['rag_selection_seconds'] = .001
            p['latency_seconds'].update(three_task_plus_retrieval=3.001, three_task_plus_retrieval_and_preflight=3.01)
            versions[variant][s['product_id']] = p
    mapping = experiment.mapping_for(list(samples.values()))
    frozen = {'statistics': {'products': 100}}
    monkeypatch.setattr(experiment, 'paired_data', lambda: (cfg, frozen, samples, versions, mapping))
    monkeypatch.setattr(experiment, 'verify_protocol', lambda: None)
    monkeypatch.setattr(experiment, 'ROOT', tmp_path)
    monkeypatch.setattr(experiment, 'resolve', lambda name: tmp_path if name == experiment.OUTPUT else tmp_path / name)
    (tmp_path / experiment.CONFIG).parent.mkdir(parents=True)
    (tmp_path / experiment.CONFIG).write_text(json.dumps(cfg), encoding='utf-8')
    experiment.build_review()
    return tmp_path, cfg, frozen, samples, versions, mapping


def test_workbook_blank_anonymous_100_and_pending_not_fake_formal_metrics(workbook):
    from openpyxl import load_workbook
    path, _, _, samples, versions, mapping = workbook
    wb = load_workbook(path / experiment.WORKBOOK)
    ws = wb['盲评成对复核']
    cols = {c.value: c.column for c in ws[1]}
    assert ws.max_row == 102 and ws.max_column == 49
    assert wb['匿名映射_汇总后解盲'].sheet_state == 'veryHidden'
    assert wb['自动线索_非人工指标'].sheet_state == 'hidden'
    for r in range(3, 103):
        assert ws.cell(r, cols['product_id']).data_type == 's'
        for side in ('A', 'B'):
            for m in experiment.METRICS:
                c = ws.cell(r, cols[f'{side}_{m}'])
                assert c.value is None and not c.protection.locked
            assert ws.cell(r, cols[f'{side}_input_facts']).protection.locked
    with pytest.raises(ValueError, match='Incomplete'):
        experiment.validate_review_rows(ws, samples, versions, mapping)
    summary = json.loads((path / 'automatic_summary.json').read_text())
    assert summary['human_metrics'] is None
    assert not (path / 'human_metrics_unblinded.json').exists()


def test_complete_human_archive_same_metrics_evidence_and_frozen_cells(workbook):
    from openpyxl import load_workbook
    path, _, _, samples, versions, mapping = workbook
    wb = load_workbook(path / experiment.WORKBOOK)
    ws = wb['盲评成对复核']
    cols = {c.value: c.column for c in ws[1]}
    for r in range(3, 103):
        for side in ('A', 'B'):
            for m, value in zip(experiment.METRICS, (2, 1, 0)):
                ws.cell(r, cols[f'{side}_{m}'], value)
        for p in experiment.PREFERENCES:
            ws.cell(r, cols[p], '持平')
        ws.cell(r, cols['reviewer'], 'synthetic reviewer')
        ws.cell(r, cols['review_confirmed'], 1)
    assert len(experiment.validate_review_rows(ws, samples, versions, mapping)) == 100
    ws.cell(3, cols['A_factual_error_count'], 1)
    with pytest.raises(ValueError, match='evidence'):
        experiment.validate_review_rows(ws, samples, versions, mapping)
    ws.cell(3, cols['A_factual_error_count'], 0)
    ws.cell(3, cols['A_title'], 'altered')
    with pytest.raises(ValueError, match='Frozen'):
        experiment.validate_review_rows(ws, samples, versions, mapping)
    ws.cell(3, cols['A_title'], experiment.treatment.review_fixed(samples[mapping[0]['product_id']], mapping[0], versions)['A_title'])
    wb.save(path / experiment.WORKBOOK)
    experiment.summarize()
    human = json.loads((path / 'human_metrics_unblinded.json').read_text())
    assert human['versions'][experiment.VARIANTS[1]]['core_attribute_hit_rate'] == 1
    assert human['preferences']['title_quality']['tie'] == 100
    assert (path / 'blind_summary.json').exists()
    assert (path / 'final_test_report.md').exists()
    assert (path / 'human_review_final.csv').read_bytes().startswith(b'\xef\xbb\xbf')
    with pytest.raises(FileExistsError):
        experiment.summarize()
