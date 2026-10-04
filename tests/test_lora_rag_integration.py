"""Synthetic-only integration checks; never open test or holdout archives."""
import copy
import json
from collections import Counter
from pathlib import Path

import pytest

from src.generation import lora_rag_integration as experiment
from src.generation.lora_data import build_instruction_text
from tests.test_generation import PromptPreflightTests


def synthetic_sample(pid="1", category="保温杯"):
    core = {"品牌": ["甲牌"], "容量": ["500mL"]}
    sample = {"product_id": pid, "category_l1": "家居日用", "category_l2": category,
              "source_title": "SOURCE_TITLE_SENTINEL", "source_attributes": copy.deepcopy(core),
              "evaluation_attributes": copy.deepcopy(core), "core_attribute_count": 2,
              "used_attributes": copy.deepcopy(core), "source_quality": {"status": "PASS", "note": "合成测试"},
              "input_quality_actions": [], "tasks": {}, "rag_selection_seconds": .001}
    for task in experiment.TASK_TYPES:
        messages = [{"role": "system", "content": "你是中文电商文案编辑，只能使用输入中的品类和可靠核心属性。"},
                    {"role": "user", "content": build_instruction_text("家居日用", category, task, core)}]
        sample["tasks"][task] = {"messages": messages, "prompt_sha256": experiment.json_digest(messages),
                                 "seed": experiment.seed_for(pid, task), "preflight": {"input_token_count": 100}}
    fact = lambda fid, field, values: {"fact_id": fid, "canonical_field": field, "normalized_values": values,
                                     "product_id": pid, "quality_status": "eligible", "is_core_attribute": field in core}
    sample["rag_context"] = {"product_id": pid, "identity_facts": [fact("brand", "品牌", ["甲牌"])],
        "mandatory_core_facts": [fact("f1", "容量", ["500mL"])], "supplemental_facts": [], "negative_constraint_facts": [],
        "mandatory_injection_expected_fact_ids": ["f1"], "identity_mandatory_coverage_fact_ids": ["brand"]}
    return sample


def configuration():
    return experiment.read_json(experiment.resolve(experiment.CONFIG))


def populate_rag_tasks(sample, cfg):
    sample["rag_tasks"] = {}
    for task in experiment.TASK_TYPES:
        messages = experiment.integration_messages(sample, task, sample["rag_context"], cfg)
        sample["rag_tasks"][task] = {"messages": messages, "seed": sample["tasks"][task]["seed"],
                                      "prompt_sha256": experiment.json_digest(messages)}
    return sample


def test_original_core_task_requirements_and_seed_preserved_no_title_leak():
    cfg = configuration()
    s = populate_rag_tasks(synthetic_sample(), cfg)
    for task in experiment.TASK_TYPES:
        original = s["tasks"][task]
        combined = s["rag_tasks"][task]
        assert combined["messages"][1]["content"].startswith(original["messages"][1]["content"])
        assert combined["seed"] == original["seed"] == experiment.seed_for("1", task, 42)
        assert "SOURCE_TITLE_SENTINEL" not in json.dumps(combined["messages"])
        assert "必须完整包含" not in combined["messages"][1]["content"]
        assert "不要求各任务重复" in combined["messages"][1]["content"]


@pytest.mark.parametrize("change", ["foreign", "review", "conflict", "duplicate", "wrong_expected", "withheld"])
def test_invalid_new_facts_never_enter_prompt(change):
    s = synthetic_sample()
    c = s["rag_context"]
    if change == "foreign":
        c["mandatory_core_facts"][0]["product_id"] = "another"
    elif change in ("review", "conflict"):
        c["mandatory_core_facts"][0]["quality_status"] = "withhold_review" if change == "review" else "blocked_conflict"
    elif change == "duplicate":
        c["supplemental_facts"] = copy.deepcopy(c["mandatory_core_facts"])
    elif change == "wrong_expected":
        c["mandatory_injection_expected_fact_ids"] = ["different"]
    else:
        del s["used_attributes"]["容量"]
    with pytest.raises(ValueError):
        experiment.integration_messages(s, "title", c, configuration())


def test_actual_tokenizer_completeness_zero_denominator_and_overflow():
    cfg = configuration()
    s = populate_rag_tasks(synthetic_sample(), cfg)
    tokenizer = PromptPreflightTests.FakeTokenizer()
    _, meta = experiment.preflight(tokenizer, s, "title", cfg)
    assert meta["mandatory_core_expected"] == meta["mandatory_core_verified"] == 1
    assert meta["identity_core_expected"] == meta["identity_core_verified"] == 1
    assert meta["all_rag_facts_verified"] == 2
    assert meta["original_core_payload_complete"] and not any(tokenizer.truncation_values)
    with pytest.raises(ValueError, match="tokenizer"):
        experiment.preflight(PromptPreflightTests.FakeTokenizer(drop_mandatory_when_tokenized=True), s, "title", cfg)
    cfg["prompt"]["max_input_tokens"] = 10
    with pytest.raises(ValueError, match="上限"):
        experiment.preflight(tokenizer, s, "title", cfg)
    cfg = configuration()
    s["rag_context"]["mandatory_core_facts"] = []
    s["rag_context"]["mandatory_injection_expected_fact_ids"] = []
    populate_rag_tasks(s, cfg)
    _, meta = experiment.preflight(tokenizer, s, "title", cfg)
    assert meta["mandatory_core_injection_rate"] is None
    assert s["core_attribute_count"] == 2  # Source missing/policy exclusion never changes denominator.


def test_control_output_seed_or_prompt_mismatch_rejected():
    s = synthetic_sample()
    product = {"product_id": "1", "variant": "LoRA", "tasks": copy.deepcopy(s["tasks"])}
    experiment.verify_control(s, product)
    product["tasks"]["title"]["seed"] += 1
    with pytest.raises(ValueError, match="seed"):
        experiment.verify_control(s, product)


def test_restricted_paths_checked_before_hashing(monkeypatch):
    cfg = configuration()
    cfg["control_outputs"] = "forbidden/test.jsonl"
    monkeypatch.setattr(experiment, "read_json", lambda _: cfg)
    monkeypatch.setattr(experiment, "digest", lambda _: pytest.fail("Cannot hash forbidden input"))
    with pytest.raises(ValueError, match="Forbidden path"):
        experiment.checked_config()


def test_guard_missing_keys_rejected_before_hashing(monkeypatch):
    cfg = configuration()
    del cfg["guards"]["holdout_generation"]
    monkeypatch.setattr(experiment, "read_json", lambda _: cfg)
    monkeypatch.setattr(experiment, "digest", lambda _: pytest.fail("Guard failure before data read"))
    with pytest.raises(ValueError, match="guards"):
        experiment.checked_config()


def twenty_four():
    return [synthetic_sample(str(i * 3 + j), f"category{i}") for i in range(8) for j in range(3)]


def test_selection_and_blinding_fixed_before_outputs_no_new_sample():
    samples = twenty_four()
    development = {"status": "completed", "generation_output_holdout_executed": False,
                   "sample_selection": {"protected_holdout_overlap_count": 0}, "results": [{"product_id": s["product_id"]} for s in samples]}
    selected = experiment.select_samples(development, list(reversed(samples)))
    assert [s["product_id"] for s in selected] == [s["product_id"] for s in samples]
    first = experiment.mapping_for(samples, 41008)
    assert first == experiment.mapping_for(list(reversed(samples)), 41008)
    assert Counter(p["A"] for p in first)[experiment.VARIANTS[0]] == 12
    assert all(p["A"] != p["B"] for p in first)
    development["sample_selection"]["protected_holdout_overlap_count"] = 1
    with pytest.raises(ValueError, match="non-holdout"):
        experiment.select_samples(development, samples)


def test_grounding_sees_original_core_even_when_rag_excludes_it():
    s = synthetic_sample()
    s["used_attributes"]["耳机类别"] = ["头戴式", "入耳式"]
    context = experiment.grounding_context(s, with_rag=True)
    assert any(f["canonical_field"] == "耳机类别" for f in context["selected_facts"])


def test_p7_withdrawals_intersect_before_rag_selection(monkeypatch):
    s = synthetic_sample()
    facts = copy.deepcopy(s["rag_context"]["mandatory_core_facts"])
    del s["used_attributes"]["容量"]
    observed = []
    def selector(units, policy, top_k):
        observed.extend(units)
        context = copy.deepcopy(s["rag_context"])
        context["mandatory_core_facts"] = []
        context["mandatory_injection_expected_fact_ids"] = []
        return context
    monkeypatch.setattr(experiment, "select_rag_v2_context", selector)
    context, withheld = experiment.gated_context(s, facts, {})
    assert observed == []
    assert withheld == [{"fact_id": "f1", "field": "容量", "reason": "already_withheld_in_frozen_P7_input"}]
    assert facts[0]["quality_status"] == "eligible"  # No mutation of frozen RAG facts/policy.
    assert s["core_attribute_count"] == 2


@pytest.fixture
def workbook(tmp_path, monkeypatch):
    cfg = configuration()
    samples = {s["product_id"]: s for s in twenty_four()}
    versions = {v: {} for v in experiment.VARIANTS}
    for s in samples.values():
        populate_rag_tasks(s, cfg)
        for variant in experiment.VARIANTS:
            tasks = {}
            for t in experiment.TASK_TYPES:
                raw = '["品牌：甲牌", "容量：500mL", "品类：保温杯"]' if t == "selling_points" else "甲牌500mL保温杯"
                tasks[t] = {"raw_text": raw, "parsed": experiment.parse_task(t, raw), "latency_seconds": 1,
                    "memory": {"peak_allocated_mib": 1, "peak_reserved_mib": 2}, "hit_max_new_tokens": False,
                    "preflight": {"input_token_count": 100, "mandatory_core_expected": 1, "mandatory_core_verified": 1} if variant == experiment.VARIANTS[1] else {"input_token_count": 50}}
            product = experiment.assemble(s, tasks, variant)
            product["rag_selection_seconds"] = .001
            product["latency_seconds"].update(three_task_plus_retrieval=3.001, three_task_plus_retrieval_and_preflight=3.01)
            versions[variant][s["product_id"]] = product
    mapping = experiment.mapping_for(list(samples.values()), 41008)
    frozen = {"statistics": {"products": 24}}
    monkeypatch.setattr(experiment, "paired_data", lambda: (cfg, frozen, samples, versions, mapping))
    monkeypatch.setattr(experiment, "verify_protocol", lambda: None)
    monkeypatch.setattr(experiment, "ROOT", tmp_path)
    monkeypatch.setattr(experiment, "resolve", lambda _: tmp_path)
    experiment.build_review()
    return tmp_path, samples, versions, mapping


def test_workbook_anonymity_blank_labels_immutable_inputs_and_pending_metrics(workbook):
    from openpyxl import load_workbook
    path, samples, versions, mapping = workbook
    wb = load_workbook(path / "lora_rag_development24_blind_review.xlsx")
    ws = wb["盲评成对复核"]
    columns = {c.value: c.column for c in ws[1]}
    assert ws.max_row == 26
    assert wb["匿名映射_汇总后解盲"].sheet_state == "veryHidden"
    assert wb["自动线索_非人工指标"].sheet_state == "hidden"
    assert ws.cell(3, columns["product_id"]).data_type == "s"
    for side in ("A", "B"):
        for metric in experiment.METRICS:
            cell = ws.cell(3, columns[f"{side}_{metric}"])
            assert cell.value is None and not cell.protection.locked
        assert ws.cell(3, columns[f"{side}_input_facts"]).protection.locked
    summary = json.loads((path / "automatic_summary.json").read_text(encoding="utf-8"))
    assert summary["human_metrics"] is None and summary["human_preferences"] is None
    with pytest.raises(ValueError, match="Incomplete"):
        experiment.validate_review_rows(ws, samples, versions, mapping)


def test_human_review_complete_required_original_metrics_and_no_fixed_mutation(workbook):
    from openpyxl import load_workbook
    path, samples, versions, mapping = workbook
    wb = load_workbook(path / "lora_rag_development24_blind_review.xlsx")
    ws = wb["盲评成对复核"]
    cols = {c.value: c.column for c in ws[1]}
    for row in range(3, 27):
        for side in ("A", "B"):
            for metric, value in zip(experiment.METRICS, (2, 1, 0)):
                ws.cell(row, cols[f"{side}_{metric}"], value)
        for p in experiment.PREFERENCES:
            ws.cell(row, cols[p], "持平")
        ws.cell(row, cols["reviewer"], "人工")
        ws.cell(row, cols["review_confirmed"], 1)
    rows = experiment.validate_review_rows(ws, samples, versions, mapping)
    assert len(rows) == 24
    ws.cell(3, cols["A_title"], "tampered")
    with pytest.raises(ValueError, match="Frozen review data"):
        experiment.validate_review_rows(ws, samples, versions, mapping)


def test_existing_generation_directory_rejects_retries(tmp_path, monkeypatch):
    folder = tmp_path / "lora_rag"
    folder.mkdir()
    monkeypatch.setattr(experiment, "verify_protocol", lambda: ({}, {}, {}))
    monkeypatch.setattr(experiment, "resolve", lambda _: tmp_path)
    with pytest.raises(FileExistsError, match="rerun"):
        experiment.generate()
