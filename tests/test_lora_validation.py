"""P7 checks use synthetic records; never open project_test or holdout archives."""
import copy
import json
from pathlib import Path

import pytest

from src.generation import lora_validation as inference
from src.generation import lora_validation_review as review
from src.generation.evaluation import GenerationJudgment, evaluate_judgments


def configuration():
    return inference.read_json(inference.resolve(inference.CONFIG_PATH))


def synthetic(n=1):
    cfg = configuration()
    cfg["source"]["products"] = n
    p5 = inference.read_json(inference.resolve(cfg["p5_config"]["path"]))
    definition = inference.read_json(inference.resolve(cfg["evaluation_definition"]["path"]))
    records, facts = [], []
    for i in range(n):
        pid = str(610000000000 + i)
        records.append({"product_id": pid, "split": "validation", "category_l1": "家居日用", "category_l2": "保温杯",
                        "title": "SOURCE_TITLE_SENTINEL", "attributes": {"品牌": ["甲牌"], "型号": ["A1"],
                        "材质": ["304不锈钢"], "容量": ["500mL"], "杯子样式": ["直身杯"]}})
        facts.append({"product_id": pid, "split": "validation", "field_name": "容量", "canonical_field": "容量",
                      "quality_reasons": [], "evidence": {"audit_status": "PASS", "audit_note": "合成测试"}})
    return records, facts, p5, definition, cfg


def test_seed_and_prompt_pairing_no_source_title():
    args = synthetic()
    sample = inference.build_inputs(*args)[0]
    for task, value in sample["tasks"].items():
        assert value["seed"] == inference.seed_for(sample["product_id"], task)
        assert "SOURCE_TITLE_SENTINEL" not in json.dumps(value["messages"])
        assert value["messages"][1]["content"] == inference.build_instruction_text(
            sample["category_l1"], sample["category_l2"], task, sample["used_attributes"])
    assert len({v["seed"] for v in sample["tasks"].values()}) == 3
    assert inference.seed_for("1", "title") != inference.seed_for("2", "title")
    with pytest.raises(ValueError):
        inference.seed_for("1", "joint_copy")


@pytest.mark.parametrize("reason", ["field_requires_review", "field_conflict", "identity_or_category_conflict"])
def test_gate_does_not_change_formal_denominator_or_remove_product(reason):
    args = synthetic()
    before = inference.build_inputs(*args)[0]
    args[1][0]["quality_reasons"] = [reason]
    args[1][0]["evidence"]["audit_status"] = "CONFLICT"
    after = inference.build_inputs(*args)[0]
    assert after["evaluation_attributes"] == before["evaluation_attributes"]
    assert after["core_attribute_count"] == before["core_attribute_count"]
    assert "容量" not in after["used_attributes"]
    if reason == "identity_or_category_conflict":
        assert after["used_attributes"] == {}


def test_rag_generic_multivalue_reason_is_not_new_input_policy():
    args = synthetic()
    before = inference.build_inputs(*args)[0]
    args[1][0]["quality_reasons"] = ["multivalue_requires_review"]
    assert inference.build_inputs(*args)[0]["used_attributes"] == before["used_attributes"]


def test_only_project_validation_records_and_audit():
    args = synthetic()
    args[0][0]["split"] = "test"
    with pytest.raises(ValueError, match="Non-validation"):
        inference.build_inputs(*args)
    args = synthetic()
    args[1][0]["split"] = "test"
    with pytest.raises(ValueError, match="validation-only"):
        inference.build_inputs(*args)


def test_all_input_paths_checked_before_any_data_read(monkeypatch):
    cfg = configuration()
    cfg["validation_quality_facts"]["path"] = "forbidden/holdout.jsonl"
    monkeypatch.setattr(inference, "read_json", lambda _: cfg)
    monkeypatch.setattr(inference, "digest", lambda _: pytest.fail("No data read before safety guard"))
    with pytest.raises(ValueError, match="Forbidden input"):
        inference.checked_config()


def test_blinding_balanced_reproducible_independent_of_source_order():
    samples = [{"product_id": str(i)} for i in range(200)]
    first = inference.blind_mapping(samples, 41007)
    assert first == inference.blind_mapping(list(reversed(samples)), 41007)
    assert sum(p["A"] == "Base" for p in first) == 100
    assert all(p["A"] != p["B"] for p in first)
    assert len({p["pair_id"] for p in first}) == 200


@pytest.mark.parametrize("text", ['["品牌：甲","容量：500mL","材质：钢"]', "1. 品牌：甲\n2. 容量：500mL\n3. 材质：钢", "- 品牌：甲\n- 容量：500mL\n- 材质：钢"])
def test_three_point_parsing_does_not_require_prompt_json(text):
    result = inference.parse_task("selling_points", text)
    assert result["success"] and len(result["value"]) == 3


def test_assembly_preserves_failed_raw_output_without_fallback_or_retry():
    tasks = {t: {"raw_text": "原始文本", "parsed": inference.parse_task(t, "原始文本"), "latency_seconds": 1.0} for t in inference.TASK_TYPES}
    result = inference.assemble({"product_id": "1"}, tasks, "Base")
    assert result["assembled"]["selling_points"] is None
    assert result["assembled"]["complete_raw_text"] == "原始文本\n原始文本\n原始文本"
    assert not result["assembled"]["structure_success"]
    assert result["latency_seconds"]["three_task_total"] == 3.0


def test_human_metrics_equal_frozen_evaluation():
    rows = [{"product_id": "1", "core_attribute_count": 5, "matched_attribute_count": 4, "fluency_pass": 1, "factual_error_count": 2},
            {"product_id": "2", "core_attribute_count": 7, "matched_attribute_count": 3, "fluency_pass": 0, "factual_error_count": 0}]
    expected = evaluate_judgments([GenerationJudgment(**r, category_style_pass=0) for r in rows])
    actual = review.official_metrics(rows)
    for key in expected.keys() - {"category_style_pass_rate"}:
        assert actual[key] == expected[key]
    with pytest.raises(ValueError, match="duplicate"):
        review.official_metrics(rows + rows)


@pytest.fixture
def mock_review(tmp_path, monkeypatch):
    samples = inference.build_inputs(*synthetic(200))
    by_id = {s["product_id"]: s for s in samples}
    mapping = inference.blind_mapping(samples, 41007)
    versions = {}
    for variant in ("Base", "LoRA"):
        versions[variant] = {}
        folder = tmp_path / variant.lower()
        folder.mkdir()
        (folder / "run_report.json").write_text("{}", encoding="utf-8")
        for s in samples:
            tasks = {t: {"raw_text": "甲牌A1保温杯", "parsed": inference.parse_task(t, "甲牌A1保温杯"),
                         "latency_seconds": 1, "hit_max_new_tokens": False} for t in inference.TASK_TYPES}
            versions[variant][s["product_id"]] = inference.assemble(s, tasks, variant)
    inference.atomic_json(tmp_path / "protocol_manifest.json", {})
    inference.atomic_json(tmp_path / "blind_mapping.json", {"pairs": mapping})
    cfg = {"output": str(tmp_path)}
    monkeypatch.setattr(review, "paired_data", lambda: (cfg, by_id, versions, mapping))
    monkeypatch.setattr(review, "resolve", lambda p: Path(p) if Path(p).is_absolute() else tmp_path / p)
    review.build_workbook()
    return tmp_path


def test_workbook_ids_blank_human_labels_and_hidden_mapping(mock_review):
    from openpyxl import load_workbook
    wb = load_workbook(mock_review / "project_validation200_blind_review.xlsx")
    ws = wb["盲评成对复核"]
    assert ws.max_row == 202
    cols = {c.value: c.column for c in ws[1]}
    assert ws.cell(3, cols["product_id"]).data_type == "s"
    assert wb["匿名映射_汇总后解盲"].sheet_state == "veryHidden"
    assert wb["自动线索_非人工指标"].sheet_state == "hidden"
    for side in ("A", "B"):
        for metric in review.METRICS:
            c = ws.cell(3, cols[f"{side}_{metric}"])
            assert c.value is None and not c.protection.locked
    assert ws.cell(3, cols["core_attribute_count"]).protection.locked


def test_incomplete_review_cannot_unblind_or_create_metrics(mock_review):
    with pytest.raises(ValueError, match="Incomplete human"):
        review.summarize_workbook()
    assert not (mock_review / "human_metrics_unblinded.json").exists()
    assert not (mock_review / "blind_summary.json").exists()


def test_complete_review_preserves_diagnostics_and_unblinds_after_summary(mock_review):
    from openpyxl import load_workbook
    path = mock_review / "project_validation200_blind_review.xlsx"
    wb = load_workbook(path)
    ws = wb["盲评成对复核"]
    cols = {c.value: c.column for c in ws[1]}
    for row in range(3, 203):
        for side in ("A", "B"):
            for metric, value in zip(review.METRICS, (3, 1, 0)):
                ws.cell(row, cols[f"{side}_{metric}"], value)
            ws.cell(row, cols[f"{side}_field_stacking"], 1)
            ws.cell(row, cols[f"{side}_notes"], "合成测试依据")
        for preference in review.PREFERENCES:
            ws.cell(row, cols[preference], "持平")
        ws.cell(row, cols["reviewer"], "synthetic_test")
        ws.cell(row, cols["review_confirmed"], 1)
    wb.save(path)
    review.summarize_workbook()
    result = inference.read_json(mock_review / "human_metrics_unblinded.json")
    assert result["official_metrics"]["LoRA"]["sample_count"] == 200
    assert result["paired_preferences"]["title_quality"]["tie"] == 200
    assert result["diagnostics_NOT_OFFICIAL_METRICS"]["Base"]["field_stacking"]["positive_samples"] == 200
    assert (mock_review / "human_review_final.csv").read_bytes().startswith(b"\xef\xbb\xbf")


def test_mapping_tampering_is_rejected(mock_review):
    from openpyxl import load_workbook
    path = mock_review / "project_validation200_blind_review.xlsx"
    wb = load_workbook(path)
    wb["匿名映射_汇总后解盲"].cell(2, 3, "changed")
    wb.save(path)
    with pytest.raises(ValueError, match="mapping changed"):
        review.summarize_workbook()
