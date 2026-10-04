"""Quick-review checks use synthetic 8x25 source IDs, no generation/test data."""
from pathlib import Path

import pytest

from scripts import review_lora_project_validation_quick32 as quick
from src.generation.lora_training import atomic_json, read_json


def synthetic_source():
    return [{"product_id": str(610000000000 + c * 25 + i), "category_l2": f"category{c}", "split": "validation"}
            for c in range(8) for i in range(25)]


def test_stratified_selection_fixed_and_output_independent():
    records = synthetic_source()
    cfg = {"seed": 42, "blinding_seed": 41008}
    selected, mapping = quick.choose(records, cfg)
    for row in records:
        row["generated_output_ignored"] = "different hypothetical quality"
        row["source_quality_ignored"] = "CONFLICT"
    assert (selected, mapping) == quick.choose(list(reversed(records)), cfg)
    assert len(selected) == 32 and len({r["product_id"] for r in selected}) == 32
    for c in range(8):
        assert sum(r["category_l2"] == f"category{c}" for r in selected) == 4
        assert sum(r["category_l2"] == f"category{c}" and r["A"] == "Base" for r in mapping) == 2
    assert quick.choose(records, {"seed": 43, "blinding_seed": 41008})[0] != selected


def test_invalid_scope_rejected():
    rows = synthetic_source()
    rows[0]["split"] = "test"
    with pytest.raises(ValueError, match="Non-validation"):
        quick.choose(rows, {"seed": 42, "blinding_seed": 41008})


@pytest.fixture
def review_fixture(tmp_path, monkeypatch):
    selected, mapping = quick.choose(synthetic_source(), {"seed": 42, "blinding_seed": 41008})
    samples, versions = {}, {"Base": {}, "LoRA": {}}
    for row in selected:
        pid = row["product_id"]
        samples[pid] = {"product_id": pid, "category_l1": "synthetic", "category_l2": row["category_l2"],
            "source_title": "源标题", "source_attributes": {"品牌": ["甲"]}, "evaluation_attributes": {"品牌": ["甲"]},
            "core_attribute_count": 1, "used_attributes": {"品牌": ["甲"]}, "input_quality_actions": [],
            "source_quality": {"status": "PASS", "note": "synthetic"}}
        for variant in versions:
            versions[variant][pid] = {"assembled": {"complete_raw_text": "甲牌文案", "structure_success": True},
                "tasks": {t: {"raw_text": "甲牌文案", "hit_max_new_tokens": False} for t in ("title", "selling_points", "short_description")}}
    atomic_json(tmp_path / "sampling_manifest.json", {})
    for variant in ("base", "lora"):
        (tmp_path / variant).mkdir()
        atomic_json(tmp_path / variant / "run_report.json", {})
    monkeypatch.setattr(quick, "verify_sample", lambda: ({"output": str(tmp_path)}, {}, tmp_path, mapping))
    monkeypatch.setattr(quick, "paired_data", lambda: ({}, samples, versions, []))
    monkeypatch.setattr(quick, "resolve", lambda p: Path(p) if Path(p).is_absolute() else tmp_path / p)
    quick.workbook()
    return tmp_path


def test_quick_workbook_32_rows_no_labels_and_partial_summary_blocked(review_fixture):
    from openpyxl import load_workbook
    wb = load_workbook(review_fixture / "project_validation32_blind_review.xlsx")
    ws = wb["盲评成对复核"]
    assert ws.max_row == 34
    assert wb["匿名映射_汇总后解盲"].sheet_state == "veryHidden"
    cols = {c.value: c.column for c in ws[1]}
    assert ws.cell(3, cols["product_id"]).data_type == "s"
    assert ws.cell(3, cols["A_factual_error_count"]).value is None
    assert not ws.column_dimensions["H"].hidden
    with pytest.raises(ValueError, match="Incomplete human"):
        quick.summarize()
    assert not (review_fixture / "blind_summary.json").exists()


def test_quick_complete_summary_scoped_to_32_and_preserves_template_diagnosis(review_fixture):
    from openpyxl import load_workbook
    path = review_fixture / "project_validation32_blind_review.xlsx"
    wb = load_workbook(path)
    ws = wb["盲评成对复核"]
    cols = {c.value: c.column for c in ws[1]}
    for row in range(3, 35):
        for side in ("A", "B"):
            for metric, value in zip(quick.METRICS, (1, 1, 0)):
                ws.cell(row, cols[f"{side}_{metric}"], value)
            ws.cell(row, cols[f"{side}_mechanical_template"], 0)
        for p in quick.PREFERENCES:
            ws.cell(row, cols[p], "持平")
        ws.cell(row, cols["reviewer"], "synthetic")
        ws.cell(row, cols["review_confirmed"], 1)
    wb.save(path)
    quick.summarize()
    result = read_json(review_fixture / "human_metrics_unblinded.json")
    assert result["official_metrics"]["Base"]["sample_count"] == 32
    assert result["official_metrics"]["LoRA"]["core_attribute_total"] == 32
    assert result["diagnostics_NOT_FORMAL_METRICS"]["Base"]["mechanical_template"]["annotated"] == 32
    assert result["project_test_run"] is False
