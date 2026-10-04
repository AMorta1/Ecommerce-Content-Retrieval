"""Synthetic formal-run guards/review tests; no GPU, real test outputs or labels."""
import ast
import copy
import inspect
import json
import re
from pathlib import Path

import pytest

from src.generation import lora_formal_test as formal
from src.generation import lora_formal_review as review
from src.generation import lora_validation as previous


def synthetic():
    cfg = previous.read_json(previous.resolve(previous.CONFIG_PATH))
    p5 = previous.read_json(previous.resolve(cfg["p5_config"]["path"]))
    definition = previous.read_json(previous.resolve(cfg["evaluation_definition"]["path"]))
    records, audit = [], {}
    for i in range(100):
        pid = str(710000000000 + i)
        record = {"product_id": pid, "split": "test", "category_l1": "家居日用", "category_l2": "保温杯",
            "title": "SOURCE_TITLE_SENTINEL", "attributes": {"品牌": ["甲牌"], "型号": ["A1"], "材质": ["304不锈钢"], "容量": ["500mL"], "杯子样式": ["直身杯"]}}
        record["generation_input"] = {"attributes": formal.select_core_attributes(record, definition)}
        records.append(record)
        audit[pid] = {"status": "PASS", "blocked_fields": [], "identity_block": False, "note": "合成测试"}
    return records, audit, p5, definition, cfg


def test_frozen_prompt_seed_no_title_and_all_samples_retained():
    args = synthetic()
    samples = formal.build_inputs(*args)
    assert len(samples) == 100
    for sample in samples:
        for task, value in sample["tasks"].items():
            assert value["seed"] == previous.seed_for(sample["product_id"], task)
            assert value["messages"][1]["content"] == previous.build_instruction_text(sample["category_l1"], sample["category_l2"], task, sample["used_attributes"])
            assert "SOURCE_TITLE_SENTINEL" not in json.dumps(value["messages"])
    assert len({t["seed"] for s in samples for t in s["tasks"].values()}) == 300


@pytest.mark.parametrize("status,identity", [("REVIEW", False), ("CONFLICT", False), ("CONFLICT", True)])
def test_audit_gate_unchanged_denominator(status, identity):
    args = synthetic()
    before = formal.build_inputs(*args)[0]
    args[1][args[0][0]["product_id"]].update(status=status, blocked_fields=["容量"], identity_block=identity)
    after = formal.build_inputs(*args)[0]
    assert before["evaluation_attributes"] == after["evaluation_attributes"]
    assert before["core_attribute_count"] == after["core_attribute_count"]
    assert "容量" not in after["used_attributes"]
    if identity:
        assert after["used_attributes"] == {}


def test_no_filter_reorder_or_validation_masquerade():
    args = synthetic()
    with pytest.raises(ValueError, match="100"):
        formal.build_inputs(args[0][:-1], *args[1:])
    args[0][0]["split"] = "validation"
    with pytest.raises(ValueError, match="Non-test"):
        formal.build_inputs(*args)
    args = synthetic()
    args[0][0]["generation_input"]["attributes"] = {}
    with pytest.raises(ValueError, match="denominator"):
        formal.build_inputs(*args)


def test_input_path_guard_before_digest(monkeypatch):
    cfg = formal.read_json(formal.resolve(formal.CONFIG))
    cfg["source"]["path"] = "forbidden/holdout.jsonl"
    monkeypatch.setattr(formal, "read_json", lambda _: cfg)
    monkeypatch.setattr(formal, "digest", lambda _: pytest.fail("Must not hash forbidden path"))
    with pytest.raises(ValueError, match="Forbidden input"):
        formal.checked_config()


def test_single_attempt_consumed_even_before_model_load(tmp_path):
    path = formal.claim_run(tmp_path, "Base")
    assert formal.read_json(path / "attempt_started.json")["single_attempt_consumed"]
    with pytest.raises(FileExistsError):
        formal.claim_run(tmp_path, "Base")
    assert formal.claim_run(tmp_path, "LoRA").exists()


def test_p7_inference_execution_body_is_identical():
    # Compare every execution AST node from model imports through final protocol
    # verification; only the inherited configuration variable name differs.
    def body(function, rename=False):
        text = inspect.getsource(function)
        if rename:
            text = re.sub(r"\bparent\b", "cfg", text)
        tree = ast.parse(text)
        statements = next(n for n in ast.walk(tree) if isinstance(n, ast.Try)).body
        stop = next(i for i, n in enumerate(statements) if isinstance(n, ast.Expr) and isinstance(n.value, ast.Call)
                    and isinstance(n.value.func, ast.Name) and n.value.func.id == "verify_protocol")
        return [ast.dump(n, include_attributes=False) for n in statements[:stop + 1]]
    assert body(previous.generate) == body(formal.generate, rename=True)


@pytest.fixture
def mock_review(tmp_path, monkeypatch):
    samples_list = formal.build_inputs(*synthetic())
    samples = {s["product_id"]: s for s in samples_list}
    mapping = previous.blind_mapping(samples_list, 41007)
    versions = {}
    for variant in ("Base", "LoRA"):
        versions[variant] = {}
        for s in samples_list:
            tasks = {t: {"raw_text": "甲牌A1保温杯", "parsed": previous.parse_task(t, "甲牌A1保温杯"),
                        "latency_seconds": 1, "hit_max_new_tokens": False,
                        "memory": {"peak_allocated_mib": 100, "peak_reserved_mib": 120}} for t in previous.TASK_TYPES}
            versions[variant][s["product_id"]] = previous.assemble(s, tasks, variant)
    previous.atomic_json(tmp_path / "protocol_manifest.json", {})
    previous.atomic_json(tmp_path / "review_manifest.json", {"protocol_manifest_sha256": previous.digest(tmp_path / "protocol_manifest.json"), "files": []})
    monkeypatch.setattr(review, "paired_data", lambda: ({"test_history_disclosure": "synthetic"}, {}, {}, samples, versions, mapping))
    monkeypatch.setattr(review, "resolve", lambda _: tmp_path)
    review.build_workbook(tmp_path, samples, versions, mapping)
    return tmp_path


def test_workbook100_blank_labels_hidden_mapping(mock_review):
    from openpyxl import load_workbook
    w = load_workbook(mock_review / review.WORKBOOK)
    s = w["盲评成对复核"]
    assert s.max_row == 102
    cols = {c.value: c.column for c in s[1]}
    assert w["匿名映射_汇总后解盲"].sheet_state == "veryHidden"
    assert w["自动线索_非人工指标"].sheet_state == "hidden"
    for i in range(3,103):
        assert s.cell(i,cols["product_id"]).data_type == "s"
        for k in (*review.PREFERENCES,"reviewer","review_confirmed",*[f"{side}_{m}" for side in ("A","B") for m in review.METRICS]):
            assert s.cell(i,cols[k]).value is None


def test_incomplete_test_review_never_unblinds(mock_review):
    with pytest.raises(ValueError, match="Incomplete human"):
        review.summarize()
    assert not (mock_review / "blind_summary.json").exists()
    assert not (mock_review / "human_metrics_unblinded.json").exists()


def test_completed_test_review_reuses_metrics_and_finalizes(mock_review):
    from openpyxl import load_workbook
    path = mock_review / review.WORKBOOK
    w = load_workbook(path)
    s = w["盲评成对复核"]
    cols = {c.value:c.column for c in s[1]}
    for i in range(3,103):
        for side in ("A","B"):
            for metric,value in zip(review.METRICS,(3,1,0)):
                s.cell(i,cols[f"{side}_{metric}"],value)
        for pref in review.PREFERENCES:
            s.cell(i,cols[pref],"持平")
        s.cell(i,cols["reviewer"],"synthetic_test")
        s.cell(i,cols["review_confirmed"],1)
    w.save(path)
    review.summarize()
    result=previous.read_json(mock_review / "human_metrics_unblinded.json")
    assert result["official_metrics"]["LoRA"]["sample_count"]==100
    assert result["paired_preferences"]["title_quality"]["tie"]==100
    assert (mock_review / "final_test_manifest.json").exists()
    assert (mock_review / "human_review_final.csv").read_bytes().startswith(b"\xef\xbb\xbf")
    with pytest.raises(FileExistsError):
        review.summarize()


def test_mapping_tampering_and_source_edits_rejected(mock_review):
    from openpyxl import load_workbook
    path=mock_review / review.WORKBOOK
    w=load_workbook(path)
    w["匿名映射_汇总后解盲"].cell(2,3,"tampered")
    w.save(path)
    with pytest.raises(ValueError,match="mapping"):
        review.summarize()


def test_range_and_effect_hints_never_equal_fact_judgments():
    sample={"evaluation_attributes":{"保温时长":["[6,12)小时"]}}
    product={"assembled":{"complete_raw_text":"保温6到12小时，适合儿童，确保使用方便"}}
    hints=review.extra_hints(sample,product)
    assert hints["range_boundary_clues_REVIEW_ONLY"]
    assert hints["effect_or_scope_terms_REVIEW_ONLY"]
    assert "factual_error_count" not in hints
