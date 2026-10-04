"""Archive user-confirmed development24 review; no inference/training/test access."""
from __future__ import annotations

import csv
import json
import os
import shutil
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.generation.lora_rag_integration import (OUTPUT, VARIANTS, PREFERENCES, DIAGNOSTICS, METRICS,
    paired_data, summarize, validate_review_rows, verify_protocol)
from src.generation.lora_training import ROOT, atomic_json, digest, now, read_json, resolve


def main():
    out = resolve(OUTPUT)
    path = out / "lora_rag_development24_blind_review.xlsx"
    backup = path.with_name(path.stem + ".user_completed.xlsx")
    temporary = path.with_name(path.stem + ".confirmation.tmp.xlsx")
    for name in ("human_review_confirmation_record.json", "blind_summary.json", "human_metrics_unblinded.json", "final_review_manifest.json"):
        if (out / name).exists():
            raise FileExistsError("Do not overwrite an existing finalization")
    if backup.exists() or temporary.exists():
        raise FileExistsError("Existing user backup/temporary; inspect before continuing")
    _, _, samples, versions, mapping = paired_data()
    review = read_json(out / "review_manifest.json")
    for entry in review["files"]:
        if digest(resolve(entry["path"])) != entry["sha256"]:
            raise ValueError("Frozen review generation changed")
    for entry in read_json(out / "automatic_phase_manifest.json")["files"]:
        original = out / "lora_rag_development24_blind_review.pre_ai_import.xlsx" if resolve(entry["path"]) == path else resolve(entry["path"])
        if digest(original) != entry["sha256"]:
            raise ValueError("Historical automatic archive changed")
    before_sha = digest(path)
    wb = load_workbook(path, data_only=False)
    ws = wb["盲评成对复核"]
    columns = {c.value: c.column for c in ws[1]}
    raw_values = list(ws.values)
    sheets = {s.title: (s.sheet_state, list(s.values)) for s in wb if s.title != ws.title}
    if list(wb["匿名映射_汇总后解盲"].values)[1:] != [tuple(p[k] for k in ("pair_id", "product_id", "A", "B")) for p in mapping]:
        raise ValueError("Sealed mapping changed")
    if ws.max_row != 26 or any(str(ws.cell(r, columns["review_confirmed"]).value) != "1" for r in range(3, 27)):
        raise ValueError("All 24 rows must be explicitly confirmed")
    supplied_reviewers = Counter(ws.cell(r, columns["reviewer"]).value for r in range(3, 27))
    filled = []
    for r in range(3, 27):
        cell = ws.cell(r, columns["reviewer"])
        if cell.value is None or not str(cell.value).strip():
            cell.value = "用户复核确认（ChatGPT初审 / Codex修正）"
            cell.data_type = "s"
            filled.append(ws.cell(r, columns["pair_id"]).value)
    rows = validate_review_rows(ws, samples, versions, mapping)
    original = load_workbook(out / "lora_rag_development24_blind_review.chatgpt_rechecked_source.xlsx", data_only=False)[ws.title]
    original_columns = {c.value: c.column for c in original[1]}
    amendments = read_json(out / "ai_review_amendments_v1.json")
    edits = []
    for r in range(3, 27):
        pair = ws.cell(r, columns["pair_id"]).value
        patches = amendments["corrections"].get(pair, {}).get("fields", {})
        for key in (*PREFERENCES, *[f"{side}_{k}" for side in ("A", "B") for k in (*METRICS, *DIAGNOSTICS, "notes")], "pair_notes"):
            expected = patches[key]["value"] if key in patches else original.cell(r, original_columns[key]).value
            actual = ws.cell(r, columns[key]).value
            if actual != expected:
                edits.append({"pair_id": pair, "field": key, "ai_import_value": expected, "user_final_value": actual})
    if digest(path) != before_sha:
        raise ValueError("Concurrent workbook update")
    shutil.copy2(path, backup)
    if digest(backup) != before_sha:
        raise ValueError("User backup hash mismatch")
    wb.save(temporary)
    reopened = load_workbook(temporary, data_only=False)
    saved = reopened[ws.title]
    if reopened.sheetnames != wb.sheetnames or any((reopened[t].sheet_state, list(reopened[t].values)) != v for t, v in sheets.items()):
        raise ValueError("Non-review sheets changed")
    for r, values in enumerate(raw_values, 1):
        for c, value in enumerate(values, 1):
            if r >= 3 and c == columns["reviewer"]:
                continue
            if saved.cell(r, c).value != value:
                raise ValueError("Metadata completion changed user scores/fixed content")
    validate_review_rows(saved, samples, versions, mapping)
    if digest(path) != before_sha:
        raise ValueError("Concurrent workbook update before replacement")
    os.replace(temporary, path)
    atomic_json(out / "human_review_confirmation_record.json", {
        "status": "user_confirmed_24_rows", "created_at_utc": now(), "user_message": "我已复核完成",
        "confirmed_rows": 24, "reviewer_metadata_completed_pairs": filled,
        "supplied_reviewer_values": [{"value": k, "count": v} for k, v in supplied_reviewers.items()],
        "metadata_completion_changes_scores_preferences_or_notes": False,
        "user_saved_workbook": str(backup.relative_to(ROOT)).replace("\\", "/"), "user_saved_sha256": before_sha,
        "completed_workbook_sha256": digest(path), "user_edits_vs_imported_ai_draft": edits,
        "provenance": "AI-assisted initial judgments, Codex amendment, user final review/confirmation; not independent double-human review"})
    # Existing frozen summarizer saves the anonymous aggregate BEFORE applying model mapping.
    summarize()
    metrics = read_json(out / "human_metrics_unblinded.json")
    with (out / "human_review_final.csv").open(encoding="utf-8-sig", newline="") as handle:
        exported = list(csv.DictReader(handle))
    from src.generation.lora_validation_review import official_metrics
    for v in VARIANTS:
        records = [r for r in exported if r["variant"] == v]
        if official_metrics(records) != metrics["versions"][v]:
            raise ValueError("Exported review does not reproduce human metrics")
    blind = read_json(out / "blind_summary.json")
    if blind["workbook_sha256"] != digest(path) or metrics["blind_summary_sha256"] != digest(out / "blind_summary.json"):
        raise ValueError("Anonymous-before-unblind chain mismatch")
    totals, by_category = metrics["versions"], metrics["by_category"]
    def percent(value):
        return f"{value * 100:.2f}%"
    lines = ["# LoRA + RAG development24 最终人工复核报告", "", "状态：24件用户复核完成；先匿名汇总，再解盲。独立探索性validation实验，原RAG v2 P3未进入。", "",
        "## 人工三指标", "", "| 指标 | LoRA | LoRA + RAG |", "|---|---:|---:|"]
    for title, key in (("核心属性命中率", "core_attribute_hit_rate"), ("通顺率", "fluency_pass_rate"), ("事实错误样本率", "factual_error_sample_rate")):
        values = [percent(totals[v][key]) for v in VARIANTS]
        if key == "core_attribute_hit_rate":
            values = [f"{totals[v]['matched_attribute_total']}/{totals[v]['core_attribute_total']} = {value}" for v, value in zip(VARIANTS, values)]
        lines.append(f"| {title} | {values[0]} | {values[1]} |")
    lines.append(f"| 事实错误总数 / 平均每件 | {totals[VARIANTS[0]]['factual_error_total']} / {totals[VARIANTS[0]]['average_factual_errors_per_sample']:.2f} | {totals[VARIANTS[1]]['factual_error_total']} / {totals[VARIANTS[1]]['average_factual_errors_per_sample']:.2f} |")
    lines += ["", "## 匿名表达偏好（解盲后）", "", "| 维度 | LoRA更好 | LoRA + RAG更好 | 持平 |", "|---|---:|---:|---:|"]
    labels = ("标题表达质量", "卖点结构", "短详情自然度", "整体电商专业度")
    for p, label in zip(PREFERENCES, labels):
        votes = metrics["preferences"][p]
        lines.append(f"| {label} | {votes[VARIANTS[0]]} | {votes[VARIANTS[1]]} | {votes['tie']} |")
    lines += ["", "## 分品类人工核心覆盖", "", "| 二级品类（各3件） | LoRA | LoRA + RAG |", "|---|---:|---:|"]
    for c in by_category[VARIANTS[0]]:
        vals = [by_category[v][c] for v in VARIANTS]
        lines.append(f"| {c} | {vals[0]['matched_attribute_total']}/{vals[0]['core_attribute_total']} | {vals[1]['matched_attribute_total']}/{vals[1]['core_attribute_total']} |")
    by_pid = {v: {r["product_id"]: r for r in exported if r["variant"] == v} for v in VARIANTS}
    cases = []
    for pair in mapping:
        pid = pair["product_id"]
        a, b = by_pid[VARIANTS[0]][pid], by_pid[VARIANTS[1]][pid]
        if a["matched_attribute_count"] != b["matched_attribute_count"] or any(str(r[d]) == "1" for r in (a, b) for d in DIAGNOSTICS) or any(int(r["factual_error_count"]) for r in (a, b)):
            cases.append({"pair_id": pair["pair_id"], "product_id": pid, "category_l2": samples[pid]["category_l2"],
                          "LoRA": a, "LoRA_RAG": b})
    atomic_json(out / "human_diagnostic_cases.json", {"status": "confirmed_diagnostics_not_for_retuning", "cases": cases})
    lines += ["", "## 重点覆盖/结构诊断（不是全部事实错误case）", ""]
    for case in cases:
        pid, a, b = case["product_id"], case["LoRA"], case["LoRA_RAG"]
        lines += [f"- {pid}（{case['pair_id']} / {case['category_l2']}）：核心命中 LoRA {a['matched_attribute_count']}、组合 {b['matched_attribute_count']}；",
                  f"  LoRA备注：{a['review_notes'] or '无'}；组合备注：{b['review_notes'] or '无'}。"]
    automatic = read_json(out / "automatic_summary.json")
    lines += ["", "## 自动/工程指标（不替代人工）", "", "两个版本三任务均72/72解析成功；无空输出、无max-token触顶。组合Mandatory token注入228/228，完整率100%；身份核心事实另117/117验证。",
              "原核心分母134不变；源missing、placeholder或门控排除不得从分母删除，不计作Mandatory注入失败。自动字面覆盖为116/134对112/134，不能替代人工同义/格式/事实性判断。"]
    for v in VARIANTS:
        s = automatic["versions"][v]
        lines.append(f"{v}：三任务平均 {s['latency_seconds']['three_task_total']['mean']:.3f}s，P95 {s['latency_seconds']['three_task_total']['p95']:.3f}s；峰值allocated {s['peak_allocated_mib']:.1f}MiB。")
    delta = totals[VARIANTS[1]]["core_attribute_hit_rate"] - totals[VARIANTS[0]]["core_attribute_hit_rate"]
    lines += ["", "## 结论与限制", "", f"组合相对LoRA的人工核心覆盖差值为 {delta * 100:+.2f} 个百分点。两版人工事实错误和通顺表现见上表；本轮不预设RAG一定修复输出遗漏。",
        "本批已观察development24、每类仅3件，非新holdout/正式test，不能外推总体提升或总体零幻觉。多数偏好持平，不据少量偏好胜负宣称稳定表达收益。",
        "人工结果来自ChatGPT初审、Codex修正和用户最终复核确认，不是独立双人标注；A/B隐藏版本名不等于严格双盲，实际输入/输出结构可能透露条件。",
        "组合增量包含四区事实及必要接口适配（系统约束可靠核心属性→可靠属性），不是纯检索算法的因果隔离；对照延迟是历史测量，不把时间差完全归因于RAG。",
        "保持冻结LoRA/policy/Prompt/解码/seed/模型配置不变；不根据本轮诊断再调参或重跑。历史自动阶段报告和manifest保留原样，初始工作簿hash由pre_ai_import备份承接。",
        "阶段结束；无新模型调用、训练、test或32条holdout读取/生成，无commit/push。建议仅把组合作为探索性验证记录；是否用于Demo由用户另行确认，不自行替换当前默认方案。",
        "", "## 归档", "", "human_metrics_unblinded.json：总体/分品类三指标及偏好；human_review_final.csv：48份最终文案评分；human_diagnostic_cases.json：诊断；human_review_confirmation_record.json：用户确认与原始保存备份；final_review_manifest.json和final_review_checksums.sha256：完整hash关系。"]
    (out / "final_review_report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    _, _, frozen = verify_protocol()
    paths = {resolve(item["path"]) for item in frozen["files"]}
    paths.update(p for p in out.rglob("*") if p.is_file())
    paths.update((Path(__file__).resolve(), resolve("scripts/import_lora_rag_ai_review.py")))
    files = [{"path": str(p.relative_to(ROOT)).replace("\\", "/"), "sha256": digest(p)} for p in sorted(paths)]
    manifest_path = out / "final_review_manifest.json"
    atomic_json(manifest_path, {"status": "completed_user_confirmed_development24_integration", "created_at_utc": now(),
        "products": 24, "reviewed_documents": 48, "parameter_updates": 0,
        "project_test_read": False, "holdout_archive_read": False, "holdout_generation": False,
        "rag_v2_historical_holdout_status": "frozen_not_executed", "p3_entered": False,
        "human_provenance": "AI-assisted draft, Codex amendment, user final review/confirmation",
        "protocol_manifest_sha256": digest(out / "protocol_manifest.json"), "workbook_sha256": digest(path),
        "anonymous_summary_sha256": digest(out / "blind_summary.json"), "metrics_sha256": digest(out / "human_metrics_unblinded.json"),
        "final_report_sha256": digest(out / "final_review_report.md"), "user_edits_vs_imported_ai_draft": edits, "files": files})
    files_with_manifest = files + [{"path": str(manifest_path.relative_to(ROOT)).replace("\\", "/"), "sha256": digest(manifest_path)}]
    (out / "final_review_checksums.sha256").write_text("".join(f"{item['sha256']}  {item['path']}\n" for item in files_with_manifest), encoding="utf-8")
    print(json.dumps({"status": "completed", "reviewer_metadata_filled": len(filled), "user_label_edits": len(edits),
                      "versions": totals, "preferences": metrics["preferences"], "archive_files": len(files)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
