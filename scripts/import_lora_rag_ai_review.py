"""Import the audited ChatGPT draft into the original sealed development24 workbook.

Does not load models, apply the anonymous mapping, or aggregate official results.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
SPECIFICATION = ROOT / "reports/generation/lora_rag/lora_rag_v2_development24_v1/ai_review_amendments_v1.json"
SHEET = "盲评成对复核"
PREFERENCES = ("title_quality", "selling_points_structure", "short_description_naturalness", "overall_ecommerce_professionalism")
METRICS = ("matched_attribute_count", "fluency_pass", "factual_error_count")
DIAGNOSTICS = ("field_stacking", "mechanical_template", "unsupported_evaluation_effect_scenario", "numeric_range_distortion", "identity_category_error")
ALLOWED = set(PREFERENCES) | {f"{s}_{f}" for s in ("A", "B") for f in (*METRICS, *DIAGNOSTICS, "notes")} | {"reviewer", "review_confirmed", "pair_notes"}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def sheet_snapshot(sheet):
    # Hash sealed values without interpreting or exposing model identities.
    payload = json.dumps(list(sheet.values), ensure_ascii=False, default=str).encode("utf-8")
    return sheet.sheet_state, hashlib.sha256(payload).hexdigest()


def validate_draft(ws, columns, products, denominator):
    rows = [dict(zip(columns, r)) for r in ws.iter_rows(min_row=3, values_only=True)]
    if len(rows) != products or len({r["product_id"] for r in rows}) != products:
        raise ValueError("Wrong/duplicate draft product scope")
    if sum(r["core_attribute_count"] for r in rows) != denominator:
        raise ValueError("Original core denominator changed")
    for row in rows:
        if any(row[k] not in ("A更好", "B更好", "持平") for k in PREFERENCES):
            raise ValueError("Invalid preference")
        for side in ("A", "B"):
            for field in (*METRICS, *DIAGNOSTICS):
                value = row[f"{side}_{field}"]
                limit = row["core_attribute_count"] if field == "matched_attribute_count" else None if field == "factual_error_count" else 1
                if type(value) is not int or value < 0 or (limit is not None and value > limit):
                    raise ValueError(f"Invalid score: {row['pair_id']} {side}_{field}")
            if (row[f"{side}_factual_error_count"] or any(row[f"{side}_{k}"] for k in DIAGNOSTICS)) and not row[f"{side}_notes"]:
                raise ValueError("Positive errors/diagnoses require evidence notes")
        if not row["reviewer"] or row["review_confirmed"] is not None:
            raise ValueError("AI draft must await user confirmation")


def import_draft(spec, *, dry_run=False):
    source = Path(spec["source"])
    target = ROOT / spec["target"]
    backup = target.with_name(target.stem + ".pre_ai_import.xlsx")
    archived_source = target.with_name(target.stem + ".chatgpt_rechecked_source.xlsx")
    temporary = target.with_name(target.stem + ".ai_import.tmp.xlsx")
    manifest = target.parent / "ai_review_import_manifest_v1.json"
    if any(p.exists() for p in (backup, archived_source, temporary, manifest)):
        raise FileExistsError("Do not repeat an existing import")
    if digest(source) != spec["source_sha256"] or digest(target) != spec["initial_target_sha256"]:
        raise ValueError("Source or original changed: inspect before import")
    original, incoming_book = load_workbook(target, data_only=False), load_workbook(source, data_only=False)
    ws, incoming = original[SHEET], incoming_book[SHEET]
    columns = [c.value for c in ws[1]]
    indices = {k: i for i, k in enumerate(columns, 1)}
    editable = {k for k, i in indices.items() if not ws.cell(3, i).protection.locked}
    if ws.max_row != 26 or ws.max_column != 49 or incoming.max_row != 26 or incoming.max_column != 49 or editable != ALLOWED:
        raise ValueError("Original 24-pair/49-column boundary changed")
    if [c.value for c in incoming[1]] != columns or [c.value for c in incoming[2]] != [c.value for c in ws[2]]:
        raise ValueError("Incoming column definitions changed")
    untouched = {s.title: sheet_snapshot(s) for s in original if s.title != SHEET}
    fixed = {(r, c): (ws.cell(r, c).value, ws.cell(r, c).data_type, ws.cell(r, c).style_id)
             for r in range(1, 27) for c, k in enumerate(columns, 1) if r <= 2 or k not in editable}
    incoming_rows = {incoming.cell(r, indices["pair_id"]).value: r for r in range(3, 27)}
    pair_ids = {ws.cell(r, indices["pair_id"]).value for r in range(3, 27)}
    if len(incoming_rows) != 24 or len(pair_ids) != 24 or set(incoming_rows) != pair_ids:
        raise ValueError("Duplicate/missing anonymous pairs")
    if not set(spec["corrections"]) <= pair_ids:
        raise ValueError("Unknown corrected pair")
    amendments = []
    imported = 0
    for r in range(3, 27):
        pair = ws.cell(r, indices["pair_id"]).value
        other = incoming_rows[pair]
        if any(ws.cell(r, indices[k]).value != incoming.cell(other, indices[k]).value for k in columns if k not in editable):
            raise ValueError(f"Frozen source or output changed: {pair}")
        if any(ws.cell(r, indices[k]).value is not None for k in editable):
            raise ValueError(f"Existing user review would be overwritten: {pair}")
        correction = spec["corrections"].get(pair, {})
        if correction and str(ws.cell(r, indices["product_id"]).value) != correction["product_id"]:
            raise ValueError("Correction product ID mismatch")
        patches = dict(spec["metadata"])
        for field, item in correction.get("fields", {}).items():
            if field not in editable or incoming.cell(other, indices[field]).value != item["expected_source"]:
                raise ValueError("Unexpected correction field/source value")
            patches[field] = item["value"]
        for key in editable:
            supplied = incoming.cell(other, indices[key]).value
            value = patches.get(key, supplied)
            cell = ws.cell(r, indices[key])
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"
            imported += 1
            if value != supplied:
                reason = correction.get("fields", {}).get(key, {}).get("reason", "AI初审不得代替用户最终确认")
                amendments.append({"pair_id": pair, "product_id": ws.cell(r, indices["product_id"]).value,
                                   "field": key, "source_value": supplied, "written_value": value, "reason": reason})
                if key not in spec["metadata"]:
                    cell.fill = PatternFill("solid", fgColor="FCE4D6")
                    cell.comment = Comment("Codex复查修正，待用户确认。" + reason, "Codex review")
        ws.cell(r, indices["review_confirmed"]).comment = Comment("AI初审已填写；请你复核本行后填1，当前不计作正式人工确认。", "Codex review")
    validate_draft(ws, columns, spec["products"], spec["core_denominator"])
    score_changes = [a for a in amendments if a["field"].endswith(tuple("_" + m for m in METRICS))]
    if score_changes:
        raise ValueError("No formal-score changes were authorized by the audited amendments")
    result = {"status": "ai_draft_pending_user_confirmation", "created_at_utc": datetime.now(timezone.utc).isoformat(),
              "source": str(source), "source_sha256": spec["source_sha256"], "target": spec["target"],
              "products": 24, "documents": 48, "core_denominator": 134, "imported_editable_cells": imported,
              "formal_score_changes": 0, "preference_votes_changed": sum(a["field"] in PREFERENCES for a in amendments),
              "user_confirmed_rows": 0, "amended_fields_vs_supplied": amendments,
              "official_aggregation_performed": False, "unblinding_performed": False, "inference_performed": False}
    if dry_run:
        return result
    # Only the explicitly named review workbook is replaced, after recoverable backups.
    shutil.copy2(target, backup)
    shutil.copy2(source, archived_source)
    if digest(backup) != spec["initial_target_sha256"] or digest(archived_source) != spec["source_sha256"]:
        raise ValueError("Backup/source archival failed")
    original.save(temporary)
    reopened = load_workbook(temporary, data_only=False)
    if reopened.sheetnames != original.sheetnames or any(sheet_snapshot(reopened[t]) != s for t, s in untouched.items()):
        raise ValueError("Guide/sealed mapping/automatic hints changed")
    saved = reopened[SHEET]
    if any((saved.cell(r, c).value, saved.cell(r, c).data_type, saved.cell(r, c).style_id) != state for (r, c), state in fixed.items()):
        raise ValueError("Immutable cell content/type/style changed")
    validate_draft(saved, columns, spec["products"], spec["core_denominator"])
    if any(saved.cell(r, indices[k]).value != ws.cell(r, indices[k]).value for r in range(3, 27) for k in editable):
        raise ValueError("Imported labels changed on save")
    if digest(source) != spec["source_sha256"] or digest(target) != spec["initial_target_sha256"]:
        raise ValueError("Concurrent source/target change; refusing replacement")
    result.update(target_after_sha256=digest(temporary), original_backup=str(backup.relative_to(ROOT)).replace("\\", "/"),
                  backup_sha256=digest(backup), source_archive=str(archived_source.relative_to(ROOT)).replace("\\", "/"),
                  source_archive_sha256=digest(archived_source), immutable_cells_verified=True, all_original_sheets_preserved=True,
                  sealed_mapping_preserved=True, amendments_sha256=digest(SPECIFICATION), importer_sha256=digest(Path(__file__)))
    os.replace(temporary, target)
    manifest.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    spec = json.loads(SPECIFICATION.read_text(encoding="utf-8"))
    result = import_draft(spec, dry_run=args.dry_run)
    print(json.dumps({k: v for k, v in result.items() if k != "amended_fields_vs_supplied"}, ensure_ascii=False))
