"""Import the audited test100 AI draft without unblinding or finalizing results."""
from __future__ import annotations

import argparse
import json
import os
import shutil
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from import_lora_rag_ai_review import (
    ALLOWED, DIAGNOSTICS, METRICS, PREFERENCES, digest, sheet_snapshot, validate_draft,
)

ROOT = Path(__file__).resolve().parents[1]
DIRECTORY = ROOT / "reports/generation/lora_rag/lora_rag_project_test100_v1"
SPECIFICATION = DIRECTORY / "ai_review_amendments_v1.json"
TARGET = DIRECTORY / "lora_rag_project_test100_blind_review.xlsx"
SOURCE = Path("E:/WebDownload/lora_rag_project_test100_blind_review_rechecked.xlsx")
SHEET = "盲评成对复核"


def structure(ws):
    """Compare workbook layout/protection without interpreting anonymous identities."""
    return (ws.sheet_state, str(ws.freeze_panes), str(ws.protection),
            str(ws.data_validations), tuple(str(r) for r in ws.merged_cells.ranges),
            tuple((k, v.hidden, v.width) for k, v in ws.column_dimensions.items()))


def import_draft(spec, *, dry_run=False):
    source, target = Path(spec["source"]), ROOT / spec["target"]
    if source.resolve() != SOURCE.resolve() or target.resolve() != TARGET.resolve():
        raise ValueError("Only the explicitly authorized source/target may be imported")
    if spec["products"] != 100 or spec["core_denominator"] != 625:
        raise ValueError("Frozen test100 scope/denominator changed")
    if set(spec["metadata"]) != {"reviewer", "review_confirmed"} or spec["metadata"]["review_confirmed"] is not None:
        raise ValueError("AI draft cannot supply final user confirmation")
    backup = target.with_name(target.stem + ".pre_ai_import.xlsx")
    archive = target.with_name(target.stem + ".chatgpt_rechecked_source.xlsx")
    temporary = target.with_name(target.stem + ".ai_import.tmp.xlsx")
    manifest = DIRECTORY / "ai_review_import_manifest_v1.json"
    if any(p.exists() for p in (backup, archive, temporary, manifest)):
        raise FileExistsError("Refusing repeated import/overwrite of archived review")
    if digest(source) != spec["source_sha256"] or digest(target) != spec["initial_target_sha256"]:
        raise ValueError("Source/target hash changed; re-audit required")
    book, incoming_book = load_workbook(target), load_workbook(source)
    ws, incoming = book[SHEET], incoming_book[SHEET]
    columns = [c.value for c in ws[1]]
    indices = {k: i for i, k in enumerate(columns, 1)}
    if len(indices) != 49 or (ws.max_row, ws.max_column, incoming.max_row, incoming.max_column) != (102, 49, 102, 49):
        raise ValueError("Frozen 100-pair/49-column layout changed")
    if [c.value for c in incoming[1]] != columns or [c.value for c in incoming[2]] != [c.value for c in ws[2]]:
        raise ValueError("Column definitions changed")
    editable = {k for k, i in indices.items() if not ws.cell(3, i).protection.locked}
    if editable != ALLOWED:
        raise ValueError("Unexpected editable fields")
    if sheet_snapshot(book.worksheets[0]) != sheet_snapshot(incoming_book.worksheets[0]):
        raise ValueError("Incoming review instructions differ")
    untouched = {s.title: sheet_snapshot(s) for s in book if s.title != SHEET}
    structures = {s.title: structure(s) for s in book}
    fixed = {(r, c): (ws.cell(r, c).value, ws.cell(r, c).data_type, ws.cell(r, c).style_id)
             for r in range(1, 103) for c, k in enumerate(columns, 1) if r <= 2 or k not in editable}
    incoming_rows = {incoming.cell(r, indices["pair_id"]).value: r for r in range(3, 103)}
    pair_ids = {ws.cell(r, indices["pair_id"]).value for r in range(3, 103)}
    if len(incoming_rows) != 100 or len(pair_ids) != 100 or set(incoming_rows) != pair_ids:
        raise ValueError("Missing/duplicate anonymous pairs")
    if not set(spec["corrections"]) <= pair_ids:
        raise ValueError("Unknown correction pair")
    amendments = []
    for r in range(3, 103):
        pair = ws.cell(r, indices["pair_id"]).value
        other = incoming_rows[pair]
        if any(ws.cell(r, indices[k]).value != incoming.cell(other, indices[k]).value for k in columns if k not in editable):
            raise ValueError(f"Frozen source/output changed: {pair}")
        if any(ws.cell(r, indices[k]).value is not None for k in editable):
            raise ValueError(f"Would overwrite existing user review: {pair}")
        correction = spec["corrections"].get(pair, {})
        if correction and str(ws.cell(r, indices["product_id"]).value) != correction["product_id"]:
            raise ValueError(f"Correction product ID mismatch: {pair}")
        patches = dict(spec["metadata"])
        for field, item in correction.get("fields", {}).items():
            if field not in editable or incoming.cell(other, indices[field]).value != item["expected_source"]:
                raise ValueError(f"Unexpected correction field/source value: {pair} {field}")
            patches[field] = item["value"]
        for field in sorted(editable):
            supplied = incoming.cell(other, indices[field]).value
            value = patches.get(field, supplied)
            cell = ws.cell(r, indices[field])
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"  # Never interpret an imported string as an Excel formula.
            if value != supplied:
                reason = correction.get("fields", {}).get(field, {}).get("reason", "AI初审不替代用户最终确认。")
                amendments.append({"pair_id": pair, "product_id": ws.cell(r, indices["product_id"]).value,
                                   "field": field, "source_value": supplied, "written_value": value, "reason": reason})
                if field not in spec["metadata"]:
                    cell.fill = PatternFill("solid", fgColor="FCE4D6")
                    cell.comment = Comment("Codex复查修正，待用户确认。" + reason, "Codex review")
        for field, reason in correction.get("review_flags", {}).items():
            if field not in editable:
                raise ValueError("Review flag cannot target a frozen field")
            cell = ws.cell(r, indices[field])
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
            cell.comment = Comment("边界判断待用户裁决；当前保留原评分。" + reason, "Codex review")
        ws.cell(r, indices["review_confirmed"]).comment = Comment(
            "这是AI初审及复查草稿。请你复核本行后填1；当前未计作最终人工确认。", "Codex review")
    validate_draft(ws, columns, 100, 625)
    score_changes = [a for a in amendments if a["field"].endswith(tuple("_" + m for m in METRICS))]
    vote_changes = [a for a in amendments if a["field"] in PREFERENCES]
    if len(score_changes) != 3 or len(vote_changes) != 1:
        raise ValueError("Audited change scope must be exactly three formal cells and one preference")
    result = {"status": "ai_draft_pending_user_confirmation", "created_at_utc": datetime.now(timezone.utc).isoformat(),
              "source": str(source), "source_sha256": spec["source_sha256"], "target": spec["target"],
              "products": 100, "documents": 200, "core_denominator": 625,
              "imported_editable_cells": 100 * len(editable), "formal_score_changes": len(score_changes),
              "preference_votes_changed": len(vote_changes), "user_confirmed_rows": 0,
              "amended_fields_vs_supplied": amendments,
              "review_flags": {k: v["review_flags"] for k, v in spec["corrections"].items() if v.get("review_flags")},
              "official_aggregation_performed": False, "unblinding_performed": False, "inference_performed": False,
              "historical_final_reviews_modified": False}
    if dry_run:
        return result
    shutil.copy2(target, backup)
    shutil.copy2(source, archive)
    if digest(backup) != spec["initial_target_sha256"] or digest(archive) != spec["source_sha256"]:
        raise ValueError("Backup/archive validation failed")
    book.save(temporary)
    reopened = load_workbook(temporary)
    if reopened.sheetnames != book.sheetnames or any(sheet_snapshot(reopened[t]) != s for t, s in untouched.items()):
        raise ValueError("Untouched guide/mapping/automatic sheets changed")
    if any(structure(reopened[t]) != s for t, s in structures.items()):
        raise ValueError("Layout/protection/validation/visibility changed")
    saved = reopened[SHEET]
    if any((saved.cell(r, c).value, saved.cell(r, c).data_type, saved.cell(r, c).style_id) != state for (r, c), state in fixed.items()):
        raise ValueError("Frozen cell content/type/style changed")
    validate_draft(saved, columns, 100, 625)
    if any(saved.cell(r, indices[k]).value != ws.cell(r, indices[k]).value for r in range(3, 103) for k in editable):
        raise ValueError("Imported values changed during serialization")
    if digest(source) != spec["source_sha256"] or digest(target) != spec["initial_target_sha256"]:
        raise ValueError("Concurrent source/target change; refusing replacement")
    result.update(target_after_sha256=digest(temporary), original_backup=str(backup.relative_to(ROOT)).replace("\\", "/"),
                  backup_sha256=digest(backup), source_archive=str(archive.relative_to(ROOT)).replace("\\", "/"),
                  source_archive_sha256=digest(archive), immutable_cells_verified=True, all_original_sheets_preserved=True,
                  sealed_mapping_preserved=True, amendments_sha256=digest(SPECIFICATION), importer_sha256=digest(Path(__file__)),
                  helper_sha256=digest(ROOT / "scripts/import_lora_rag_ai_review.py"))
    os.replace(temporary, target)
    manifest.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    result = import_draft(json.loads(SPECIFICATION.read_text(encoding="utf-8")), dry_run=args.dry_run)
    print(json.dumps({k: v for k, v in result.items() if k != "amended_fields_vs_supplied"}, ensure_ascii=False))
