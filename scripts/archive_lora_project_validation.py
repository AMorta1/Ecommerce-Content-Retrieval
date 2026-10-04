"""Verify and archive P7 engineering/latency results, never infer human metrics."""
import csv
import json
import math
import statistics
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.generation.lora_validation import TASK_TYPES, load_jsonl, verify_protocol
from src.generation.lora_validation_review import METRICS, PREFERENCES, paired_data, review_columns
from src.generation.lora_validation_review import automatic_hints
from scripts.review_lora_project_validation_quick32 import verify_sample
from src.generation.lora_training import atomic_json, digest, now, read_json, resolve


def percentile(values, quantile):
    ordered = sorted(values)
    index = (len(ordered) - 1) * quantile
    low, high = math.floor(index), math.ceil(index)
    return ordered[low] + (ordered[high] - ordered[low]) * (index - low)


def main():
    from openpyxl import load_workbook
    cfg, frozen = verify_protocol()
    _, samples, versions, _ = paired_data()
    _, quick_cfg, review_out, mapping = verify_sample()
    out = resolve(cfg["output"])
    if (out / "automatic_phase_manifest.json").exists():
        raise FileExistsError("Refusing to overwrite automatic-phase archive")
    review_manifest = read_json(review_out / "review_manifest.json")
    for entry in review_manifest["files"]:
        if digest(resolve(entry["path"])) != entry["sha256"]:
            raise ValueError("Frozen generation file changed")
    workbook = review_out / "project_validation32_blind_review.xlsx"
    if digest(workbook) != review_manifest["initial_workbook_sha256"]:
        raise ValueError("Initial workbook changed before engineering acceptance")
    wb = load_workbook(workbook, read_only=False)
    ws = wb["盲评成对复核"]
    columns = {key: i for i, (key, _) in enumerate(review_columns(), 1)}
    if ws.max_row != 34 or wb["匿名映射_汇总后解盲"].sheet_state != "veryHidden":
        raise ValueError("Workbook coverage/anonymity mismatch")
    for row in range(3, 35):
        if ws.cell(row, columns["product_id"]).data_type != "s":
            raise ValueError("Product IDs must be preserved as text")
        for field in [*PREFERENCES, "reviewer", "review_confirmed", *[f"{side}_{m}" for side in ("A", "B") for m in METRICS]]:
            if ws.cell(row, columns[field]).value is not None:
                raise ValueError("Human fields must not be prefilled")
    wb.close()
    summaries, latency_rows = {}, []
    for variant in ("Base", "LoRA"):
        tasks = load_jsonl(out / variant.lower() / "task_outputs.jsonl")
        if len(tasks) != 600 or len({(t["product_id"], t["task_type"]) for t in tasks}) != 600:
            raise ValueError("Independent task coverage mismatch")
        for item in tasks:
            expected = samples[item["product_id"]]["tasks"][item["task_type"]]
            if any(item[k] != expected[k] for k in ("seed", "prompt_sha256", "preflight")):
                raise ValueError("Generation does not match frozen paired inputs")
            if not math.isfinite(item["latency_seconds"]) or item["latency_seconds"] <= 0:
                raise ValueError("Invalid latency")
        timings = {}
        for task in (*TASK_TYPES, "three_task_total"):
            values = [p["latency_seconds"][task] for p in versions[variant].values()]
            timings[task] = {"samples": len(values), "mean_seconds": statistics.mean(values),
                "p50_seconds": percentile(values, .5), "p95_seconds": percentile(values, .95)}
        for pid, product in versions[variant].items():
            latency_rows.append({"product_id": pid, "variant": variant, **product["latency_seconds"]})
        report = read_json(out / variant.lower() / "run_report.json")
        hints = [automatic_hints(samples[pid], product) for pid, product in versions[variant].items()]
        summaries[variant] = {"products": 200, "task_calls": 600, "wall_seconds": report["wall_seconds"],
            "latency": timings, "parse_success_by_task_NOT_FLUENCY": {t: sum(i["parsed"]["success"] for i in tasks if i["task_type"] == t) for t in TASK_TYPES},
            "assembled_structure_success_NOT_FLUENCY": sum(p["assembled"]["structure_success"] for p in versions[variant].values()),
            "hit_token_limit_tasks": sum(i["hit_max_new_tokens"] for i in tasks),
            "empty_task_outputs": sum(not i["raw_text"] for i in tasks),
            "peak_allocated_mib": max(i["memory"]["peak_allocated_mib"] for i in tasks),
            "peak_reserved_mib": max(i["memory"]["peak_reserved_mib"] for i in tasks),
            "model_parameters_unchanged": report["model_parameters_unchanged"],
            "literal_coverage_NOT_HUMAN": {"matched_fields": sum(h["literal_matched_count_NOT_HUMAN"] for h in hints),
                "core_attribute_total": frozen["source_statistics"]["core_attribute_total"],
                "ratio": sum(h["literal_matched_count_NOT_HUMAN"] for h in hints) / frozen["source_statistics"]["core_attribute_total"]},
            "claim_terms_REVIEW_ONLY_samples": sum(bool(h["claim_terms_REVIEW_ONLY"]) for h in hints),
            "number_terms_REVIEW_ONLY_samples": sum(bool(h["number_terms_REVIEW_ONLY"]) for h in hints)}
    a = read_json(out / "inference_base.json")
    b = read_json(out / "inference_lora.json")
    if {k for k in a if a[k] != b[k]} != {"variant", "adapter"}:
        raise ValueError("Inference configs differ beyond adapter identity")
    atomic_json(out / "automatic_phase_summary.json", {"status": "awaiting_human_review_NOT_FINAL_P7",
        "source_statistics": frozen["source_statistics"], "versions": summaries,
        "human_official_metrics": None, "human_paired_preferences": None,
        "human_sample": {"products": 32, "by_category": {c: 4 for c in sorted({p['category_l2'] for p in mapping})},
            "seed": quick_cfg["seed"], "core_attribute_total": sum(samples[p["product_id"]]["core_attribute_count"] for p in mapping),
            "status": "frozen_awaiting_human_review", "selection": "source_only_stratified_not_output_selected"},
        "paired_seed_prompt_preflight_checks": 600, "input_and_decoding_same": True,
        "first_call_cold_start_included": True, "old_single_call_baseline_compared": False,
        "project_test_read": False, "rag_holdout_files_read": False, "rag_v2_generation": False, "parameter_updates": 0})
    with (out / "latency_by_product.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(latency_rows[0]))
        writer.writeheader()
        writer.writerows(latency_rows)
    paths = sorted(p for p in out.rglob("*") if p.is_file())
    paths += [resolve("scripts/archive_lora_project_validation.py"), resolve("configs/lora_project_validation_quick32_v1.json"),
              resolve("scripts/review_lora_project_validation_quick32.py"), resolve("tests/test_lora_validation_quick32.py"),
              resolve("docs/lora_project_validation_quick32_v1.md")]
    items = [{"path": str(p.relative_to(resolve("."))).replace("\\", "/"), "sha256": digest(p)} for p in paths]
    atomic_json(out / "automatic_phase_manifest.json", {"status": "automatic_generation_complete_human_review_pending",
        "created_at_utc": now(), "files": items, "code_at_inference": "protocol_manifest.json", "workbook_hash_role": "initial_blank_review_only_expected_to_change_by_human",
        "holdout_boundary": cfg["holdout_boundary_user_confirmed"], "no_test_execution": True})
    (out / "automatic_phase_checksums.sha256").write_text("".join(f"{i['sha256']}  {i['path']}\n" for i in items)
        + f"{digest(out / 'automatic_phase_manifest.json')}  {str((out / 'automatic_phase_manifest.json').relative_to(resolve('.'))).replace(chr(92), '/')}\n", encoding="utf-8")
    print(json.dumps({"status": "automatic_generation_complete_human_review_pending", "versions": summaries}, ensure_ascii=False))


if __name__ == "__main__":
    main()
