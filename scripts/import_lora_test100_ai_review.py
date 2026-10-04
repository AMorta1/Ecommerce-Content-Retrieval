"""Import the reviewed AI draft into the original sealed test100 workbook.

Only unlocked review cells are written. No model, protocol preparation,
aggregation, model-identity lookup or unblinding is performed.
"""
from __future__ import annotations

import hashlib
import json
import os
import shutil
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
AMENDMENTS = ROOT / "reports/generation/lora/project_test100_v1/ai_review_amendments_v1.json"
SHEET = "盲评成对复核"
PREFERENCES = ("title_quality", "selling_points_structure", "short_description_naturalness", "overall_ecommerce_professionalism")
METRICS = ("matched_attribute_count", "fluency_pass", "factual_error_count")
DIAGNOSTICS = ("field_stacking", "mechanical_template", "unsupported_evaluation_effect_scenario", "numeric_range_distortion", "identity_category_error")


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def sheet_digest(sheet):
    # Compare sealed content without exposing or applying any model mapping.
    encoded = json.dumps(list(sheet.values), ensure_ascii=False, default=str).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def validate_rows(sheet, headers):
    reviewed = []
    for row in sheet.iter_rows(min_row=3, values_only=True):
        d = dict(zip(headers, row))
        if any(d[key] not in ("A更好", "B更好", "持平") for key in PREFERENCES):
            raise ValueError(f"Invalid preference: {d['pair_id']}")
        for side in ("A", "B"):
            for field in (*METRICS, *DIAGNOSTICS):
                value = d[f"{side}_{field}"]
                if type(value) is not int or value < 0:
                    raise ValueError(f"Invalid score: {d['pair_id']} {side}_{field}")
                limit = d["core_attribute_count"] if field == "matched_attribute_count" else (None if field == "factual_error_count" else 1)
                if limit is not None and value > limit:
                    raise ValueError(f"Score out of bounds: {d['pair_id']} {side}_{field}")
            if (d[f"{side}_factual_error_count"] or any(d[f"{side}_{f}"] for f in DIAGNOSTICS)) and not d[f"{side}_notes"]:
                raise ValueError(f"Evidence missing: {d['pair_id']} {side}")
        if not d["reviewer"] or d["review_confirmed"] != 0:
            raise ValueError("AI draft must remain pending user confirmation")
        reviewed.append(d)
    if len(reviewed) != 100 or sum(d["core_attribute_count"] for d in reviewed) != 625:
        raise ValueError("Review scope/core denominator changed")
    return reviewed


def main():
    specification = json.loads(AMENDMENTS.read_text(encoding="utf-8"))
    source = Path(specification["source"])
    target = ROOT / specification["target"]
    backup = target.with_name(target.stem + ".pre_ai_import.xlsx")
    manifest_path = target.parent / "ai_review_import_manifest_v1.json"
    temporary = target.with_name(target.stem + ".ai_import.tmp.xlsx")
    if any(p.exists() for p in (backup, manifest_path, temporary)):
        raise FileExistsError("No repeated import or overwrite of import archive")
    if digest(source) != specification["source_sha256"] or digest(target) != specification["initial_target_sha256"]:
        raise ValueError("Source/blank original changed; inspect before importing")

    original = load_workbook(target, data_only=False)
    supplied = load_workbook(source, data_only=False)
    ws, incoming = original[SHEET], supplied[SHEET]
    headers = [c.value for c in ws[1]]
    if ws.max_row != 102 or ws.max_column != 45 or [c.value for c in incoming[1]] != headers or incoming.max_row != 102:
        raise ValueError("Expected original 100-pair, 45-column structure")
    if any(c.value != incoming.cell(2, c.column).value for c in ws[2]):
        raise ValueError("Column descriptions changed")
    indices = {key: i for i, key in enumerate(headers, 1)}
    editable = {key for key, i in indices.items() if not ws.cell(3, i).protection.locked}
    allowed = set(PREFERENCES) | {f"{s}_{f}" for s in ("A", "B") for f in (*METRICS, *DIAGNOSTICS, "notes")} | {"reviewer", "review_confirmed", "pair_notes"}
    if editable != allowed:
        raise ValueError("Editable boundary differs from frozen workbook")
    untouched = {s.title: (s.sheet_state, sheet_digest(s)) for s in original if s.title != SHEET}
    fixed = {(r, c): (ws.cell(r, c).value, ws.cell(r, c).data_type, ws.cell(r, c).style_id)
             for r in range(1, 103) for c, key in enumerate(headers, 1) if r <= 2 or key not in editable}
    supplied_rows = {incoming.cell(r, indices["pair_id"]).value: r for r in range(3, 103)}
    pair_ids = {ws.cell(r, indices["pair_id"]).value for r in range(3, 103)}
    if len(supplied_rows) != 100 or len(pair_ids) != 100 or set(supplied_rows) != pair_ids:
        raise ValueError("Duplicate/missing pair IDs")
    if not set(specification["corrections"]).issubset(pair_ids):
        raise ValueError("Unknown correction pair")
    changes = []
    revised_pairs, revised_scores, changed_preferences = set(), set(), 0
    imported = 0
    for r in range(3, 103):
        pid = ws.cell(r, indices["pair_id"]).value
        other = supplied_rows[pid]
        if any(ws.cell(r, indices[k]).value != incoming.cell(other, indices[k]).value for k in headers if k not in editable):
            raise ValueError(f"Frozen source/output differs: {pid}")
        if any(ws.cell(r, indices[k]).value is not None for k in editable):
            raise ValueError(f"Existing user review would be overwritten: {pid}")
        values = {k: incoming.cell(other, indices[k]).value for k in editable}
        patches = dict(specification["metadata"])
        correction = specification["corrections"].get(pid, {})
        for side, item in correction.items():
            if side not in ("A", "B"):
                raise ValueError("Only anonymous sides may be patched")
            parts = []
            if "error_points" in item:
                points = item["error_points"]
                if len(points) != len(set(points)) or any(not isinstance(p, str) or not p for p in points):
                    raise ValueError(f"Invalid error ledger: {pid} {side}")
                patches[f"{side}_factual_error_count"] = len(points)
                parts.append(f"独立错误{len(points)}点：" + ("；".join(f"{n}. {p}" for n, p in enumerate(points, 1)) if points else "未发现事实错误。"))
            for key, value in item.get("overrides", {}).items():
                if key not in (*METRICS, *DIAGNOSTICS):
                    raise ValueError(f"Forbidden override: {key}")
                patches[f"{side}_{key}"] = value
            parts.append(item["note"])
            patches[f"{side}_notes"] = "AI复查：" + " ".join(parts)
        if any(k.endswith(tuple("_" + m for m in METRICS)) and values[k] != v for k, v in patches.items()):
            combined = {**values, **patches}
            denominator = ws.cell(r, indices["core_attribute_count"]).value
            patches["pair_notes"] = (f"AI复核后：核心命中A {combined['A_matched_attribute_count']}/{denominator}、B {combined['B_matched_attribute_count']}/{denominator}；"
                                     f"独立事实错误A {combined['A_factual_error_count']}、B {combined['B_factual_error_count']}；表达偏好沿用原稿，待用户确认。")
        for key in editable:
            old, value = values[key], patches.get(key, values[key])
            cell = ws.cell(r, indices[key])
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"
            imported += 1
            if old != value:
                changes.append({"pair_id": pid, "product_id": ws.cell(r, indices["product_id"]).value,
                                "field": key, "from_user_file": old, "written_value": value})
                if key not in specification["metadata"]:
                    revised_pairs.add(pid)
                    cell.fill = PatternFill("solid", fgColor="FCE4D6")
                    cell.comment = Comment("Codex复核修改；AI草稿，待用户确认。依据见本行notes和ai_review_amendments_v1.json。", "Codex review")
                    if key.endswith(tuple("_" + m for m in METRICS)):
                        revised_scores.add(pid)
                if key in PREFERENCES:
                    changed_preferences += 1
    if changed_preferences:
        raise ValueError("Preference import unexpectedly altered votes")
    validate_rows(ws, headers)
    # No mutation above has touched disk; keep the complete original first.
    shutil.copy2(target, backup)
    if digest(backup) != specification["initial_target_sha256"]:
        raise ValueError("Backup failed")
    original.save(temporary)
    reopened = load_workbook(temporary, data_only=False)
    for title, snapshot in untouched.items():
        if (reopened[title].sheet_state, sheet_digest(reopened[title])) != snapshot:
            raise ValueError("Sealed mapping/guide/automatic sheet changed")
    for (r, c), value in fixed.items():
        cell = reopened[SHEET].cell(r, c)
        if (cell.value, cell.data_type, cell.style_id) != value:
            raise ValueError(f"Locked cell changed: {r} {c}")
    validate_rows(reopened[SHEET], headers)
    if digest(source) != specification["source_sha256"] or digest(target) != specification["initial_target_sha256"]:
        raise ValueError("Source/original changed during import")
    reopened.close()
    original.close()
    supplied.close()
    os.replace(temporary, target)
    result = {
        "status": specification["status"], "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "source_path": str(source), "source_sha256": digest(source), "target_path": str(target), "target_sha256": digest(target),
        "blank_backup_path": str(backup), "blank_backup_sha256": digest(backup),
        "amendments_sha256": digest(AMENDMENTS), "import_script_sha256": digest(Path(__file__)),
        "products": 100, "model_documents": 200, "core_attribute_total": 625, "imported_editable_cells": imported,
        "revised_pairs_excluding_metadata": len(revised_pairs), "pairs_with_formal_score_corrections": len(revised_scores),
        "preference_votes_changed": changed_preferences, "user_confirmed_rows": 0,
        "immutable_cells_verified": True, "all_original_sheets_preserved": True, "sealed_mapping_preserved": True,
        "original_delivery_workbook_hash_preserved_in_backup": True,
        "official_aggregation_performed": False, "unblinding_performed": False, "inference_performed": False,
        "amended_fields_vs_supplied": changes,
    }
    manifest_path.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: result[k] for k in ("status", "products", "imported_editable_cells", "revised_pairs_excluding_metadata", "pairs_with_formal_score_corrections", "preference_votes_changed", "user_confirmed_rows", "immutable_cells_verified", "all_original_sheets_preserved", "sealed_mapping_preserved")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
