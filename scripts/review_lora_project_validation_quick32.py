"""Freeze a source-only stratified32 review; workbook/summary never run test."""
import argparse
import csv
import hashlib
import json
import math
import statistics
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.generation.lora_validation import load_jsonl, verify_protocol
from src.generation.lora_validation_review import (
    CHINESE, CHOICES, DIAGNOSTICS, GUIDE, METRICS, PREFERENCES, automatic_hints,
    fixed_values, official_metrics, paired_data, review_columns,
)
from src.generation.lora_training import atomic_json, digest, now, read_json, resolve
from src.generation.evaluation import parse_binary

CONFIG = "configs/lora_project_validation_quick32_v1.json"
SCRIPT = "scripts/review_lora_project_validation_quick32.py"
DOCUMENT = "docs/lora_project_validation_quick32_v1.md"


def choose(records, cfg):
    groups = defaultdict(list)
    seen = set()
    for row in records:
        pid = str(row["product_id"])
        if row["split"] != "validation" or pid in seen:
            raise ValueError("Non-validation/duplicate source")
        seen.add(pid)
        groups[row["category_l2"]].append(pid)
    if len(seen) != 200 or len(groups) != 8 or any(len(v) != 25 for v in groups.values()):
        raise ValueError("Expected frozen 8x25 project_validation")
    rank = lambda namespace, seed, category, pid: hashlib.sha256(f"{namespace}:{seed}:{category}:{pid}".encode()).hexdigest()
    selected, mapping = [], []
    for category, ids in sorted(groups.items()):
        ids = sorted(ids, key=lambda p: rank("p7_quick32_v1", cfg["seed"], category, p))[:4]
        selected.extend({"product_id": pid, "category_l2": category} for pid in ids)
        assignment = sorted(ids, key=lambda p: rank("p7_quick32_ab", cfg["blinding_seed"], category, p))
        for i, pid in enumerate(assignment):
            mapping.append({"product_id": pid, "category_l2": category, "A": "Base" if i < 2 else "LoRA", "B": "LoRA" if i < 2 else "Base"})
    mapping.sort(key=lambda p: rank("p7_quick32_order", cfg["blinding_seed"], p["category_l2"], p["product_id"]))
    for i, pair in enumerate(mapping, 1):
        pair["pair_id"] = f"QV{i:02d}"
    return selected, mapping


def freeze():
    parent, _ = verify_protocol()
    cfg = read_json(resolve(CONFIG))
    if cfg["source"] != parent["source"]["path"] or cfg["source_sha256"] != parent["source"]["sha256"]:
        raise ValueError("Only frozen project_validation may be sampled")
    if cfg["project_test_read"] or cfg["rag_holdout_files_read"]:
        raise ValueError("Forbidden scope")
    selected, mapping = choose(load_jsonl(resolve(cfg["source"])), cfg)
    out = resolve(cfg["output"])
    if out.exists():
        raise FileExistsError("No resampling/overwrite")
    parent_out = resolve(parent["output"])
    if any((parent_out / name).exists() for name in ("human_metrics_unblinded.json", "blind_summary.json")):
        raise ValueError("Sampling must precede human quality results")
    out.mkdir()
    atomic_json(out / "sample_ids.json", {"status": "frozen_before_human_quality_review", "selected": selected})
    atomic_json(out / "blind_mapping.json", {"status": "sealed_until_review_complete", "pairs": mapping})
    files = [resolve(CONFIG), resolve(SCRIPT), resolve(DOCUMENT), resolve(cfg["source"]), parent_out / "protocol_manifest.json", out / "sample_ids.json", out / "blind_mapping.json"]
    atomic_json(out / "sampling_manifest.json", {"status": "frozen_before_human_quality_review", "created_at_utc": now(),
        "config": cfg, "products": 32, "by_category": dict(Counter(r["category_l2"] for r in selected)),
        "selection_uses": ["source_product_id", "source_category_l2"], "generated_text_or_human_results_used_for_selection": False,
        "timing": cfg["sampling_timing"], "code_snapshot_hex": resolve(SCRIPT).read_bytes().hex(),
        "files": [{"path": str(p.relative_to(resolve("."))).replace("\\", "/"), "sha256": digest(p)} for p in files]})
    print(json.dumps({"status": "sample32_frozen", "by_category": dict(Counter(r["category_l2"] for r in selected))}, ensure_ascii=False))


def verify_sample():
    parent, _ = verify_protocol()
    cfg = read_json(resolve(CONFIG))
    out = resolve(cfg["output"])
    manifest = read_json(out / "sampling_manifest.json")
    if cfg != manifest["config"]:
        raise ValueError("Sampling config changed")
    for item in manifest["files"]:
        if digest(resolve(item["path"])) != item["sha256"]:
            raise ValueError("Frozen sample/code/document changed")
    return parent, cfg, out, read_json(out / "blind_mapping.json")["pairs"]


def workbook():
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    parent, cfg, out, mapping = verify_sample()
    _, samples, versions, _ = paired_data()
    path = out / "project_validation32_blind_review.xlsx"
    if path.exists():
        raise FileExistsError("No workbook overwrite")
    wb = Workbook()
    guide = wb.active
    guide.title = "复核说明"
    instructions = ["本轮仅复核固定分层32件，每类4件、两模型64份文案。采样在人工质量结果前固定，没有依据输出挑样本。",
        "先读下面口径。每行源属性、核心清单、实际输入与A/B三任务同一行显示；黄色区域可填写，其余不改。",
        "若要调整列宽/行高，审阅→撤销工作表保护（无密码）。不要改事实、ID、分母或输出。"]
    instructions += [t for t in GUIDE if not any(x in t for x in ("200", "所有200", "blind_mapping.json"))]
    instructions += ["匿名映射封存在veryHidden页与本快速复核目录blind_mapping.json，完成前勿查看；需全32行填写完整再解盲。",
        "汇总仅代表该分层32件样本，不冒充全200人工结果。是否进入project_test由你另行确认，脚本不会执行test。"]
    guide.append(["顺序", "操作与口径"])
    for i, text in enumerate(instructions, 1):
        guide.append([i, text])
        guide.row_dimensions[i + 1].height = 48
    guide.column_dimensions["A"].width = 8
    guide.column_dimensions["B"].width = 140
    ws = wb.create_sheet("盲评成对复核")
    columns = review_columns()
    ws.append([k for k, _ in columns])
    ws.append([v for _, v in columns])
    indices = {k: i for i, (k, _) in enumerate(columns, 1)}
    fixed_keys = None
    for pair in mapping:
        values = fixed_values(samples[pair["product_id"]], pair, versions)
        fixed_keys = set(values)
        ws.append([values.get(k) for k, _ in columns])
    for row in ws:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if cell.row <= 2:
                cell.fill = PatternFill("solid", fgColor="244062")
                cell.font = Font(color="FFFFFF", bold=True)
            else:
                editable = columns[cell.column - 1][0] not in fixed_keys
                cell.protection = Protection(locked=not editable)
                cell.fill = PatternFill("solid", fgColor="FFF2CC" if editable else "E7E6E6")
    for key, index in indices.items():
        ws.column_dimensions[get_column_letter(index)].width = 55 if key.endswith(("title", "selling_points", "short_description", "complete")) else (44 if "attributes" in key or "notes" in key else 18)
        ws.cell(1, index).comment = Comment(columns[index - 1][1], "P7 quick32")
        if key in ("A_complete", "B_complete", "source_title", "input_quality_actions"):
            ws.column_dimensions[get_column_letter(index)].hidden = True
    for row in range(3, 35):
        ws.row_dimensions[row].height = 230
        ws.cell(row, indices["product_id"]).number_format = "@"
    for preference in PREFERENCES:
        dv = DataValidation(type="list", formula1='"A更好,B更好,持平"', allow_blank=True)
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(indices[preference])}3:{get_column_letter(indices[preference])}34")
    for side in ("A", "B"):
        for field in (*METRICS, *DIAGNOSTICS):
            dv = DataValidation(type="whole", operator="greaterThanOrEqual" if field == "factual_error_count" else "between",
                formula1="0", formula2=None if field == "factual_error_count" else ("$I3" if field == "matched_attribute_count" else "1"), allow_blank=True)
            dv.showErrorMessage = True
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(indices[f'{side}_{field}'])}3:{get_column_letter(indices[f'{side}_{field}'])}34")
    ws.freeze_panes = "F3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(columns))}34"
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    hints = wb.create_sheet("自动线索_非人工指标")
    hints.append(["pair_id", "product_id", "side", "diagnostic_ONLY"])
    for pair in mapping:
        for side in ("A", "B"):
            hints.append([pair["pair_id"], pair["product_id"], side, json.dumps(automatic_hints(samples[pair["product_id"]], versions[pair[side]][pair["product_id"]]), ensure_ascii=False)])
    hints.sheet_state = "hidden"
    sealed = wb.create_sheet("匿名映射_汇总后解盲")
    sealed.append(["pair_id", "product_id", "A", "B"])
    for pair in mapping:
        sealed.append([pair[k] for k in ("pair_id", "product_id", "A", "B")])
    sealed.sheet_state = "veryHidden"
    for row in guide:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)
    parent_out = resolve(parent["output"])
    outputs = [p for name in ("base", "lora") for p in (parent_out / name).iterdir() if p.is_file()]
    atomic_json(out / "review_manifest.json", {"status": "awaiting_human_blinded_review", "products": 32, "documents": 64,
        "initial_workbook_sha256": digest(path), "sampling_manifest_sha256": digest(out / "sampling_manifest.json"),
        "files": [{"path": str(p.relative_to(resolve("."))).replace("\\", "/"), "sha256": digest(p)} for p in outputs],
        "human_labels_prefilled": False})
    print(f"32-pair blinded workbook: {path}")


def summarize():
    from openpyxl import load_workbook
    _, _, out, mapping = verify_sample()
    _, samples, versions, _ = paired_data()
    manifest = read_json(out / "review_manifest.json")
    for item in manifest["files"]:
        if digest(resolve(item["path"])) != item["sha256"]:
            raise ValueError("Generation changed")
    if digest(out / "sampling_manifest.json") != manifest["sampling_manifest_sha256"]:
        raise ValueError("Sampling manifest changed")
    for name in ("blind_summary.json", "human_metrics_unblinded.json", "human_review_final.csv"):
        if (out / name).exists():
            raise FileExistsError("No human result overwrite")
    path = out / "project_validation32_blind_review.xlsx"
    wb = load_workbook(path, data_only=False)
    ws = wb["盲评成对复核"]
    columns = review_columns()
    if ws.max_row != 34 or [c.value for c in ws[1]] != [k for k, _ in columns]:
        raise ValueError("32-pair workbook coverage/columns changed")
    if list(wb["匿名映射_汇总后解盲"].values)[1:] != [tuple(p[k] for k in ("pair_id", "product_id", "A", "B")) for p in mapping]:
        raise ValueError("Anonymous mapping changed")
    number = lambda v: str(int(v)) if isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v) and int(v) == v else str(v).strip()
    rows = []
    for i, pair in enumerate(mapping, 3):
        row = {k: ws.cell(i, j).value for j, (k, _) in enumerate(columns, 1)}
        if any(row[k] != v for k, v in fixed_values(samples[pair["product_id"]], pair, versions).items()):
            raise ValueError("Frozen input/output changed")
        if number(row["review_confirmed"]) != "1" or not row["reviewer"]:
            raise ValueError("Incomplete human confirmation")
        if any(row[p] not in CHOICES for p in PREFERENCES):
            raise ValueError("Incomplete paired preferences")
        for side in ("A", "B"):
            for metric in METRICS:
                row[f"{side}_{metric}"] = number(row[f"{side}_{metric}"])
            official_metrics([{"product_id": row["product_id"], "core_attribute_count": row["core_attribute_count"], **{m: row[f"{side}_{m}"] for m in METRICS}}])
            for diagnosis in DIAGNOSTICS:
                if row[f"{side}_{diagnosis}"] is not None:
                    row[f"{side}_{diagnosis}"] = parse_binary(number(row[f"{side}_{diagnosis}"]), diagnosis, row["product_id"])
            if (int(row[f"{side}_factual_error_count"]) > 0 or any(row[f"{side}_{d}"] == 1 for d in DIAGNOSTICS)) and not row[f"{side}_notes"]:
                raise ValueError("Fact errors/diagnostics require evidence")
        rows.append(row)
    # Complete anonymous aggregation is written before applying model identities.
    atomic_json(out / "blind_summary.json", {"products": 32, "preferences": {p: dict(Counter(r[p] for r in rows)) for p in PREFERENCES}, "workbook_sha256": digest(path)})
    judgments, votes = {"Base": [], "LoRA": []}, {p: Counter() for p in PREFERENCES}
    exported = []
    for row, pair in zip(rows, mapping):
        for side in ("A", "B"):
            item = {"product_id": row["product_id"], "category_l2": samples[row["product_id"]]["category_l2"], "variant": pair[side],
                "pair_id": pair["pair_id"], "anonymous_side": side, "core_attribute_count": row["core_attribute_count"],
                **{m: row[f"{side}_{m}"] for m in METRICS}, **{d: row[f"{side}_{d}"] for d in DIAGNOSTICS},
                **{p: row[p] for p in PREFERENCES}, "reviewer": row["reviewer"], "notes": row[f"{side}_notes"] or "", "pair_notes": row["pair_notes"] or ""}
            judgments[pair[side]].append(item)
            exported.append(item)
        for p in PREFERENCES:
            votes[p]["tie" if row[p] == "持平" else pair[row[p][0]]] += 1
    atomic_json(out / "human_metrics_unblinded.json", {"status": "final_human_review_sample32_only", "products_per_model": 32,
        "official_metrics": {v: official_metrics(r) for v, r in judgments.items()},
        "paired_preferences": {p: {v: votes[p][v] for v in ("Base", "LoRA", "tie")} for p in PREFERENCES},
        "by_category": {v: {c: official_metrics([r for r in judgments[v] if r["category_l2"] == c]) for c in sorted({r["category_l2"] for r in judgments[v]})} for v in judgments},
        "diagnostics_NOT_FORMAL_METRICS": {v: {d: {"annotated": sum(r[d] is not None for r in judgments[v]), "positive": sum(r[d] == 1 for r in judgments[v])} for d in DIAGNOSTICS} for v in judgments},
        "sampling_manifest_sha256": digest(out / "sampling_manifest.json"), "workbook_sha256": digest(path),
        "blind_summary_sha256": digest(out / "blind_summary.json"), "project_test_run": False})
    with (out / "human_review_final.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(exported[0]))
        writer.writeheader()
        writer.writerows(exported)
    print("Complete32 human summary saved. No test execution.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--mode", choices=["freeze", "workbook", "summarize"], required=True)
    args = parser.parse_args()
    {"freeze": freeze, "workbook": workbook, "summarize": summarize}[args.mode]()
