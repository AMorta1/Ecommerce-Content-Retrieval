"""Additional fixed test100 comparison, reusing immutable development treatment.

This is NOT a fresh unseen test or the original RAG v2 P3. Old test labels are
never imported into the new anonymous review; only hash-verified control text is
reused as a comparator, never as model input.
"""
from __future__ import annotations

import copy
import csv
import hashlib
import importlib.metadata
import json
import math
import random
import subprocess
import sys
import time
from collections import Counter, defaultdict

from . import lora_rag_integration as treatment
from .lora_formal_test import verify_protocol as verify_test
from .lora_formal_review import extra_hints
from .lora_training import ROOT, atomic_json, digest, json_digest, memory, now, read_json, resolve
from .lora_validation import (TASK_TYPES, PACKAGE_NAMES, append_jsonl, assemble,
                              load_jsonl, model_snapshot, parse_task)
from .lora_validation_review import CHOICES, DIAGNOSTICS, GUIDE, METRICS, PREFERENCES, official_metrics
from .qwen import preflight_chat_prompt
from .qlora_memory import prepare_kbit_with_cpu_staged_large_layers
from .rag import build_fact_units, load_quality_audit

CONFIG = 'configs/lora_rag_project_test100_v1.json'
DOCUMENT = 'docs/lora_rag_project_test100_v1_protocol.md'
OUTPUT = 'reports/generation/lora_rag/lora_rag_project_test100_v1'
WORKBOOK = 'lora_rag_project_test100_blind_review.xlsx'
CODE = ('scripts/run_lora_rag_project_test.py', 'src/generation/lora_rag_project_test.py',
        'tests/test_lora_rag_project_test.py')
VARIANTS = treatment.VARIANTS
DEVELOPMENT = 'reports/generation/lora_rag/lora_rag_v2_development24_v1'
TEST = 'reports/generation/lora/project_test100_v1'
ALLOWED = {
    'development_protocol': DEVELOPMENT + '/protocol_manifest.json',
    'test_protocol': TEST + '/protocol_manifest.json',
    'test_archive': TEST + '/automatic_phase_manifest.json',
    'test_inputs': TEST + '/frozen_inputs.jsonl',
    'control_outputs': TEST + '/lora/product_outputs.jsonl',
    'control_tasks': TEST + '/lora/task_outputs.jsonl',
    'control_report': TEST + '/lora/run_report.json',
    'control_inference': TEST + '/inference_lora.json',
}
PROTOCOL_HASHES = {
    'development_protocol': '17d573d3bddc8fa564075abc14dcd94a2530a4300e30c09b642874a34addedc7',
    'test_protocol': 'e8754d37ed27789d02801b710184ee9a034cff2e6996a42da8a29eadde0f7850',
}
GUARDS = {k: False for k in ('parameter_updates', 'output_based_adjustment', 'rerun',
          'holdout_archive_read', 'holdout_generation', 'original_rag_v2_p3_entered',
          'old_outputs_or_labels_as_model_inputs')}


def checked_config():
    cfg = read_json(resolve(CONFIG))
    # Reject forbidden paths BEFORE hashing/opening any configurable input.
    for key, name in ALLOWED.items():
        if cfg[key]['path'] != name:
            raise ValueError(f'Forbidden path before read: {key}')
    if (cfg['scope'] != 'additional_fixed_comparison_on_previously_observed_test100'
            or cfg['output'] != OUTPUT or cfg['guards'] != GUARDS
            or cfg['review'] != {'products': 100, 'selection': 'all_original_frozen100_no_resampling',
                                 'blind_seed': 41009, 'prefill_human_labels': False}
            or cfg['generation'] != {'new_products': 100, 'new_task_calls': 300,
                'single_attempt': True, 'resume': False,
                'control': 'reuse_hash_verified_formal_LoRA_three_task_outputs'}):
        raise ValueError('Fixed test scope/guards changed')
    for key, sha in PROTOCOL_HASHES.items():
        if cfg[key]['sha256'] != sha:
            raise ValueError('Parent protocol hash changed')
    for key, name in ALLOWED.items():
        if digest(resolve(name)) != cfg[key]['sha256']:
            raise ValueError(f'Frozen parent input changed: {key}')
    dev_cfg, dev_parent, dev_frozen = treatment.verify_protocol()
    test_cfg, inherited, test_frozen = verify_test()
    if dev_parent != inherited:
        raise ValueError('Control and combination model/decoding/training provenance differ')
    archived = {e['path']: e['sha256'] for e in read_json(resolve(cfg['test_archive']['path']))['files']}
    for key in ('test_inputs', 'control_outputs', 'control_tasks', 'control_report', 'control_inference'):
        if archived.get(cfg[key]['path']) != cfg[key]['sha256']:
            raise ValueError('Control input/output not pinned by original test archive')
    if dev_frozen['environment'] != test_frozen['environment']:
        raise ValueError('Parent inference environments differ')
    return cfg, dev_cfg, inherited, dev_frozen, test_cfg, test_frozen


def environment():
    import torch
    return {'python': sys.version, 'packages': {n: importlib.metadata.version(n) for n in PACKAGE_NAMES},
            'cuda': torch.version.cuda}


def mapping_for(samples, seed=41009):
    """Balanced within each category (difference <=1), 50/50 overall, no output use."""
    groups = defaultdict(list)
    for s in samples:
        groups[s['category_l2']].append(s)
    rank = lambda ns, value: hashlib.sha256(f'{ns}:{seed}:{value}'.encode()).hexdigest()
    quotas = {c: len(rows) // 2 for c, rows in groups.items()}
    extra = len(samples) // 2 - sum(quotas.values())
    odd = sorted((c for c in groups if len(groups[c]) % 2), key=lambda c: rank('category', c))
    for c in odd[:extra]:
        quotas[c] += 1
    controls_a = {s['product_id'] for c, rows in groups.items()
                  for s in sorted(rows, key=lambda s: rank('assignment', s['product_id']))[:quotas[c]]}
    order = sorted(samples, key=lambda s: rank('order', s['product_id']))
    return [{'pair_id': f'CT{i:03d}', 'product_id': s['product_id'],
             'A': VARIANTS[0] if s['product_id'] in controls_a else VARIANTS[1],
             'B': VARIANTS[1] if s['product_id'] in controls_a else VARIANTS[0]}
            for i, s in enumerate(order, 1)]


def validate_control(samples, products, tasks, report):
    ids = {s['product_id'] for s in samples}
    by_id = {r['product_id']: r for r in products}
    if (len(samples) != 100 or len(ids) != 100 or len(products) != 100 or set(by_id) != ids
            or len(tasks) != 300 or len({(t['product_id'], t['task_type']) for t in tasks}) != 300
            or report['status'] != 'passed' or report['products'] != 100 or report['task_calls'] != 300
            or not report['adapter_loaded'] or not report['model_parameters_unchanged']):
        raise ValueError('Incomplete original frozen100 control')
    for sample in samples:
        treatment.verify_control(sample, by_id[sample['product_id']])
    for t in tasks:
        if t != by_id[t['product_id']]['tasks'][t['task_type']]:
            raise ValueError('Original independent control task differs from assembly')
    return by_id


def test_facts(records, audit, policy, source):
    """Apply frozen policy unchanged; scope normalization only in an audit COPY."""
    ids = [str(r['product_id']) for r in records]
    if len(ids) != 100 or len(set(ids)) != 100 or set(ids) != set(audit):
        raise ValueError('Must use exactly original frozen100 source and audit')
    facts = []
    for record in records:
        pid = str(record['product_id'])
        row = copy.deepcopy(audit[pid])
        if record['split'] != 'test' or row['scope'] != 'test100':
            raise ValueError('Forbidden source/audit split')
        row['scope'] = 'test'
        facts.extend(build_fact_units(record, row, policy, dataset_version='week1_v3',
                     source_path=source['path'], source_sha256=source['sha256']))
    return facts


def gated_context(sample, facts, policy):
    # The development cohort had no identity-blocked products. On the frozen
    # test cohort, keep blocked units as selector metadata (never injections),
    # so an empty four-zone context still carries its product ID. Same policy.
    if sample['source_quality'].get('identity_block'):
        if any(f['quality_status'] != 'blocked_conflict' for f in facts):
            raise ValueError('Identity-conflict facts must all remain blocked')
        context = treatment.select_rag_v2_context(facts, policy, top_k=3)
        treatment.assert_context(sample, context)
        if any(context[z] for z in ('identity_facts', 'mandatory_core_facts',
                                    'supplemental_facts', 'negative_constraint_facts')):
            raise ValueError('Identity conflict cannot inject facts')
        return context, [{'fact_id': f['fact_id'], 'field': f['canonical_field'],
                          'reason': 'already_withheld_in_frozen_P7_input'} for f in facts]
    return treatment.gated_context(sample, facts, policy)


def prepare():
    from transformers import AutoTokenizer
    cfg, dev_cfg, inherited, dev_frozen, test_cfg, test_frozen = checked_config()
    out = resolve(OUTPUT)
    if out.exists():
        raise FileExistsError('Never overwrite fixed test freeze')
    if environment() != test_frozen['environment']:
        raise ValueError('Frozen inference environment changed')
    samples = copy.deepcopy(load_jsonl(resolve(cfg['test_inputs']['path'])))
    by_id = validate_control(samples, load_jsonl(resolve(cfg['control_outputs']['path'])),
        load_jsonl(resolve(cfg['control_tasks']['path'])), read_json(resolve(cfg['control_report']['path'])))
    policy = read_json(resolve(dev_cfg['fact_policy']['path']))
    records = load_jsonl(resolve(test_cfg['source']['path']))
    if [str(r['product_id']) for r in records] != [s['product_id'] for s in samples]:
        raise ValueError('Original100 order changed')
    for r, s in zip(records, samples):
        if (r['attributes'] != s['source_attributes'] or r['generation_input']['attributes'] != s['evaluation_attributes']
                or len(s['evaluation_attributes']) != s['core_attribute_count']):
            raise ValueError('Original source/core denominator differs from frozen test')
    facts = test_facts(records, load_quality_audit(resolve(test_cfg['source_audit']['path']), 'test100'),
                       policy, test_cfg['source'])
    grouped = defaultdict(list)
    for f in facts:
        grouped[str(f['product_id'])].append(f)
    tokenizer = AutoTokenizer.from_pretrained(str(model_snapshot(inherited)), local_files_only=True)
    if json_digest(tokenizer.chat_template) != dev_frozen['chat_template_sha256']:
        raise ValueError('Frozen chat template changed')
    controls, disagreements = [], []
    for sample in samples:
        tick = time.perf_counter()
        context, withheld = gated_context(sample, grouped[sample['product_id']], policy)
        sample['rag_selection_seconds'] = time.perf_counter() - tick
        sample['rag_context'], sample['integration_gate_withheld'] = context, withheld
        sample['rag_tasks'] = {}
        for task in TASK_TYPES:
            messages = treatment.integration_messages(sample, task, context, dev_cfg)
            sample['rag_tasks'][task] = {'messages': messages, 'seed': sample['tasks'][task]['seed'],
                                       'prompt_sha256': json_digest(messages)}
            _, meta = treatment.preflight(tokenizer, sample, task, dev_cfg)
            sample['rag_tasks'][task]['preflight'] = meta
            _, control_meta = preflight_chat_prompt(tokenizer, sample['tasks'][task]['messages'],
                max_input_tokens=inherited['prompt']['max_input_tokens'],
                mandatory_lines=[json.dumps(sample['used_attributes'], ensure_ascii=False, sort_keys=False)])
            if control_meta != sample['tasks'][task]['preflight']:
                raise ValueError('Historical control preflight changed')
        disagreements.extend({'product_id': sample['product_id'], 'field': f['canonical_field'],
            'rag_quality': f['quality_status'], 'rag_reasons': f['quality_reasons'],
            'action': 'retain frozen control core unchanged; do not add this fact in RAG zones'}
            for f in grouped[sample['product_id']] if f['quality_status'] != 'eligible'
            and f['canonical_field'] in sample['used_attributes'])
        control = copy.deepcopy(by_id[sample['product_id']])
        control.update(source_variant=control['variant'], variant=VARIANTS[0], generation_reused=True,
                       rag_selection_seconds=0.0)
        controls.append(control)
    out.mkdir(parents=True)
    for filename, rows in (('frozen_inputs.jsonl', samples), ('test100_fact_units.jsonl', facts),
                           ('control_product_outputs.jsonl', controls)):
        for row in rows:
            append_jsonl(out / filename, row)
    # Independent control tasks remain verbatim, anchored to the original archive.
    (out / 'control_task_outputs.jsonl').write_bytes(resolve(cfg['control_tasks']['path']).read_bytes())
    atomic_json(out / 'sample_ids.json', {'project_test100': [s['product_id'] for s in samples],
        'selection': 'all original frozen100; no resampling', 'previously_observed': True})
    atomic_json(out / 'blind_mapping.json', {'status': 'sealed_until_complete_review',
        'pairs': mapping_for(samples, cfg['review']['blind_seed'])})
    atomic_json(out / 'inference_config.json', {'model': inherited['model'], 'adapter': inherited['adapter'],
        'decoding': inherited['decoding'], 'decoding_resolved': read_json(resolve(cfg['control_inference']['path']))['decoding_resolved'],
        'original_prompt': inherited['prompt'], 'integration_prompt': dev_cfg['prompt'], 'pairing': inherited['pairing'],
        'chat_template_sha256': dev_frozen['chat_template_sha256']})
    freeze = subprocess.run([sys.executable, '-m', 'pip', 'freeze'], capture_output=True, text=True, check=True).stdout
    (out / 'environment_freeze.txt').write_text(freeze, encoding='utf-8')
    git = lambda *args: subprocess.run(['git', *args], cwd=ROOT, capture_output=True, text=True,
                                      encoding='utf-8', check=True).stdout.strip()
    paths = {e['path'] for parent in (dev_frozen, test_frozen) for e in parent['files']}
    paths.update((*CODE, CONFIG, DOCUMENT, *(e['path'] for e in cfg.values() if isinstance(e, dict) and 'path' in e)))
    paths.update(p.relative_to(ROOT).as_posix() for p in out.iterdir() if p.is_file())
    stats = {'products': 100, 'by_category': dict(Counter(s['category_l2'] for s in samples)),
        'original_core_denominator': sum(s['core_attribute_count'] for s in samples),
        'source_quality': dict(Counter(s['source_quality']['status'] for s in samples)),
        'original_input_core_fields': sum(len(s['used_attributes']) for s in samples),
        'mandatory_core_facts_per_product_sum': sum(len(s['rag_context']['mandatory_core_facts']) for s in samples),
        'identity_core_facts_sum': sum(len(s['rag_context']['identity_mandatory_coverage_fact_ids']) for s in samples),
        'max_input_tokens': max(t['preflight']['input_token_count'] for s in samples for t in s['rag_tasks'].values())}
    if stats['original_core_denominator'] != 625:
        raise ValueError('Frozen100 denominator must remain625')
    manifest = {'version': cfg['version'], 'status': 'frozen_before_new_test_generation', 'created_at_utc': now(),
        'config': cfg, 'parent_config': inherited, 'treatment_config': dev_cfg,
        'git_head': git('rev-parse', 'HEAD'), 'git_status': git('-c', 'core.quotepath=false', 'status', '--short'),
        'code_original_bytes_hex': {p: resolve(p).read_bytes().hex() for p in CODE},
        'files': [{'path': p, 'sha256': digest(resolve(p))} for p in sorted(paths)],
        'environment': environment(), 'model_metadata_sha256': dev_frozen['model_metadata_sha256'],
        'base_model_revision': inherited['model']['revision'], 'base_snapshot': str(model_snapshot(inherited)),
        'chat_template_sha256': dev_frozen['chat_template_sha256'],
        'prompt_bundle_sha256': json_digest([{'product_id': s['product_id'], 'tasks': s['rag_tasks']} for s in samples]),
        'statistics': stats, 'existing_input_vs_rag_gate_disagreements': disagreements,
        'integration_gate_withheld': [{'product_id': s['product_id'], **f} for s in samples for f in s['integration_gate_withheld']],
        'audit_scope_normalization': 'test100 -> test in memory only; status/fields/notes unchanged',
        'test_history_disclosure': cfg['test_history_disclosure'], 'guards': cfg['guards'],
        'comparison_limit': 'frozen RAG-context plus necessary interface treatment; control latency historical',
        'historical_rag_v2_stage': 'original P3 not entered; 32 holdout remains frozen_not_executed'}
    atomic_json(out / 'protocol_manifest.json', manifest)
    print(json.dumps({'status': manifest['status'], **stats}, ensure_ascii=False), flush=True)


def verify_protocol():
    cfg, dev_cfg, inherited, _, _, _ = checked_config()
    frozen = read_json(resolve(OUTPUT) / 'protocol_manifest.json')
    if (frozen['config'] != cfg or frozen['treatment_config'] != dev_cfg or frozen['parent_config'] != inherited
            or frozen['status'] != 'frozen_before_new_test_generation'):
        raise ValueError('Frozen additional-test protocol changed')
    for item in frozen['files']:
        if digest(resolve(item['path'])) != item['sha256']:
            raise ValueError(f"Frozen additional-test file changed: {item['path']}")
    if {n: digest(model_snapshot(inherited) / n) for n in frozen['model_metadata_sha256']} != frozen['model_metadata_sha256']:
        raise ValueError('Local base metadata changed')
    return cfg, dev_cfg, inherited, frozen


def claim_run(out):
    folder = out / 'lora_rag'
    folder.mkdir()  # One attempt, even if it later fails. No resume/rerun.
    atomic_json(folder / 'attempt_started.json', {'started_at_utc': now(), 'single_attempt_consumed': True,
                'resume_allowed': False, 'scope': 'additional_fixed_previously_observed_test100'})
    return folder


def generate():
    cfg, dev_cfg, inherited, frozen = verify_protocol()
    out = resolve(OUTPUT)
    folder = claim_run(out)
    report = {'status': 'running', 'started_at_utc': now(), 'scope': cfg['scope'], 'parameter_updates': 0,
              'project_test_read': True, 'holdout_archive_read': False, 'holdout_generation': False}
    started_run = time.perf_counter()
    try:
        import torch
        from peft import PeftModel, prepare_model_for_kbit_training
        from transformers import AutoModelForCausalLM, AutoTokenizer, BitsAndBytesConfig, GenerationConfig
        if environment() != frozen['environment'] or not torch.cuda.is_available() or not torch.cuda.is_bf16_supported():
            raise RuntimeError('Frozen CUDA/BF16 environment changed; no fallback')
        snapshot = model_snapshot(inherited)
        tokenizer = AutoTokenizer.from_pretrained(str(snapshot), local_files_only=True)
        if json_digest(tokenizer.chat_template) != frozen['chat_template_sha256']:
            raise ValueError('Chat template changed')
        generation = GenerationConfig.from_dict(read_json(out / 'inference_config.json')['decoding_resolved'])
        random.seed(42)
        torch.manual_seed(42)
        torch.cuda.manual_seed_all(42)
        quant = BitsAndBytesConfig(load_in_4bit=True, bnb_4bit_quant_type='nf4',
            bnb_4bit_use_double_quant=True, bnb_4bit_compute_dtype=torch.bfloat16)
        model = AutoModelForCausalLM.from_pretrained(str(snapshot), local_files_only=True, device_map={'': 0},
            quantization_config=quant, dtype=torch.bfloat16, low_cpu_mem_usage=True)
        model, _ = prepare_kbit_with_cpu_staged_large_layers(model, prepare_model_for_kbit_training,
                                                            use_gradient_checkpointing=False)
        model.config.use_cache = True
        model = PeftModel.from_pretrained(model, str(resolve(inherited['adapter']['path'])),
                                         is_trainable=False, local_files_only=True)
        model.eval()
        if any(p.requires_grad for p in model.parameters()) or not model.is_loaded_in_4bit:
            raise RuntimeError('Expected frozen4-bit inference-only model')
        before = {n: p._version for n, p in model.named_parameters()}
        samples = load_jsonl(out / 'frozen_inputs.jsonl')
        report.update(gpu=torch.cuda.get_device_name(0), adapter_loaded=True, base_snapshot=str(snapshot),
                      model_load_seconds=time.perf_counter() - started_run)
        for index, sample in enumerate(samples, 1):
            tasks = {}
            for task in TASK_TYPES:
                expected = sample['rag_tasks'][task]
                tick = time.perf_counter()
                inputs, meta = treatment.preflight(tokenizer, sample, task, dev_cfg)
                preflight_seconds = time.perf_counter() - tick
                if meta != expected['preflight']:
                    raise ValueError('Frozen tokenizer preflight changed')
                inputs = {k: v.to('cuda') for k, v in inputs.items()}
                random.seed(expected['seed'])
                torch.manual_seed(expected['seed'])
                torch.cuda.manual_seed_all(expected['seed'])
                torch.cuda.reset_peak_memory_stats()
                torch.cuda.synchronize()
                tick = time.perf_counter()
                with torch.no_grad(), torch.autocast('cuda', dtype=torch.bfloat16):
                    generated = model.generate(**inputs, generation_config=generation)
                torch.cuda.synchronize()
                latency = time.perf_counter() - tick
                ids = generated[0, inputs['input_ids'].shape[1]:].tolist()
                text = tokenizer.decode(ids, skip_special_tokens=True).strip()
                row = {'product_id': sample['product_id'], 'variant': VARIANTS[1], 'task_type': task,
                    'seed': expected['seed'], 'prompt_sha256': expected['prompt_sha256'], 'preflight': meta,
                    'preflight_seconds': preflight_seconds, 'raw_text': text, 'parsed': parse_task(task, text),
                    'generated_tokens': len(ids), 'hit_max_new_tokens': len(ids) >= inherited['decoding']['max_new_tokens'],
                    'latency_seconds': latency, 'memory': memory()}
                if not text:
                    raise RuntimeError('Empty output; stop without retry')
                tasks[task] = row
                append_jsonl(folder / 'task_outputs.jsonl', row)
                del generated, inputs
            product = assemble(sample, tasks, VARIANTS[1])
            product['rag_selection_seconds'] = sample['rag_selection_seconds']
            product['latency_seconds']['three_task_plus_retrieval'] = product['latency_seconds']['three_task_total'] + sample['rag_selection_seconds']
            product['latency_seconds']['three_task_plus_retrieval_and_preflight'] = product['latency_seconds']['three_task_plus_retrieval'] + sum(t['preflight_seconds'] for t in tasks.values())
            append_jsonl(folder / 'product_outputs.jsonl', product)
            print(json.dumps({'phase': 'additional_fixed_test100', 'products_complete': index, 'total': 100,
                              'elapsed_seconds': round(time.perf_counter() - started_run, 2)}), flush=True)
        if before != {n: p._version for n, p in model.named_parameters()} or any(p.grad is not None for p in model.parameters()):
            raise RuntimeError('Inference changed parameters or accumulated gradients')
        verify_protocol()
        report.update(status='passed', completed_at_utc=now(), products=100, task_calls=300,
            parsed_tasks=sum(t['parsed']['success'] for t in load_jsonl(folder / 'task_outputs.jsonl')),
            model_parameters_unchanged=True, wall_seconds=time.perf_counter() - started_run)
        atomic_json(folder / 'run_report.json', report)
    except Exception as error:
        report.update(status='failed', error=str(error), failed_at_utc=now())
        atomic_json(folder / 'run_report.json', report)
        raise


def paired_data():
    cfg, _, _, frozen = verify_protocol()
    out = resolve(OUTPUT)
    report = read_json(out / 'lora_rag/run_report.json')
    if report['status'] != 'passed' or report['products'] != 100 or report['task_calls'] != 300 or not report['model_parameters_unchanged']:
        raise ValueError('Complete one-shot generation required')
    samples = {s['product_id']: s for s in load_jsonl(out / 'frozen_inputs.jsonl')}
    if len(samples) != 100:
        raise ValueError('Frozen100 coverage changed')
    versions = {}
    for variant, path in zip(VARIANTS, (out / 'control_product_outputs.jsonl', out / 'lora_rag/product_outputs.jsonl')):
        rows = load_jsonl(path)
        versions[variant] = {r['product_id']: r for r in rows}
        if len(rows) != 100 or set(versions[variant]) != set(samples):
            raise ValueError('Missing/duplicate generated products')
        for row in rows:
            if row['variant'] != variant:
                raise ValueError('Model variant changed')
            for task in TASK_TYPES:
                item = row['tasks'][task]
                expected = samples[row['product_id']]['tasks' if variant == VARIANTS[0] else 'rag_tasks'][task]
                if any(item[k] != expected[k] for k in ('seed', 'prompt_sha256', 'preflight')):
                    raise ValueError('Output seed/input/preflight changed')
                if not math.isfinite(item['latency_seconds']) or item['latency_seconds'] <= 0:
                    raise ValueError('Invalid task latency')
    for variant, path in zip(VARIANTS, (out / 'control_task_outputs.jsonl', out / 'lora_rag/task_outputs.jsonl')):
        tasks = load_jsonl(path)
        if len(tasks) != 300 or len({(t['product_id'], t['task_type']) for t in tasks}) != 300:
            raise ValueError('Independent task coverage mismatch')
        for t in tasks:
            if t != versions[variant][t['product_id']]['tasks'][t['task_type']]:
                raise ValueError('Independent task differs from assembled task')
    mapping = read_json(out / 'blind_mapping.json')['pairs']
    if mapping != mapping_for(list(samples.values()), cfg['review']['blind_seed']):
        raise ValueError('Anonymous mapping changed')
    return cfg, frozen, samples, versions, mapping


def save_manifest(out, filename, status, *, extra=None):
    excluded = {filename, filename.replace('.json', '_checksums.sha256')}
    paths = [p for p in out.rglob('*') if p.is_file() and p.name not in excluded]
    files = [{'path': p.relative_to(ROOT).as_posix(), 'sha256': digest(p)} for p in sorted(paths)]
    manifest = {'status': status, 'created_at_utc': now(), 'files': files,
                'test_history_disclosure': read_json(resolve(CONFIG))['test_history_disclosure'],
                'guards': GUARDS, **(extra or {})}
    atomic_json(out / filename, manifest)
    entries = files + [{'path': (out / filename).relative_to(ROOT).as_posix(), 'sha256': digest(out / filename)}]
    (out / filename.replace('.json', '_checksums.sha256')).write_text(
        ''.join(f"{i['sha256']}  {i['path']}\n" for i in entries), encoding='utf-8')


def write_report(out, cfg, frozen, summaries, human=None):
    lines = ['# LoRA + RAG 追加固定 test100 正式对照', '', cfg['test_history_disclosure'], '',
             '## 状态与可比性', '',
             '同原100件、同625核心字段分母。control复用已冻结LoRA三任务输出，组合仅生成一次300任务。',
             '两臂同adapter/基座、task、max tokens、解码与逐商品逐任务seed；处理差异仅为已冻结的四区上下文及必要system接口适配。',
             '历史control耗时非同期测量，不直接作纯检索因果归因；自动线索不是事实错误率或通顺率。', '',
             '| 自动辅助指标 | 冻结LoRA | LoRA + RAG |', '|---|---:|---:|']
    a, b = [summaries[v] for v in VARIANTS]
    for label, key in (('组装结构成功/100', 'assembled_structure_success'), ('空任务/300', 'empty_tasks'),
                       ('max-token任务/300', 'max_token_tasks'), ('疑似评价样本（非事实错误）', 'claim_suspect_samples_NOT_FACTUAL_RATE'),
                       ('疑似数值样本（非事实错误）', 'number_suspect_samples_NOT_FACTUAL_RATE'),
                       ('峰值allocated MiB', 'peak_allocated_mib')):
        lines.append(f'| {label} | {a[key]} | {b[key]} |')
    for task in TASK_TYPES:
        lines.append(f"| {task} 解析成功/100 | {a['task_parse_success'][task]} | {b['task_parse_success'][task]} |")
    for v in VARIANTS:
        s = summaries[v]
        lines += ['', f'### {v}', '', f"字面覆盖（非人审）：{s['literal_core_coverage_NOT_HUMAN']}",
                  f"各任务和三任务总耗时：{json.dumps(s['latency_seconds'], ensure_ascii=False)}",
                  f"输入tokens：mean={s['input_tokens_mean']:.2f}, max={s['input_tokens_max']}",
                  f"Mandatory实际token注入：{s['mandatory_core_injection_per_task']}",
                  f"效果/范围疑似样本={s['effect_scope_suspect_samples_REVIEW_ONLY']}；场景疑似样本={s['scenario_suspect_samples_REVIEW_ONLY']}；范围边界线索样本={s['range_suspect_samples_REVIEW_ONLY']}",
                  f"分品类自动辅助：{json.dumps(s['by_category'], ensure_ascii=False)}"]
    lines += ['', '## 正式人工指标', '']
    if human is None:
        lines += ['待完成100行A/B人工复核；目前不能给出组合正式核心命中率、通顺率、事实错误样本率或偏好结论。',
                  f'填写{WORKBOOK}黄色列：两臂各三指标、四项偏好、reviewer、review_confirmed=1。',
                  '旧LoRA正式标签不预填到本轮工作簿，避免破坏配对匿名评审；既有结果保持原样。']
    else:
        lines += ['| 人工指标 | 冻结LoRA | LoRA + RAG |', '|---|---:|---:|']
        for key in ('core_attribute_hit_rate', 'fluency_pass_rate', 'factual_error_sample_rate', 'average_factual_errors_per_sample'):
            lines.append(f"| {key} | {human['versions'][VARIANTS[0]][key]} | {human['versions'][VARIANTS[1]][key]} |")
        lines += ['', '匿名偏好：', json.dumps(human['preferences'], ensure_ascii=False, indent=2), '',
                  '分品类人审：', json.dumps(human['by_category'], ensure_ascii=False, indent=2), '',
                  '重点事实错误/诊断记录见human_diagnostic_cases.json；不据case调整或重跑。']
    lines += ['', '## 冻结与停止', '', f"生成前统计：{json.dumps(frozen['statistics'], ensure_ascii=False)}",
              '协议/代码/数据/Prompt/adapter/环境hash见protocol_manifest.json，原始代码字节也已独立归档。',
              '原RAG v2 P3未进入；32条holdout清单未读取、未生成，原归档不变。没有训练、调参、commit或push。',
              '本轮无论结果好坏只归档，禁止修改后重跑同批test。']
    (out / ('final_test_report.md' if human else 'automatic_test_report.md')).write_text('\n'.join(lines) + '\n', encoding='utf-8')


def build_review():
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    cfg, frozen, samples, versions, mapping = paired_data()
    out = resolve(OUTPUT)
    if (out / WORKBOOK).exists() or (out / 'automatic_summary.json').exists():
        raise FileExistsError('Never overwrite review/results')
    summaries, diagnostics = treatment.automatic_summary(samples, versions)
    for variant in VARIANTS:
        extras = [extra_hints(samples[pid], p) for pid, p in versions[variant].items()]
        for pid, product in versions[variant].items():
            diagnostics[variant][pid]['additional_REVIEW_ONLY'] = extra_hints(samples[pid], product)
        for key, source in (('effect_scope_suspect_samples_REVIEW_ONLY', 'effect_or_scope_terms_REVIEW_ONLY'),
                            ('scenario_suspect_samples_REVIEW_ONLY', 'scenario_terms_REVIEW_ONLY_may_be_supported'),
                            ('range_suspect_samples_REVIEW_ONLY', 'range_boundary_clues_REVIEW_ONLY')):
            summaries[variant][key] = sum(bool(e[source]) for e in extras)
    wb = Workbook()
    guide = wb.active
    guide.title = '复核说明'
    instructions = ['追加固定project_test100：原100件全部成对复核、200份文案，不根据输出抽样。',
        cfg['test_history_disclosure'], 'A/B仅版本匿名；实际输入不同可能揭示增强方式，不宣称完全双盲。',
        '黄色列人工填写，灰色事实/原625分母/输出不得改；不导入旧人工标签。两臂实际输入事实与输出在同一行。',
        *[t.replace('全量200件', '全量100件') for t in GUIDE[2:18]],
        '全部100行两臂正式三指标、四项偏好、reviewer和review_confirmed=1完成后才汇总、解盲。',
        '不要查看veryHidden映射或带版本名文件；不根据本轮test修改模型、Prompt或数据，不允许修改后重跑。']
    guide.append(['顺序', '复核口径与操作'])
    for i, instruction in enumerate(instructions, 1):
        guide.append([i, instruction])
        guide.row_dimensions[i + 1].height = 50
    guide.column_dimensions['B'].width = 140
    for row in guide:
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    columns = treatment.integration_columns()
    indices = {k: i for i, (k, _) in enumerate(columns, 1)}
    ws = wb.create_sheet('盲评成对复核')
    ws.append([k for k, _ in columns])
    ws.append([label for _, label in columns])
    for pair in mapping:
        values = treatment.review_fixed(samples[pair['product_id']], pair, versions)
        ws.append([values.get(k) for k, _ in columns])
    fixed_keys = set(values)
    for row in ws:
        for c in row:
            editable = c.row > 2 and columns[c.column - 1][0] not in fixed_keys
            c.alignment = Alignment(wrap_text=True, vertical='top')
            c.protection = Protection(locked=not editable)
            c.fill = PatternFill('solid', fgColor='FFF2CC' if editable else 'E7E6E6')
            if isinstance(c.value, str):
                c.data_type = 's'
            if c.row <= 2:
                c.fill = PatternFill('solid', fgColor='244062')
                c.font = Font(color='FFFFFF', bold=True)
    for key, index in indices.items():
        ws.column_dimensions[get_column_letter(index)].width = 62 if key.endswith((*TASK_TYPES, '_complete')) else 42 if any(s in key for s in ('attributes', 'notes', 'input_facts')) else 18
        ws.cell(1, index).comment = Comment(columns[index - 1][1], 'fixed additional test100')
        if key.endswith('_task_messages') or key in ('source_title', 'input_quality_actions'):
            ws.column_dimensions[get_column_letter(index)].hidden = True
    for r in range(3, 103):
        ws.row_dimensions[r].height = 220
        ws.cell(r, indices['product_id']).number_format = '@'
    for p in PREFERENCES:
        dv = DataValidation(type='list', formula1='"A更好,B更好,持平"', allow_blank=True)
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(f'{get_column_letter(indices[p])}3:{get_column_letter(indices[p])}102')
    for side in ('A', 'B'):
        for m in (*METRICS, *DIAGNOSTICS):
            dv = DataValidation(type='whole', operator='greaterThanOrEqual' if m == 'factual_error_count' else 'between',
                formula1='0', formula2=None if m == 'factual_error_count' else '$I3' if m == 'matched_attribute_count' else '1', allow_blank=True)
            dv.showErrorMessage = True
            ws.add_data_validation(dv)
            col = get_column_letter(indices[f'{side}_{m}'])
            dv.add(f'{col}3:{col}102')
    dv = DataValidation(type='list', formula1='"1"', allow_blank=True)
    ws.add_data_validation(dv)
    col = get_column_letter(indices['review_confirmed'])
    dv.add(f'{col}3:{col}102')
    ws.freeze_panes = 'F3'
    ws.auto_filter.ref = f'A2:{get_column_letter(len(columns))}102'
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    sealed = wb.create_sheet('匿名映射_汇总后解盲')
    sealed.append(['pair_id', 'product_id', 'A', 'B'])
    for pair in mapping:
        sealed.append([pair[k] for k in ('pair_id', 'product_id', 'A', 'B')])
    sealed.sheet_state = 'veryHidden'
    hints = wb.create_sheet('自动线索_非人工指标')
    hints.append(['pair_id', 'side', 'product_id', 'diagnostics_REVIEW_ONLY'])
    for pair in mapping:
        for side in ('A', 'B'):
            hints.append([pair['pair_id'], side, pair['product_id'], json.dumps(diagnostics[pair[side]][pair['product_id']], ensure_ascii=False)])
    hints.sheet_state = 'hidden'
    wb.save(out / WORKBOOK)
    (out / WORKBOOK.replace('.xlsx', '.initial_blank.xlsx')).write_bytes((out / WORKBOOK).read_bytes())
    atomic_json(out / 'automatic_summary.json', {'status': 'human_review_pending_NOT_FINAL', 'versions': summaries,
        'statistics': frozen['statistics'], 'human_metrics': None, 'human_preferences': None,
        'test_history_disclosure': cfg['test_history_disclosure'], 'historical_control_latency': True})
    atomic_json(out / 'automatic_diagnostics.json', diagnostics)
    write_report(out, cfg, frozen, summaries)
    files = [{'path': p.relative_to(ROOT).as_posix(), 'sha256': digest(p)}
             for p in sorted(out.rglob('*')) if p.is_file() and p.name != WORKBOOK]
    atomic_json(out / 'review_manifest.json', {'status': 'human_review_pending', 'products': 100, 'documents': 200,
        'files': files, 'workbook': (out / WORKBOOK).relative_to(ROOT).as_posix(),
        'initial_workbook_sha256': digest(out / WORKBOOK), 'labels_prefilled': False})
    verify_protocol()
    save_manifest(out, 'automatic_phase_manifest.json', 'generation_complete_human_review_pending')
    print(json.dumps({'workbook': str(out / WORKBOOK), 'versions': summaries}, ensure_ascii=False), flush=True)


def validate_review_rows(ws, samples, versions, mapping):
    columns = treatment.integration_columns()
    if ws.max_row != 102 or ws.max_column != len(columns) or [c.value for c in ws[1]] != [k for k, _ in columns]:
        raise ValueError('Review coverage/columns changed')
    rows = []
    for index, pair in enumerate(mapping, 3):
        row = {k: ws.cell(index, j).value for j, (k, _) in enumerate(columns, 1)}
        for key, expected in treatment.review_fixed(samples[pair['product_id']], pair, versions).items():
            if row[key] != expected:
                raise ValueError(f"Frozen review data changed: {pair['pair_id']} {key}")
        if str(row['review_confirmed']).strip() != '1' or not str(row['reviewer'] or '').strip():
            raise ValueError('Incomplete human confirmation')
        if any(row[p] not in CHOICES for p in PREFERENCES):
            raise ValueError('Incomplete/invalid anonymous preferences')
        for side in ('A', 'B'):
            official_metrics([{'product_id': pair['product_id'], 'core_attribute_count': row['core_attribute_count'],
                              **{m: row[f'{side}_{m}'] for m in METRICS}}])
            for d in DIAGNOSTICS:
                value = row[f'{side}_{d}']
                if value is not None and str(value) not in ('0', '1'):
                    raise ValueError('Invalid diagnosis')
                if str(value) == '1' and not row[f'{side}_notes']:
                    raise ValueError('Diagnosis needs evidence')
            if int(row[f'{side}_factual_error_count']) > 0 and not row[f'{side}_notes']:
                raise ValueError('Factual errors need evidence')
        rows.append(row)
    return rows


def summarize():
    from openpyxl import load_workbook
    cfg, frozen, samples, versions, mapping = paired_data()
    out = resolve(OUTPUT)
    for name in ('blind_summary.json', 'human_metrics_unblinded.json', 'human_review_final.csv', 'final_test_manifest.json'):
        if (out / name).exists():
            raise FileExistsError('Never overwrite finalized human metrics')
    for item in read_json(out / 'review_manifest.json')['files']:
        if digest(resolve(item['path'])) != item['sha256']:
            raise ValueError('Frozen review outputs changed')
    path = out / WORKBOOK
    wb = load_workbook(path, data_only=False)
    if list(wb['匿名映射_汇总后解盲'].values)[1:] != [tuple(p[k] for k in ('pair_id', 'product_id', 'A', 'B')) for p in mapping]:
        raise ValueError('Workbook mapping changed')
    rows = validate_review_rows(wb['盲评成对复核'], samples, versions, mapping)
    atomic_json(out / 'blind_summary.json', {'status': 'complete_anonymous_aggregate_before_unblinding',
        'preferences': {p: dict(Counter(r[p] for r in rows)) for p in PREFERENCES}, 'workbook_sha256': digest(path)})
    judgments = {v: [] for v in VARIANTS}
    prefs = {p: Counter() for p in PREFERENCES}
    cases = []
    for row, pair in zip(rows, mapping):
        for side in ('A', 'B'):
            item = {'pair_id': pair['pair_id'], 'product_id': pair['product_id'], 'variant': pair[side],
                'category_l2': samples[pair['product_id']]['category_l2'], 'core_attribute_count': row['core_attribute_count'],
                **{m: int(row[f'{side}_{m}']) for m in METRICS}, **{d: row[f'{side}_{d}'] for d in DIAGNOSTICS},
                'reviewer': row['reviewer'], 'review_notes': row[f'{side}_notes'] or ''}
            judgments[pair[side]].append(item)
            if item['factual_error_count'] or any(str(item[d]) == '1' for d in DIAGNOSTICS):
                cases.append({**item, 'output': versions[pair[side]][pair['product_id']]['assembled'],
                              'source_attributes': samples[pair['product_id']]['source_attributes']})
        for p in PREFERENCES:
            prefs[p]['tie' if row[p] == '持平' else pair[row[p][0]]] += 1
    human = {'status': 'completed_human_review', 'versions': {v: official_metrics(r) for v, r in judgments.items()},
        'by_category': {v: {c: official_metrics([r for r in judgments[v] if r['category_l2'] == c])
                        for c in sorted({s['category_l2'] for s in samples.values()})} for v in VARIANTS},
        'preferences': {p: {k: prefs[p][k] for k in (*VARIANTS, 'tie')} for p in PREFERENCES},
        'workbook_sha256': digest(path), 'blind_summary_sha256': digest(out / 'blind_summary.json'),
        'test_history_disclosure': cfg['test_history_disclosure']}
    atomic_json(out / 'human_metrics_unblinded.json', human)
    atomic_json(out / 'human_diagnostic_cases.json', {'cases': cases, 'use': 'archive only; never retune or rerun'})
    exported = [r for rs in judgments.values() for r in rs]
    with (out / 'human_review_final.csv').open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(exported[0]))
        writer.writeheader()
        writer.writerows(exported)
    write_report(out, cfg, frozen, read_json(out / 'automatic_summary.json')['versions'], human)
    save_manifest(out, 'final_test_manifest.json', 'final_fixed_test100_human_review_complete',
                  extra={'human_metrics_sha256': digest(out / 'human_metrics_unblinded.json')})
    print('Final fixed test100 archived. Stop; no tuning, rerun, training or holdout.')
