"""Blinded review workbook and human-only summaries for frozen P7 outputs."""
from __future__ import annotations

import csv
import json
import math
import re
import unicodedata
from collections import Counter

from .evaluation import parse_binary, parse_nonnegative_integer
from .lora_validation import TASK_TYPES, load_jsonl, verify_protocol
from .lora_training import atomic_json, digest, read_json, resolve

PREFERENCES = ("title_quality", "selling_points_structure", "short_description_naturalness", "overall_ecommerce_professionalism")
DIAGNOSTICS = ("field_stacking", "mechanical_template", "unsupported_evaluation_effect_scenario", "numeric_range_distortion", "identity_category_error")
CHOICES = ("A更好", "B更好", "持平")
METRICS = ("matched_attribute_count", "fluency_pass", "factual_error_count")
CHINESE = {
    "title_quality": "标题表达质量", "selling_points_structure": "卖点结构", "short_description_naturalness": "短详情自然度",
    "overall_ecommerce_professionalism": "整体电商专业度", "field_stacking": "字段堆叠", "mechanical_template": "机械模板化",
    "unsupported_evaluation_effect_scenario": "无依据评价/效果/场景", "numeric_range_distortion": "数值或范围失真", "identity_category_error": "身份/品类错误",
    "matched_attribute_count": "整份文案正确命中的核心属性数（0至原核心总数）", "fluency_pass": "整份文案通顺：1通过/0不通过",
    "factual_error_count": "整份文案独立事实错误点数；同一错误重复只计一次",
}
GUIDE = [
    "本工作簿为原 project_validation 全量200件、400份文案；每行一对A/B。没有抽样，没有预填人工判断。",
    "先看A/B标题、卖点、短详情、完整文案；输入事实、原核心清单和源问题在同一行，可展开隐藏的源信息列。",
    "黄色单元格为人工填写。正式三指标分别填写A与B；不要改核心属性总数、源属性、生成文本或匿名映射。",
    "四个成对比较均填 A更好 / B更好 / 持平；差异不明确时持平，不为了区分而强选一方。",
    "标题表达质量：识别信息清晰、简洁自然、无赘语；不是越短越好，也不是营销形容词越多越好。",
    "卖点结构：分条清楚、信息组织合理、不无谓重复；字段：值本身可以是合格卖点，不因中性而自动扣分。",
    "短详情自然度：连接和语序自然、读起来像段落；不鼓励为润色添加效果、体验或场景。",
    "整体电商专业度：准确、规范、品类表达合适，非只比较字数或营销强度。",
    "matched_attribute_count：以原核心清单为分母，全文任一位置正确表达即可；同一字段最多一次，正常同义/格式表达可命中。",
    "多值核心字段沿用既有规则：有任一可靠值被正确表达可计该字段一次；错误扩大兼容范围/数值边界不计正确命中。",
    "fluency_pass：整体语序、搭配、衔接可读，允许标题短语及三条中性卖点；事实正确不意味着通顺，事实错误不自动判不通顺。",
    "factual_error_count：无来源事实、效果/评价/场景，反转或扭曲源事实、数值范围失真、身份品类错误；同一独立错误跨区重复只计一次。",
    "判断事实时对照同商品完整源属性和可靠性备注，不能把未进入输入但源数据确有支持的事实机械判错。原标题只供风格/源质量审查，不单独作为营销断言的事实依据。",
    "直接低风险品类语义或正常同义转述不算事实错误；不能由拖把品类推出具体场景、人群或宣称某种效果。",
    "REVIEW/CONFLICT的源争议单独写进notes；Source Fact Quality不机械等同于模型幻觉。原核心分母、全量200件统计不因此改变。",
    "五项诊断0/1可选填写，1表示人工确认存在；只是解释项，不替代正式三指标。出现错误或诊断1，请写具体原文与依据。",
    "字段堆叠指影响可读性的无组织串接，不是出现字段标签就算；机械模板化指僵硬重复，固定结构本身不自动算。",
    "先完成盲评，再按需查看隐藏的自动线索（仍仅A/B）；线索/字面命中/疑似词不是正式人工结论。",
    "完成每行后填写reviewer及review_confirmed=1；所有200行正式字段和四个成对比较完成后才允许正式汇总。",
    "匿名映射为veryHidden页并另存blind_mapping.json；隐藏不是安全加密。评审完成前不要查看映射、带版本名输出或速度日志。",
    "评审程序先保存不解盲的A/B汇总，再应用封存映射形成Base/LoRA汇总。不完整评审不会产生正式metrics。",
    "若分批复核，直接保存此xlsx继续填写，不改抽样范围；不根据输出挑子集。不要运行project_test或RAG holdout。",
]


def normalized(value):
    return "".join(c for c in unicodedata.normalize("NFKC", str(value)).casefold() if c.isalnum())


def automatic_hints(sample, product):
    text = product["assembled"]["complete_raw_text"]
    norm = normalized(text)
    matched = [f for f, values in sample["evaluation_attributes"].items() if any(normalized(v) and normalized(v) in norm for v in values)]
    terms = [t for t in ("优质", "高效", "耐用", "方便", "安全", "舒适", "精准", "稳定", "提升", "确保", "带来", "体验", "各种", "广泛") if t in text]
    source = normalized(json.dumps(sample["source_attributes"], ensure_ascii=False))
    numbers = [m.group() for m in re.finditer(r"\d+(?:\.\d+)?\s*(?:mAh|mL|ml|mm|cm|kg|dpi|W|V|小时|毫升|毫安|升)", text, re.IGNORECASE) if normalized(m.group()) not in source]
    return {"literal_matched_count_NOT_HUMAN": len(matched), "literal_matched_fields": matched,
            "literal_unmatched_fields_NOT_OMISSION_VERDICT": [f for f in sample["evaluation_attributes"] if f not in matched],
            "claim_terms_REVIEW_ONLY": terms, "number_terms_REVIEW_ONLY": numbers,
            "structure_success": product["assembled"]["structure_success"],
            "tasks_hit_token_limit": [t for t in TASK_TYPES if product["tasks"][t]["hit_max_new_tokens"]]}


def paired_data():
    cfg, frozen = verify_protocol()
    out = resolve(cfg["output"])
    samples = {s["product_id"]: s for s in load_jsonl(out / "frozen_inputs.jsonl")}
    versions = {}
    for variant in ("Base", "LoRA"):
        report = read_json(out / variant.lower() / "run_report.json")
        if report["status"] != "passed" or report["products"] != 200 or report["task_calls"] != 600:
            raise ValueError("Both variants must complete before creating final review")
        rows = load_jsonl(out / variant.lower() / "product_outputs.jsonl")
        by_id = {r["product_id"]: r for r in rows}
        if len(rows) != 200 or set(by_id) != set(samples):
            raise ValueError("Missing/duplicate generation products")
        versions[variant] = by_id
    for pid in samples:
        for task in TASK_TYPES:
            a, b = (versions[v][pid]["tasks"][task] for v in ("Base", "LoRA"))
            if any(a[k] != b[k] for k in ("seed", "prompt_sha256", "preflight")):
                raise ValueError("Base/LoRA pairing mismatch")
    mapping = read_json(out / "blind_mapping.json")["pairs"]
    if len(mapping) != 200 or {m["product_id"] for m in mapping} != set(samples):
        raise ValueError("Blind mapping coverage mismatch")
    return cfg, samples, versions, mapping


def review_columns():
    columns = [("pair_id", "匿名配对编号"), ("product_id", "商品编号（文本，不改）"),
               ("category_l1", "一级品类"), ("category_l2", "二级品类"), ("source_quality_status", "源质量状态（不改）"),
               ("source_title", "源标题，仅风格/源审查"), ("source_attributes", "完整源属性"),
               ("evaluation_attributes", "原核心属性清单"), ("core_attribute_count", "原核心总数/正式分母（不改）"),
               ("used_attributes", "实际相同输入的可靠核心属性"), ("source_quality_note", "既有源质量审计备注"),
               ("input_quality_actions", "输入门控动作" )]
    for task, preference in zip(TASK_TYPES, PREFERENCES[:3]):
        columns.extend([(f"A_{task}", f"A：{task} 原始输出"), (f"B_{task}", f"B：{task} 原始输出"), (preference, CHINESE[preference] + "：A更好/B更好/持平")])
    columns.extend([("A_complete", "A完整文案（原始三任务串联）"), ("B_complete", "B完整文案（原始三任务串联）"),
                    (PREFERENCES[3], CHINESE[PREFERENCES[3]] + "：A更好/B更好/持平")])
    for side in ("A", "B"):
        columns += [(f"{side}_{m}", side + "：" + CHINESE[m]) for m in METRICS]
        columns += [(f"{side}_{d}", side + "：" + CHINESE[d] + "（可选0/1）") for d in DIAGNOSTICS]
        columns.append((f"{side}_notes", side + "：具体错误/诊断原文与依据；源争议单列"))
    columns += [("reviewer", "人工复核人"), ("review_confirmed", "本行复核确认：填1"), ("pair_notes", "成对比较理由/其他备注")]
    return columns


def fixed_values(sample, pair, versions):
    values = {"pair_id": pair["pair_id"], "product_id": pair["product_id"],
              "category_l1": sample["category_l1"], "category_l2": sample["category_l2"],
              "source_title": sample["source_title"], "source_attributes": sample["source_attributes"],
              "evaluation_attributes": sample["evaluation_attributes"], "core_attribute_count": sample["core_attribute_count"],
              "used_attributes": sample["used_attributes"], "source_quality_status": sample["source_quality"]["status"],
              "source_quality_note": sample["source_quality"]["note"], "input_quality_actions": sample["input_quality_actions"]}
    for side in ("A", "B"):
        product = versions[pair[side]][sample["product_id"]]
        for task in TASK_TYPES:
            values[f"{side}_{task}"] = product["tasks"][task]["raw_text"]
        values[f"{side}_complete"] = product["assembled"]["complete_raw_text"]
    return {k: json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v for k, v in values.items()}


def build_workbook():
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    cfg, samples, versions, mapping = paired_data()
    out = resolve(cfg["output"])
    path = out / "project_validation200_blind_review.xlsx"
    if path.exists():
        raise FileExistsError("Refusing to overwrite human review workbook")
    wb = Workbook()
    guide = wb.active
    guide.title = "复核说明"
    guide.append(["顺序", "复核口径与操作"])
    for i, instruction in enumerate(GUIDE, 1):
        guide.append([i, instruction])
        guide.row_dimensions[i + 1].height = 46
    guide.column_dimensions["A"].width = 8
    guide.column_dimensions["B"].width = 135
    for row in guide:
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws = wb.create_sheet("盲评成对复核")
    columns = review_columns()
    ws.append([k for k, _ in columns])
    ws.append([label for _, label in columns])
    indices = {k: i for i, (k, _) in enumerate(columns, 1)}
    for pair in mapping:
        values = fixed_values(samples[pair["product_id"]], pair, versions)
        ws.append([values.get(k) for k, _ in columns])
    fixed_keys = set(fixed_values(samples[mapping[0]["product_id"]], mapping[0], versions))
    for row in ws:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Preserve IDs and neutralize source-text formula injection.
            if cell.row <= 2:
                cell.fill = PatternFill("solid", fgColor="244062")
                cell.font = Font(color="FFFFFF", bold=True)
            else:
                key = columns[cell.column - 1][0]
                editable = key not in fixed_keys
                cell.protection = Protection(locked=not editable)
                if editable:
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
    for key, index in indices.items():
        letter = get_column_letter(index)
        is_output = key.endswith(tuple(TASK_TYPES)) or key.endswith("_complete")
        ws.column_dimensions[letter].width = 65 if is_output else (42 if "attributes" in key or "notes" in key else 19)
        if key in ("source_title", "source_attributes", "evaluation_attributes", "used_attributes", "source_quality_note", "input_quality_actions", "A_complete", "B_complete"):
            ws.column_dimensions[letter].hidden = True
        ws.cell(1, index).comment = Comment(columns[index - 1][1], "P7 protocol")
    for row in range(3, 203):
        ws.row_dimensions[row].height = 210
        ws.cell(row, indices["product_id"]).number_format = "@"
    for preference in PREFERENCES:
        dv = DataValidation(type="list", formula1='"A更好,B更好,持平"', allow_blank=True)
        dv.error, dv.errorTitle, dv.showErrorMessage = "仅选择A更好/B更好/持平", "无效比较值", True
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(indices[preference])}3:{get_column_letter(indices[preference])}202")
    for side in ("A", "B"):
        for field in (*METRICS, *DIAGNOSTICS):
            key = f"{side}_{field}"
            max_value = "$I3" if field == "matched_attribute_count" else "1"
            dv = DataValidation(type="whole", operator="greaterThanOrEqual" if field == "factual_error_count" else "between",
                                formula1="0", formula2=None if field == "factual_error_count" else max_value, allow_blank=True)
            dv.showErrorMessage = True
            dv.error = "请填写非负整数；命中数不得超过核心总数，通顺和诊断仅填0/1。"
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(indices[key])}3:{get_column_letter(indices[key])}202")
    confirm = DataValidation(type="list", formula1='"1"', allow_blank=True)
    ws.add_data_validation(confirm)
    confirm.add(f"{get_column_letter(indices['review_confirmed'])}3:{get_column_letter(indices['review_confirmed'])}202")
    ws.freeze_panes, ws.auto_filter.ref = "F3", ws.dimensions
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    automatic = wb.create_sheet("自动线索_非人工指标")
    automatic.append(["pair_id", "product_id", "anonymous_side", "automatic_hints_NOT_OFFICIAL"])
    for pair in mapping:
        for side in ("A", "B"):
            sample = samples[pair["product_id"]]
            automatic.append([pair["pair_id"], pair["product_id"], side,
                json.dumps(automatic_hints(sample, versions[pair[side]][pair["product_id"]]), ensure_ascii=False)])
    automatic.sheet_state = "hidden"
    mapping_ws = wb.create_sheet("匿名映射_汇总后解盲")
    mapping_ws.append(["pair_id", "product_id", "A", "B"])
    for pair in mapping:
        mapping_ws.append([pair[k] for k in ("pair_id", "product_id", "A", "B")])
    mapping_ws.sheet_state = "veryHidden"
    wb.save(path)
    outputs = [p for folder in ("base", "lora") for p in (out / folder).iterdir() if p.is_file()]
    atomic_json(out / "review_manifest.json", {"status": "awaiting_human_blinded_review", "workbook": str(path),
        "initial_workbook_sha256": digest(path), "products": 200, "model_documents": 400, "labels_prefilled": False,
        "files": [{"path": str(p.relative_to(resolve("."))).replace("\\", "/"), "sha256": digest(p)} for p in outputs],
        "protocol_manifest_sha256": digest(out / "protocol_manifest.json"), "mapping_sha256": digest(out / "blind_mapping.json")})
    print(f"Blinded review workbook created: {path}")


def official_metrics(rows):
    if not rows:
        raise ValueError("Empty human review")
    clean = []
    seen = set()
    for row in rows:
        pid = row["product_id"]
        if not pid or pid in seen:
            raise ValueError("Empty/duplicate reviewed product")
        seen.add(pid)
        count = parse_nonnegative_integer(str(row["core_attribute_count"]), "core_attribute_count", pid)
        matched = parse_nonnegative_integer(str(row["matched_attribute_count"]), "matched_attribute_count", pid)
        if count <= 0 or matched > count:
            raise ValueError("Invalid unchanged denominator/match count")
        fluency = parse_binary(str(row["fluency_pass"]), "fluency_pass", pid)
        errors = parse_nonnegative_integer(str(row["factual_error_count"]), "factual_error_count", pid)
        clean.append((count, matched, fluency, errors))
    total = sum(r[0] for r in clean)
    return {"sample_count": len(clean), "core_attribute_total": total, "matched_attribute_total": sum(r[1] for r in clean),
            "core_attribute_hit_rate": sum(r[1] for r in clean) / total, "fluency_pass_rate": sum(r[2] for r in clean) / len(clean),
            "factual_error_sample_rate": sum(r[3] > 0 for r in clean) / len(clean),
            "factual_error_total": sum(r[3] for r in clean), "average_factual_errors_per_sample": sum(r[3] for r in clean) / len(clean)}


def summarize_workbook():
    """No partial metrics. First aggregate anonymous votes; unblind only afterwards."""
    from openpyxl import load_workbook
    cfg, samples, versions, mapping = paired_data()
    out = resolve(cfg["output"])
    manifest = read_json(out / "review_manifest.json")
    for item in manifest["files"]:
        if digest(resolve(item["path"])) != item["sha256"]:
            raise ValueError("Frozen generation output changed")
    if digest(out / "protocol_manifest.json") != manifest["protocol_manifest_sha256"] or digest(out / "blind_mapping.json") != manifest["mapping_sha256"]:
        raise ValueError("Protocol/mapping changed")
    for name in ("blind_summary.json", "human_metrics_unblinded.json", "human_review_final.csv"):
        if (out / name).exists():
            raise FileExistsError("Refusing to overwrite finalized human review")
    path = out / "project_validation200_blind_review.xlsx"
    wb = load_workbook(path, data_only=False)
    mapping_ws = wb["匿名映射_汇总后解盲"]
    if mapping_ws.max_row != 201 or list(mapping_ws.values)[1:] != [tuple(p[k] for k in ("pair_id", "product_id", "A", "B")) for p in mapping]:
        raise ValueError("Workbook anonymous mapping changed")
    ws = wb["盲评成对复核"]
    columns = review_columns()
    if ws.max_row != 202 or [c.value for c in ws[1]] != [k for k, _ in columns]:
        raise ValueError("Workbook coverage/columns changed")
    rows = []
    numeric = lambda v: str(int(v)) if isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v) and int(v) == v else str(v).strip()
    for index, pair in enumerate(mapping, 3):
        row = {k: ws.cell(index, j).value for j, (k, _) in enumerate(columns, 1)}
        for key, value in fixed_values(samples[pair["product_id"]], pair, versions).items():
            if row[key] != value:
                raise ValueError(f"Frozen review source/output changed: {pair['pair_id']} {key}")
        if numeric(row["review_confirmed"]) != "1" or not row["reviewer"]:
            raise ValueError(f"Incomplete human confirmation: {pair['pair_id']}")
        for preference in PREFERENCES:
            if row[preference] not in CHOICES:
                raise ValueError(f"Incomplete/invalid paired preference: {pair['pair_id']} {preference}")
        for side in ("A", "B"):
            for metric in METRICS:
                row[f"{side}_{metric}"] = numeric(row[f"{side}_{metric}"])
            official_metrics([{ "product_id": row["product_id"], "core_attribute_count": row["core_attribute_count"],
                               **{m: row[f"{side}_{m}"] for m in METRICS}}])
            for diagnosis in DIAGNOSTICS:
                value = row[f"{side}_{diagnosis}"]
                if value is not None:
                    row[f"{side}_{diagnosis}"] = parse_binary(numeric(value), diagnosis, row["product_id"])
                    if row[f"{side}_{diagnosis}"] == 1 and not row[f"{side}_notes"]:
                        raise ValueError("Confirmed diagnostics require evidence notes")
            if int(row[f"{side}_factual_error_count"]) > 0 and not row[f"{side}_notes"]:
                raise ValueError("Fact errors require evidence notes")
        rows.append(row)
    anonymous = {p: dict(Counter(row[p] for row in rows)) for p in PREFERENCES}
    atomic_json(out / "blind_summary.json", {"status": "complete_blinded_aggregate_before_unblinding", "products": 200,
                                            "preferences": anonymous, "completed_workbook_sha256": digest(path)})
    # Mapping is only applied to scores after the complete anonymous aggregation above.
    judgments, preferences = {"Base": [], "LoRA": []}, {p: Counter() for p in PREFERENCES}
    exported = []
    for row, pair in zip(rows, mapping):
        for side in ("A", "B"):
            item = {"product_id": row["product_id"], "variant": pair[side], "pair_id": pair["pair_id"],
                    "core_attribute_count": row["core_attribute_count"], **{m: row[f"{side}_{m}"] for m in METRICS},
                    **{d: row[f"{side}_{d}"] for d in DIAGNOSTICS},
                    **{p: row[p] for p in PREFERENCES}, "anonymous_side": side,
                    "reviewer": row["reviewer"], "review_notes": row[f"{side}_notes"] or "", "pair_notes": row["pair_notes"] or ""}
            judgments[pair[side]].append(item)
            exported.append(item)
        for preference in PREFERENCES:
            winner = "tie" if row[preference] == "持平" else pair[row[preference][0]]
            preferences[preference][winner] += 1
    atomic_json(out / "human_metrics_unblinded.json", {"status": "final_human_review_unblinded", "products_per_model": 200,
                "official_metrics": {v: official_metrics(judgments[v]) for v in judgments},
                "paired_preferences": {p: {k: preferences[p][k] for k in ("Base", "LoRA", "tie")} for p in preferences},
                "diagnostics_NOT_OFFICIAL_METRICS": {v: {d: {"annotated_samples": sum(r[d] is not None for r in judgments[v]),
                    "positive_samples": sum(r[d] == 1 for r in judgments[v])} for d in DIAGNOSTICS} for v in judgments},
                "blind_aggregate_sha256": digest(out / "blind_summary.json"), "workbook_sha256": digest(path)})
    with (out / "human_review_final.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(exported[0]))
        writer.writeheader()
        writer.writerows(exported)
    print("Complete human metrics saved; no project_test generation performed.")
