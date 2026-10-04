"""Test100 automatic diagnostics and unchanged P7 human metric definitions."""
from __future__ import annotations

import csv
import json
import math
import re
import statistics
from collections import Counter

from .evaluation import parse_binary
from .lora_formal_test import CONFIG, OUTPUT, verify_protocol
from .lora_training import atomic_json, digest, now, read_json, resolve
from .lora_validation import TASK_TYPES, append_jsonl, load_jsonl
from .lora_validation_review import (CHOICES, DIAGNOSTICS, GUIDE, METRICS, PREFERENCES,
                                    automatic_hints, fixed_values, official_metrics, review_columns)

WORKBOOK = "project_test100_blind_review.xlsx"
VARIANTS = ("Base", "LoRA")


def paired_data():
    cfg, inherited, frozen = verify_protocol()
    out = resolve(OUTPUT)
    samples = {s["product_id"]: s for s in load_jsonl(out / "frozen_inputs.jsonl")}
    mapping = read_json(out / "blind_mapping.json")["pairs"]
    versions = {}
    if len(samples) != 100 or len(mapping) != 100 or {m["product_id"] for m in mapping} != set(samples):
        raise ValueError("Frozen test100 coverage mismatch")
    for variant in VARIANTS:
        folder = out / variant.lower()
        report = read_json(folder / "run_report.json")
        if report["status"] != "passed" or report["products"] != 100 or report["task_calls"] != 300 or not report["model_parameters_unchanged"]:
            raise ValueError("Each one-shot run must finish before review")
        rows = load_jsonl(folder / "product_outputs.jsonl")
        tasks = load_jsonl(folder / "task_outputs.jsonl")
        by_id = {r["product_id"]: r for r in rows}
        if len(rows) != 100 or set(by_id) != set(samples) or len(tasks) != 300 or len({(t["product_id"], t["task_type"]) for t in tasks}) != 300:
            raise ValueError("Missing/duplicate independent outputs")
        for t in tasks:
            expected = samples[t["product_id"]]["tasks"][t["task_type"]]
            if t != by_id[t["product_id"]]["tasks"][t["task_type"]] or t["variant"] != variant:
                raise ValueError("Independent/assembled outputs differ")
            if any(t[k] != expected[k] for k in ("seed", "prompt_sha256", "preflight")):
                raise ValueError("Task pairing changed")
            if not math.isfinite(t["latency_seconds"]) or t["latency_seconds"] <= 0:
                raise ValueError("Invalid task latency")
        versions[variant] = by_id
    a, b = [read_json(out / f"inference_{v.lower()}.json") for v in VARIANTS]
    if {k for k in a if a[k] != b[k]} != {"variant", "adapter"}:
        raise ValueError("Only adapter presence may differ")
    return cfg, inherited, frozen, samples, versions, mapping


def extra_hints(sample, product):
    """Additional REVIEW ONLY diagnostics, not validator or human judgments."""
    text = product["assembled"]["complete_raw_text"]
    effect_terms = [t for t in ("确保", "保证", "提升", "满足", "带来", "提供", "省力", "省电", "持久", "无延迟", "适合", "适用于", "专为", "各种", "广泛") if t in text]
    scenarios = [t for t in ("办公", "游戏", "旅行", "户外", "家庭", "家居", "日常", "学生", "儿童", "老人", "宿舍", "厨房", "卧室") if t in text]
    # Literal boundary absence is merely a review clue: correct prose may pass humans.
    boundaries = []
    for field, values in sample["evaluation_attributes"].items():
        for value in values:
            if re.search(r"[\[(（]\s*\d+(?:\.\d+)?\s*[,，]\s*\d+(?:\.\d+)?\s*[\])）]", str(value)) and str(value) not in text:
                boundaries.append({"field": field, "source_range": value, "literal_range_absent_NOT_ERROR": True})
    return {"effect_or_scope_terms_REVIEW_ONLY": effect_terms,
            "scenario_terms_REVIEW_ONLY_may_be_supported": scenarios,
            "range_boundary_clues_REVIEW_ONLY": boundaries}


def percentile(values, q):
    ordered = sorted(values)
    index = (len(ordered) - 1) * q
    low, high = math.floor(index), math.ceil(index)
    return ordered[low] + (ordered[high] - ordered[low]) * (index - low)


def automatic_group(samples, products):
    rows = list(products.values())
    hints = [automatic_hints(samples[p["product_id"]], p) for p in rows]
    extras = [extra_hints(samples[p["product_id"]], p) for p in rows]
    tasks = [p["tasks"][t] for p in rows for t in TASK_TYPES]
    denominator = sum(samples[p["product_id"]]["core_attribute_count"] for p in rows)
    matched = sum(h["literal_matched_count_NOT_HUMAN"] for h in hints)
    return {"products": len(rows), "task_calls": len(tasks),
        "parse_success_by_task_NOT_FLUENCY": {t: sum(p["tasks"][t]["parsed"]["success"] for p in rows) for t in TASK_TYPES},
        "assembled_structure_success_NOT_FLUENCY": sum(p["assembled"]["structure_success"] for p in rows),
        "literal_coverage_NOT_HUMAN": {"matched_fields": matched, "core_attribute_total": denominator, "ratio": matched / denominator},
        "latency": {t: {"samples": len(rows), "mean_seconds": statistics.mean(p["latency_seconds"][t] for p in rows),
            "p50_seconds": percentile([p["latency_seconds"][t] for p in rows], .5),
            "p95_seconds": percentile([p["latency_seconds"][t] for p in rows], .95)} for t in (*TASK_TYPES, "three_task_total")},
        "empty_task_outputs": sum(not t["raw_text"] for t in tasks),
        "hit_token_limit_tasks": sum(t["hit_max_new_tokens"] for t in tasks),
        "claim_terms_REVIEW_ONLY_samples": sum(bool(h["claim_terms_REVIEW_ONLY"]) for h in hints),
        "number_terms_REVIEW_ONLY_samples": sum(bool(h["number_terms_REVIEW_ONLY"]) for h in hints),
        "effect_scope_terms_REVIEW_ONLY_samples": sum(bool(h["effect_or_scope_terms_REVIEW_ONLY"]) for h in extras),
        "scenario_terms_REVIEW_ONLY_samples": sum(bool(h["scenario_terms_REVIEW_ONLY_may_be_supported"]) for h in extras),
        "range_boundary_clues_REVIEW_ONLY_samples": sum(bool(h["range_boundary_clues_REVIEW_ONLY"]) for h in extras),
        "peak_allocated_mib": max(t["memory"]["peak_allocated_mib"] for t in tasks),
        "peak_reserved_mib": max(t["memory"]["peak_reserved_mib"] for t in tasks)}


def build_workbook(out, samples, versions, mapping):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    path = out / WORKBOOK
    if path.exists():
        raise FileExistsError("Never overwrite human workbook")
    wb = Workbook()
    guide = wb.active
    guide.title = "复核说明"
    instructions = ["最终project_test范围：原冻结test100全部100件，100对/200份文案。不重新抽样，映射在生成前冻结。",
        "本轮是LoRA v1一次性正式测试，旧test100此前用于Baseline/RAG，并非从未观察的测试集。",
        "完成前不要查看版本输出、带版本名自动报告或veryHidden映射。黄色列人工填写，灰色事实/输出/分母不得修改。"]
    instructions += [t.replace("200", "100").replace("project_validation", "project_test100") for t in GUIDE if "本工作簿为" not in t]
    guide.append(["顺序", "操作与口径"])
    for i, text in enumerate(instructions, 1):
        guide.append([i, text])
        guide.row_dimensions[i + 1].height = 48
    guide.column_dimensions["A"].width = 8
    guide.column_dimensions["B"].width = 140
    ws = wb.create_sheet("盲评成对复核")
    columns = review_columns()
    indices = {k: i for i, (k, _) in enumerate(columns, 1)}
    ws.append([k for k, _ in columns])
    ws.append([v for _, v in columns])
    fixed_keys = set()
    for pair in mapping:
        values = fixed_values(samples[pair["product_id"]], pair, versions)
        fixed_keys = set(values)
        ws.append([values.get(k) for k, _ in columns])
    for sheet in (guide, ws):
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if sheet == ws:
                    editable = cell.row > 2 and columns[cell.column - 1][0] not in fixed_keys
                    cell.protection = Protection(locked=not editable)
                    cell.fill = PatternFill("solid", fgColor="FFF2CC" if editable else "E7E6E6")
                    if cell.row <= 2:
                        cell.fill = PatternFill("solid", fgColor="244062")
                        cell.font = Font(color="FFFFFF", bold=True)
    for key, index in indices.items():
        ws.column_dimensions[get_column_letter(index)].width = 55 if key.endswith(("title", "selling_points", "short_description", "complete")) else (44 if "attributes" in key or "notes" in key else 18)
        ws.cell(1, index).comment = Comment(columns[index - 1][1], "frozen test100")
        if key in ("A_complete", "B_complete", "source_title", "input_quality_actions"):
            ws.column_dimensions[get_column_letter(index)].hidden = True
    for row in range(3, 103):
        ws.row_dimensions[row].height = 230
        ws.cell(row, indices["product_id"]).number_format = "@"
    for preference in PREFERENCES:
        dv = DataValidation(type="list", formula1='"A更好,B更好,持平"', allow_blank=True)
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(indices[preference])}3:{get_column_letter(indices[preference])}102")
    for side in ("A", "B"):
        for field in (*METRICS, *DIAGNOSTICS):
            dv = DataValidation(type="whole", operator="greaterThanOrEqual" if field == "factual_error_count" else "between",
                formula1="0", formula2=None if field == "factual_error_count" else ("$I3" if field == "matched_attribute_count" else "1"), allow_blank=True)
            dv.showErrorMessage = True
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(indices[f'{side}_{field}'])}3:{get_column_letter(indices[f'{side}_{field}'])}102")
    dv = DataValidation(type="list", formula1='"1"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(indices['review_confirmed'])}3:{get_column_letter(indices['review_confirmed'])}102")
    ws.freeze_panes = "F3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(columns))}102"
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    hints = wb.create_sheet("自动线索_非人工指标")
    hints.append(["pair_id", "product_id", "side", "diagnostic_ONLY"])
    for pair in mapping:
        for side in ("A", "B"):
            sample, product = samples[pair["product_id"]], versions[pair[side]][pair["product_id"]]
            hints.append([pair["pair_id"], pair["product_id"], side, json.dumps({**automatic_hints(sample, product), **extra_hints(sample, product)}, ensure_ascii=False)])
    hints.sheet_state = "hidden"
    sealed = wb.create_sheet("匿名映射_汇总后解盲")
    sealed.append(["pair_id", "product_id", "A", "B"])
    for pair in mapping:
        sealed.append([pair[k] for k in ("pair_id", "product_id", "A", "B")])
    sealed.sheet_state = "veryHidden"
    wb.save(path)
    return path


def archive():
    cfg, inherited, frozen, samples, versions, mapping = paired_data()
    out = resolve(OUTPUT)
    if (out / "automatic_phase_manifest.json").exists() or (out / WORKBOOK).exists():
        raise FileExistsError("Never overwrite initial archive/workbook")
    summaries = {v: automatic_group(samples, versions[v]) for v in VARIANTS}
    by_category = {v: {c: automatic_group(samples, {pid: p for pid, p in versions[v].items() if samples[pid]["category_l2"] == c})
                    for c in sorted({s["category_l2"] for s in samples.values()})} for v in VARIANTS}
    result = {"status": "formal_test100_generation_complete_human_review_pending", "source_statistics": frozen["source_statistics"],
              "versions": summaries, "by_category_AUTOMATIC_ONLY": by_category, "paired_seed_prompt_preflight_checks": 300,
              "input_decoding_same": True, "human_official_metrics": None, "human_paired_preferences": None,
              "test_history_disclosure": cfg["test_history_disclosure"], "old_single_call_baseline_compared": False,
              "first_call_cold_start_included": True, "parameter_updates": 0, "rag_holdout_files_read": False}
    atomic_json(out / "automatic_phase_summary.json", result)
    latencies = []
    for v in VARIANTS:
        for pid, p in versions[v].items():
            latencies.append({"product_id": pid, "variant": v, **p["latency_seconds"]})
            append_jsonl(out / "automatic_diagnostics.jsonl", {"product_id": pid, "variant": v,
                **automatic_hints(samples[pid], p), **extra_hints(samples[pid], p)})
    with (out / "latency_by_product.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(latencies[0]))
        writer.writeheader()
        writer.writerows(latencies)
    workbook = build_workbook(out, samples, versions, mapping)
    output_paths = [p for v in VARIANTS for p in (out / v.lower()).iterdir() if p.is_file()]
    atomic_json(out / "review_manifest.json", {"status": "awaiting_final_human_blinded_review", "products": 100,
        "documents": 200, "initial_workbook_sha256": digest(workbook), "human_labels_prefilled": False,
        "protocol_manifest_sha256": digest(out / "protocol_manifest.json"),
        "files": [{"path": p.relative_to(resolve(".")).as_posix(), "sha256": digest(p)} for p in output_paths]})
    report = ["# LoRA v1 正式 test100 自动阶段报告", "", "状态：完整100件配对生成已完成，最终人工复核待完成。不是全200件test评测。", "",
        "测试样本此前用于Baseline/RAG，不宣称首次完全未观察测试。本轮Base与LoRA均为同配置三任务推理，不与旧单调用Baseline直接比较。", "",
        "| 自动辅助项 | Base | LoRA |", "| --- | ---: | ---: |"]
    for key, label in (("products", "商品数"), ("task_calls", "任务调用"), ("assembled_structure_success_NOT_FLUENCY", "组装结构成功"),
        ("empty_task_outputs", "空输出任务"), ("hit_token_limit_tasks", "max token任务"),
        ("claim_terms_REVIEW_ONLY_samples", "疑似评价词商品"), ("effect_scope_terms_REVIEW_ONLY_samples", "疑似效果/范围词商品"),
        ("scenario_terms_REVIEW_ONLY_samples", "场景词商品（可能有来源）"), ("number_terms_REVIEW_ONLY_samples", "疑似数值商品"),
        ("range_boundary_clues_REVIEW_ONLY_samples", "区间边界字面缺失线索商品")):
        report.append(f"| {label} | {summaries['Base'][key]} | {summaries['LoRA'][key]} |")
    for v in VARIANTS:
        s = summaries[v]
        report += ["", f"## {v} 自动诊断", "", f"三任务解析成功：{s['parse_success_by_task_NOT_FLUENCY']}。",
            f"字面覆盖：{s['literal_coverage_NOT_HUMAN']['matched_fields']}/{s['literal_coverage_NOT_HUMAN']['core_attribute_total']}（{s['literal_coverage_NOT_HUMAN']['ratio']:.2%}），不是人工属性命中率。",
            f"峰值allocated/reserved：{s['peak_allocated_mib']}/{s['peak_reserved_mib']}MiB。", "",
            "| 任务 | 平均秒 | P95秒 |", "| --- | ---: | ---: |"]
        for t, latency in s["latency"].items():
            report.append(f"| {t} | {latency['mean_seconds']:.3f} | {latency['p95_seconds']:.3f} |")
    report += ["", "## 下一步人工操作与结论边界", "", f"打开 `{WORKBOOK}`，对全100行分别填写两侧三正式指标、四匿名偏好、诊断及依据，reviewer、review_confirmed=1。",
        "评审完成前不要查看veryHidden映射或按版本命名的输出/报告。人工字段全部留空，自动线索不能代替人审。",
        "分品类自动结果见automatic_phase_summary.json；正式人工三指标、匿名偏好、分品类人工结果及重点bad case均待最终复核，当前不编造。",
        "未根据输出调参、重试、更新权重；不运行RAG v2或读取其holdout。保留原始解析失败输出。",
        "本轮自动完成后停止，等用户复核；后续仅汇总归档，不因test结果修改模型/Prompt。"]
    (out / "automatic_phase_report.md").write_text("\n".join(report) + "\n", encoding="utf-8")
    files = [{"path": p.relative_to(resolve(".")).as_posix(), "sha256": digest(p)} for p in sorted(out.rglob("*")) if p.is_file()]
    atomic_json(out / "automatic_phase_manifest.json", {"status": result["status"], "created_at_utc": now(), "files": files,
        "human_workbook_hash_role": "initial_blank_expected_to_change_by_user", "protocol_manifest_sha256": digest(out / "protocol_manifest.json"),
        "human_metrics_final": False, "parameter_updates": 0, "rag_holdout_files_read": False})
    (out / "automatic_phase_checksums.sha256").write_text("".join(f"{p['sha256']}  {p['path']}\n" for p in files)
        + f"{digest(out / 'automatic_phase_manifest.json')}  {(out / 'automatic_phase_manifest.json').relative_to(resolve('.')).as_posix()}\n", encoding="utf-8")
    print(json.dumps({"status": result["status"], "workbook": str(workbook), "versions": summaries}, ensure_ascii=False), flush=True)


def summarize():
    from openpyxl import load_workbook
    cfg, inherited, frozen, samples, versions, mapping = paired_data()
    out = resolve(OUTPUT)
    for name in ("blind_summary.json", "human_metrics_unblinded.json", "human_review_final.csv", "final_test_manifest.json"):
        if (out / name).exists():
            raise FileExistsError("Never overwrite final test results")
    manifest = read_json(out / "review_manifest.json")
    if digest(out / "protocol_manifest.json") != manifest["protocol_manifest_sha256"]:
        raise ValueError("Protocol changed during review")
    for item in manifest["files"]:
        if digest(resolve(item["path"])) != item["sha256"]:
            raise ValueError("Generation changed during review")
    path = out / WORKBOOK
    wb = load_workbook(path, data_only=False)
    ws = wb["盲评成对复核"]
    cols = review_columns()
    if ws.max_row != 102 or [c.value for c in ws[1]] != [k for k, _ in cols]:
        raise ValueError("Workbook scope/columns changed")
    if list(wb["匿名映射_汇总后解盲"].values)[1:] != [tuple(p[k] for k in ("pair_id", "product_id", "A", "B")) for p in mapping]:
        raise ValueError("Anonymous mapping changed")
    number = lambda v: str(int(v)) if isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v) and int(v) == v else str(v).strip()
    rows = []
    for i, pair in enumerate(mapping, 3):
        row = {k: ws.cell(i, j).value for j, (k, _) in enumerate(cols, 1)}
        if any(row[k] != v for k, v in fixed_values(samples[pair["product_id"]], pair, versions).items()):
            raise ValueError("Frozen source/output/denominator changed")
        if number(row["review_confirmed"]) != "1" or not row["reviewer"]:
            raise ValueError("Incomplete human confirmation: no unblinding")
        if any(row[p] not in CHOICES for p in PREFERENCES):
            raise ValueError("Incomplete paired preferences")
        for side in ("A", "B"):
            for m in METRICS:
                row[f"{side}_{m}"] = number(row[f"{side}_{m}"])
            official_metrics([{"product_id": row["product_id"], "core_attribute_count": row["core_attribute_count"],
                              **{m: row[f"{side}_{m}"] for m in METRICS}}])
            for d in DIAGNOSTICS:
                if row[f"{side}_{d}"] is not None:
                    row[f"{side}_{d}"] = parse_binary(number(row[f"{side}_{d}"]), d, row["product_id"])
            if (int(row[f"{side}_factual_error_count"]) > 0 or any(row[f"{side}_{d}"] == 1 for d in DIAGNOSTICS)) and not row[f"{side}_notes"]:
                raise ValueError("Errors/diagnoses require evidence")
        rows.append(row)
    atomic_json(out / "blind_summary.json", {"status": "anonymous_complete_before_unblinding", "products": 100,
        "preferences": {p: dict(Counter(r[p] for r in rows)) for p in PREFERENCES}, "workbook_sha256": digest(path)})
    judgments, votes = {v: [] for v in VARIANTS}, {p: Counter() for p in PREFERENCES}
    exported = []
    for row, pair in zip(rows, mapping):
        for side in ("A", "B"):
            item = {"product_id": row["product_id"], "category_l2": samples[row["product_id"]]["category_l2"],
                "variant": pair[side], "pair_id": pair["pair_id"], "anonymous_side": side,
                "core_attribute_count": row["core_attribute_count"], **{m: row[f"{side}_{m}"] for m in METRICS},
                **{d: row[f"{side}_{d}"] for d in DIAGNOSTICS}, **{p: row[p] for p in PREFERENCES},
                "notes": row[f"{side}_notes"] or "", "pair_notes": row["pair_notes"] or "", "reviewer": row["reviewer"]}
            judgments[pair[side]].append(item)
            exported.append(item)
        for p in PREFERENCES:
            votes[p]["tie" if row[p] == "持平" else pair[row[p][0]]] += 1
    result = {"status": "final_formal_test100_user_confirmed_human_review", "products_per_model": 100,
        "official_metrics": {v: official_metrics(judgments[v]) for v in VARIANTS},
        "paired_preferences": {p: {v: votes[p][v] for v in (*VARIANTS, "tie")} for p in PREFERENCES},
        "by_category": {v: {c: official_metrics([r for r in judgments[v] if r["category_l2"] == c]) for c in sorted({r["category_l2"] for r in judgments[v]})} for v in VARIANTS},
        "diagnostics_NOT_FORMAL_METRICS": {v: {d: {"annotated": sum(r[d] is not None for r in judgments[v]), "positive": sum(r[d] == 1 for r in judgments[v])} for d in DIAGNOSTICS} for v in VARIANTS},
        "workbook_sha256": digest(path), "blind_summary_sha256": digest(out / "blind_summary.json"),
        "test_history_disclosure": cfg["test_history_disclosure"], "parameter_updates": 0}
    atomic_json(out / "human_metrics_unblinded.json", result)
    with (out / "human_review_final.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(exported[0]))
        writer.writeheader()
        writer.writerows(exported)
    # Diagnostic examples after review, never used to change the frozen candidate.
    cases = {v: sorted(judgments[v], key=lambda r: (-int(r["factual_error_count"]), -(int(r["core_attribute_count"]) - int(r["matched_attribute_count"])), r["product_id"]))[:8] for v in VARIANTS}
    atomic_json(out / "human_bad_cases.json", cases)
    report = ["# LoRA v1 最终test100人工报告", "", "完整100对最终确认，指标与validation口径一致。", "",
              "既有test100此前用于Baseline/RAG，本轮不是首次未观察测试；不根据结果调参或重跑。", "",
              "| 人工指标 | Base | LoRA |", "| --- | ---: | ---: |"]
    for key in ("core_attribute_hit_rate", "fluency_pass_rate", "factual_error_sample_rate", "factual_error_total", "average_factual_errors_per_sample"):
        report.append(f"| {key} | {result['official_metrics']['Base'][key]} | {result['official_metrics']['LoRA'][key]} |")
    report += ["", "## 匿名偏好解盲", "", "```json", json.dumps(result["paired_preferences"], ensure_ascii=False, indent=2), "```",
               "", "分品类和诊断见human_metrics_unblinded.json，人工确认的重点案例见human_bad_cases.json（排序只用于报告，不用于挑评测样本）。",
               "自动指标、每任务和总延迟见automatic_phase_summary.json。人工标签不由自动线索替代。", "",
               "至此只归档，不根据test修改模型、Prompt、训练数据或推理参数。"]
    (out / "final_test_report.md").write_text("\n".join(report) + "\n", encoding="utf-8")
    files = [{"path": p.relative_to(resolve(".")).as_posix(), "sha256": digest(p)} for p in sorted(out.rglob("*")) if p.is_file()]
    atomic_json(out / "final_test_manifest.json", {"status": result["status"], "created_at_utc": now(), "files": files,
        "protocol_manifest_sha256": digest(out / "protocol_manifest.json"), "no_retuning_after_test": True})
    print("Final test100 human report saved. No tuning or reruns.")
