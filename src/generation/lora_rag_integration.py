"""Once-only, frozen LoRA + four-zone RAG integration on known development IDs."""
from __future__ import annotations

import copy
import csv
import hashlib
import importlib.metadata
import json
import math
import random
import statistics
import subprocess
import sys
import time
from collections import Counter, defaultdict

from .grounding import validate_generation_grounding
from .lora_training import ROOT, atomic_json, digest, json_digest, memory, now, read_json, resolve
from .lora_validation import (TASK_TYPES, PACKAGE_NAMES, append_jsonl, assemble,
                              load_jsonl, model_snapshot, parse_task, seed_for,
                              verify_protocol as verify_parent)
from .lora_validation_review import (CHOICES, DIAGNOSTICS, GUIDE, METRICS, PREFERENCES,
                                    automatic_hints, fixed_values, official_metrics, review_columns)
from .qwen import preflight_chat_prompt
from .qlora_memory import prepare_kbit_with_cpu_staged_large_layers
from .rag import select_rag_v2_context

CONFIG = "configs/lora_rag_v2_development24_v1.json"
DOCUMENT = "docs/lora_rag_v2_development24_v1_protocol.md"
OUTPUT = "reports/generation/lora_rag/lora_rag_v2_development24_v1"
VARIANTS = ("LoRA_control", "LoRA_RAG")
CODE = ("scripts/run_lora_rag_integration.py", "src/generation/lora_rag_integration.py",
        "tests/test_lora_rag_integration.py")
ALLOWED = {
    "parent_config": "configs/lora_project_validation_v1.json",
    "parent_archive": "reports/generation/lora/project_validation_v1/automatic_phase_manifest.json",
    "parent_inputs": "reports/generation/lora/project_validation_v1/frozen_inputs.jsonl",
    "control_outputs": "reports/generation/lora/project_validation_v1/lora/product_outputs.jsonl",
    "control_report": "reports/generation/lora/project_validation_v1/lora/run_report.json",
}
ENTRIES = {
    "parent_protocol": ("reports/generation/lora/project_validation_v1/protocol_manifest.json", "095cb45a3c07b86ccace377d746ffe3224a9805aaf5204f19640652e08aa1ab2"),
    "development_sample": ("reports/generation/rag/rag_v2_p2_development24.json", "78accdb37ae70916559368d32ebff439ed8f35f7998cba3d43249787eb50a8c0"),
    "fact_policy": ("configs/rag_fact_policy_v4.json", "b032bf8835a4bfa11338a6bfddba05392aaa896c103473aa402b4ead25029fea"),
}


def checked_config():
    cfg = read_json(resolve(CONFIG))
    # Resolve every permitted path before hashing any external file.
    for key, name in ALLOWED.items():
        if resolve(cfg[key]) != resolve(name):
            raise ValueError(f"Forbidden path before read: {key}")
    for key, (name, sha) in ENTRIES.items():
        if cfg[key]["path"] != name or cfg[key]["sha256"] != sha:
            raise ValueError(f"Forbidden entry before read: {key}")
    expected_guards = {key: False for key in ("parameter_updates", "project_test_read", "holdout_archive_read",
                       "holdout_generation", "old_rag_output_as_model_input", "content_tuning", "p3_entered")}
    if (cfg["output"] != OUTPUT or cfg["guards"] != expected_guards
            or cfg["scope"] != "independent_integration_validation_development24_only"):
        raise ValueError("Integration scope/guards changed")
    if (cfg["variants"] != list(VARIANTS) or cfg["rag"]["top_k"] != 3
            or cfg["prompt"]["max_input_tokens"] != 2048 or cfg["prompt"]["truncation"]
            or cfg["development_sample"]["products"] != 24 or cfg["generation"]["resume"]):
        raise ValueError("Integration design changed")
    for key, (name, sha) in ENTRIES.items():
        if digest(resolve(name)) != sha:
            raise ValueError(f"Frozen {key} changed")
    parent, frozen = verify_parent()  # Only P7 train/validation provenance; never test/holdout.
    archived = {e["path"]: e["sha256"] for e in read_json(resolve(cfg["parent_archive"]))["files"]}
    for key in ("parent_inputs", "control_outputs", "control_report"):
        name = cfg[key]
        if archived.get(name) != digest(resolve(name)):
            raise ValueError(f"P7 archived file changed: {key}")
    return cfg, parent, frozen


def select_samples(development, parent_inputs):
    if (development.get("status") != "completed"
            or development.get("generation_output_holdout_executed") is not False
            or development["sample_selection"]["protected_holdout_overlap_count"] != 0):
        raise ValueError("Development non-holdout provenance is missing")
    ids = [str(r["product_id"]) for r in development["results"]]
    by_id = {s["product_id"]: s for s in parent_inputs}
    if len(ids) != 24 or len(set(ids)) != 24 or not set(ids) <= set(by_id):
        raise ValueError("Wrong/missing development24 IDs")
    samples = [copy.deepcopy(by_id[pid]) for pid in ids]
    counts = Counter(s["category_l2"] for s in samples)
    if len(counts) != 8 or set(counts.values()) != {3}:
        raise ValueError("Must retain all 3 products per 8 categories")
    if any(s["source_quality"]["status"] != "PASS" for s in samples):
        raise ValueError("Known development source status changed")
    return samples


def assert_context(sample, context):
    pid = sample["product_id"]
    ids = []
    for zone in ("identity_facts", "mandatory_core_facts", "supplemental_facts", "negative_constraint_facts"):
        for fact in context[zone]:
            if str(fact["product_id"]) != pid or fact["quality_status"] != "eligible":
                raise ValueError("Cross-product or ineligible RAG injection")
            if (fact.get("is_core_attribute") and fact["canonical_field"] in sample["evaluation_attributes"]
                    and fact["canonical_field"] not in sample["used_attributes"]):
                raise ValueError("RAG cannot restore a P7-withheld core field")
            ids.append(fact["fact_id"])
    if len(ids) != len(set(ids)) or len(context["supplemental_facts"]) > 3:
        raise ValueError("Duplicate RAG zones or wrong Top-3")
    if context["product_id"] != pid or context["mandatory_injection_expected_fact_ids"] != [f["fact_id"] for f in context["mandatory_core_facts"]]:
        raise ValueError("Mandatory expectation mismatch")


def gated_context(sample, facts, policy):
    """Intersect the frozen RAG gate with P7 withdrawals before ranking Top-3."""
    allowed, withheld = [], []
    for fact in facts:
        if str(fact["product_id"]) != sample["product_id"]:
            raise ValueError("Cross-product facts before selection")
        blocked = (sample["source_quality"].get("identity_block", False)
                   or (fact.get("is_core_attribute") and fact["canonical_field"] in sample["evaluation_attributes"]
                       and fact["canonical_field"] not in sample["used_attributes"]))
        if blocked:
            withheld.append({"fact_id": fact["fact_id"], "field": fact["canonical_field"],
                             "reason": "already_withheld_in_frozen_P7_input"})
        else:
            allowed.append(fact)
    context = select_rag_v2_context(allowed, policy, top_k=3)
    assert_context(sample, context)
    return context, withheld


def fact_line(fact):
    return f"[{fact['fact_id']}] {fact['canonical_field']}={'、'.join(map(str, fact['normalized_values']))}"


def integration_messages(sample, task, context, cfg):
    assert_context(sample, context)
    # Exactly retain the frozen task instruction and original core payload.
    lines = ["", "同商品检索事实（仅补充输入事实，不是输出内容）："]
    for label, key in (("Identity Facts（身份信息）", "identity_facts"),
                       ("Mandatory Reliable Core Facts（输入完整提供，不要求各任务重复）", "mandatory_core_facts"),
                       ("Supplemental Top-3（可选补充细节）", "supplemental_facts"),
                       ("Negative Constraints（只读约束，不要求写成卖点）", "negative_constraint_facts")):
        lines.append(label + "：")
        lines.extend(fact_line(f) for f in context[key])
    return [{"role": "system", "content": cfg["prompt"]["system"]},
            {"role": "user", "content": sample["tasks"][task]["messages"][1]["content"] + "\n".join(lines)}]


def preflight(tokenizer, sample, task, cfg):
    context = sample["rag_context"]
    all_lines = [fact_line(f) for key in ("identity_facts", "mandatory_core_facts", "supplemental_facts", "negative_constraint_facts") for f in context[key]]
    core = json.dumps(sample["used_attributes"], ensure_ascii=False, sort_keys=False)
    inputs, meta = preflight_chat_prompt(tokenizer, sample["rag_tasks"][task]["messages"],
                                        max_input_tokens=cfg["prompt"]["max_input_tokens"], mandatory_lines=[core, *all_lines])
    # Base preflight checked every serialized fact after tokenization, not just a text count.
    expected = len(context["mandatory_core_facts"])
    identity = len(context["identity_mandatory_coverage_fact_ids"])
    meta.update(mandatory_core_expected=expected, mandatory_core_verified=expected,
                mandatory_core_injection_rate=1.0 if expected else None,
                identity_core_expected=identity, identity_core_verified=identity,
                original_core_payload_complete=True, all_rag_facts_verified=len(all_lines))
    return inputs, meta


def mapping_for(samples, seed):
    groups = defaultdict(list)
    for sample in samples:
        groups[sample["category_l2"]].append(sample)
    rank = lambda ns, pid: hashlib.sha256(f"{ns}:{seed}:{pid}".encode()).hexdigest()
    controls_as_a = set()
    for index, category in enumerate(sorted(groups)):
        rows = sorted(groups[category], key=lambda s: rank("integration_ab", s["product_id"]))
        controls_as_a.update(s["product_id"] for s in rows[:1 + index % 2])
    order = sorted(samples, key=lambda s: rank("integration_order", s["product_id"]))
    return [{"pair_id": f"IR{i:03d}", "product_id": s["product_id"],
             "A": VARIANTS[0] if s["product_id"] in controls_as_a else VARIANTS[1],
             "B": VARIANTS[1] if s["product_id"] in controls_as_a else VARIANTS[0]} for i, s in enumerate(order, 1)]


def verify_control(sample, product):
    if product["variant"] != "LoRA" or product["product_id"] != sample["product_id"]:
        raise ValueError("Wrong source control model/product")
    for task in TASK_TYPES:
        expected, actual = sample["tasks"][task], product["tasks"][task]
        for key in ("seed", "prompt_sha256", "preflight"):
            if actual[key] != expected[key]:
                raise ValueError(f"Control frozen {key} differs")
        if expected["seed"] != seed_for(sample["product_id"], task):
            raise ValueError("Control seed rule changed")


def prepare():
    import torch
    from transformers import AutoTokenizer
    cfg, parent, parent_manifest = checked_config()
    out = resolve(OUTPUT)
    if out.exists():
        raise FileExistsError("Never overwrite an existing integration freeze")
    samples = select_samples(read_json(resolve(cfg["development_sample"]["path"])), load_jsonl(resolve(cfg["parent_inputs"])))
    facts = load_jsonl(resolve(parent["validation_quality_facts"]["path"]))
    policy = read_json(resolve(cfg["fact_policy"]["path"]))
    if policy["version"] != "rag_fact_policy_v4":
        raise ValueError("Wrong RAG policy")
    facts_by_id = defaultdict(list)
    for fact in facts:
        if (fact["split"] != "validation" or fact["source_path"] != parent["source"]["path"]
                or fact["source_sha256"] != parent["source"]["sha256"]):
            raise ValueError("RAG facts not from frozen project_validation")
        facts_by_id[str(fact["product_id"])].append(fact)
    controls = load_jsonl(resolve(cfg["control_outputs"]))
    by_id = {r["product_id"]: r for r in controls}
    source_report = read_json(resolve(cfg["control_report"]))
    if len(controls) != 200 or len(by_id) != 200 or source_report["status"] != "passed" or source_report["adapter_loaded"] is not True:
        raise ValueError("Control generation incomplete")
    tokenizer = AutoTokenizer.from_pretrained(str(model_snapshot(parent)), local_files_only=True)
    if json_digest(tokenizer.chat_template) != parent_manifest["chat_template_sha256"]:
        raise ValueError("Chat template changed")
    selected_controls, disagreements = [], []
    for sample in samples:
        verify_control(sample, by_id[sample["product_id"]])
        started = time.perf_counter()
        context, withheld = gated_context(sample, facts_by_id[sample["product_id"]], policy)
        sample["rag_selection_seconds"] = time.perf_counter() - started
        assert_context(sample, context)
        sample["integration_gate_withheld"] = withheld
        sample["rag_context"] = context
        sample["rag_tasks"] = {}
        for task in TASK_TYPES:
            messages = integration_messages(sample, task, context, cfg)
            sample["rag_tasks"][task] = {"messages": messages, "seed": sample["tasks"][task]["seed"], "prompt_sha256": json_digest(messages)}
            _, meta = preflight(tokenizer, sample, task, cfg)
            sample["rag_tasks"][task]["preflight"] = meta
            # Recheck the original token stream before reusing any old output.
            _, control_meta = preflight_chat_prompt(tokenizer, sample["tasks"][task]["messages"],
                max_input_tokens=parent["prompt"]["max_input_tokens"],
                mandatory_lines=[json.dumps(sample["used_attributes"], ensure_ascii=False, sort_keys=False)])
            if control_meta != sample["tasks"][task]["preflight"]:
                raise ValueError("Original control tokenizer/preflight changed")
        disagreements.extend({"product_id": sample["product_id"], "field": f["canonical_field"], "rag_quality": f["quality_status"],
                              "rag_reasons": f["quality_reasons"], "action": "retain original P7 core unchanged; do not add this fact in RAG zones"}
                             for f in facts_by_id[sample["product_id"]] if f["quality_status"] != "eligible" and f["canonical_field"] in sample["used_attributes"])
        control = copy.deepcopy(by_id[sample["product_id"]])
        control["source_variant"] = control["variant"]
        control["variant"] = VARIANTS[0]
        control["generation_reused"] = True
        control["rag_selection_seconds"] = 0.0
        selected_controls.append(control)
    seen_categories = set()
    smoke_ids = []
    for s in samples:
        if s["category_l2"] not in seen_categories:
            smoke_ids.append(s["product_id"])
            seen_categories.add(s["category_l2"])
    environment = {"python": sys.version, "packages": {n: importlib.metadata.version(n) for n in PACKAGE_NAMES}, "cuda": torch.version.cuda}
    if environment != parent_manifest["environment"]:
        raise ValueError("Do not reuse control with a changed inference environment")
    out.mkdir(parents=True)
    for sample in samples:
        append_jsonl(out / "frozen_inputs.jsonl", sample)
    for row in selected_controls:
        append_jsonl(out / "control_product_outputs.jsonl", row)
    atomic_json(out / "sample_ids.json", {"development24": [s["product_id"] for s in samples], "smoke8": smoke_ids,
        "source_selection": "unchanged archived development24; already observed development, not a fresh holdout", "holdout_overlap": "0 per hash-pinned P2 selection provenance; holdout manifest not opened"})
    atomic_json(out / "blind_mapping.json", {"status": "sealed_until_complete_review", "pairs": mapping_for(samples, cfg["pairing"]["blind_seed"])})
    atomic_json(out / "inference_config.json", {"model": parent["model"], "adapter": parent["adapter"],
        "decoding": parent["decoding"], "decoding_resolved": read_json(resolve("reports/generation/lora/project_validation_v1/inference_lora.json"))["decoding_resolved"],
        "original_prompt": parent["prompt"], "integration_prompt": cfg["prompt"], "pairing": parent["pairing"], "chat_template_sha256": parent_manifest["chat_template_sha256"]})
    freeze = subprocess.run([sys.executable, "-m", "pip", "freeze"], capture_output=True, text=True, check=True).stdout
    (out / "environment_freeze.txt").write_text(freeze, encoding="utf-8")
    git = lambda *args: subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True, check=True, encoding="utf-8").stdout.strip()
    paths = {e["path"] for e in parent_manifest["files"]}
    paths.update((*CODE, CONFIG, DOCUMENT, cfg["parent_archive"], *ALLOWED.values(), *(v[0] for v in ENTRIES.values())))
    paths.update(str(p.relative_to(ROOT)).replace("\\", "/") for p in out.iterdir() if p.is_file())
    metadata = {name: digest(model_snapshot(parent) / name) for name in ("config.json", "tokenizer_config.json", "generation_config.json")}
    if metadata != parent_manifest["model_metadata_sha256"]:
        raise ValueError("Frozen base metadata changed")
    manifest = {"version": cfg["version"], "status": "frozen_before_new_generation", "created_at_utc": now(), "config": cfg,
        "parent_config": parent, "git_head": git("rev-parse", "HEAD"), "git_status": git("-c", "core.quotepath=false", "status", "--short"),
        "code_original_bytes_hex": {p: resolve(p).read_bytes().hex() for p in CODE},
        "files": [{"path": p, "sha256": digest(resolve(p))} for p in sorted(paths)],
        "environment": environment, "model_metadata_sha256": metadata, "chat_template_sha256": parent_manifest["chat_template_sha256"],
        "prompt_bundle_sha256": json_digest([{ "product_id": s["product_id"], "tasks": s["rag_tasks"]} for s in samples]),
        "statistics": {"products": 24, "by_category": dict(Counter(s["category_l2"] for s in samples)), "original_core_denominator": sum(s["core_attribute_count"] for s in samples),
            "original_input_core_fields": sum(len(s["used_attributes"]) for s in samples),
            "mandatory_core_facts_per_product_sum": sum(len(s["rag_context"]["mandatory_core_facts"]) for s in samples),
            "identity_core_facts_sum": sum(len(s["rag_context"]["identity_mandatory_coverage_fact_ids"]) for s in samples),
            "max_input_tokens": max(t["preflight"]["input_token_count"] for s in samples for t in s["rag_tasks"].values())},
        "existing_input_vs_rag_gate_disagreements": disagreements, "guards": cfg["guards"],
        "integration_gate_withheld": [{"product_id": s["product_id"], **f} for s in samples for f in s["integration_gate_withheld"]],
        "historical_rag_v2_stage": "unchanged stopped; original P3 not entered; 32 holdout remains frozen_not_executed",
        "comparison_limit": "RAG-context plus necessary interface treatment, not pure retrieval-algorithm isolation; control latency is historical"}
    atomic_json(out / "protocol_manifest.json", manifest)
    print(json.dumps({"status": manifest["status"], **manifest["statistics"]}, ensure_ascii=False), flush=True)


def verify_protocol():
    cfg, parent, _ = checked_config()
    frozen = read_json(resolve(OUTPUT) / "protocol_manifest.json")
    if frozen["config"] != cfg or frozen["parent_config"] != parent:
        raise ValueError("Frozen integration configuration changed")
    for item in frozen["files"]:
        if digest(resolve(item["path"])) != item["sha256"]:
            raise ValueError(f"Frozen integration file changed: {item['path']}")
    if {n: digest(model_snapshot(parent) / n) for n in frozen["model_metadata_sha256"]} != frozen["model_metadata_sha256"]:
        raise ValueError("Frozen local base metadata changed")
    return cfg, parent, frozen


def generate(*, smoke=False):
    cfg, parent, frozen = verify_protocol()
    out = resolve(OUTPUT)
    folder = out / ("smoke8" if smoke else "lora_rag")
    if folder.exists():
        raise FileExistsError("No automatic rerun or resume of integration generation")
    if not smoke and read_json(out / "smoke8/run_report.json")["status"] != "passed":
        raise ValueError("Engineering smoke must pass first")
    folder.mkdir()
    report = {"status": "running", "started_at_utc": now(), "scope": "engineering_only" if smoke else "development24",
              "parameter_updates": 0, "project_test_read": False, "holdout_archive_read": False, "holdout_generation": False}
    started_run = time.perf_counter()
    try:
        import torch
        from peft import PeftModel, prepare_model_for_kbit_training
        from transformers import AutoModelForCausalLM, AutoTokenizer, BitsAndBytesConfig, GenerationConfig
        environment = {"python": sys.version, "packages": {n: importlib.metadata.version(n) for n in PACKAGE_NAMES}, "cuda": torch.version.cuda}
        if environment != frozen["environment"] or not torch.cuda.is_available() or not torch.cuda.is_bf16_supported():
            raise RuntimeError("Frozen CUDA/BF16 environment changed; no fallback")
        snapshot = model_snapshot(parent)
        tokenizer = AutoTokenizer.from_pretrained(str(snapshot), local_files_only=True)
        if json_digest(tokenizer.chat_template) != frozen["chat_template_sha256"]:
            raise ValueError("Chat template changed")
        generation = GenerationConfig.from_dict(read_json(out / "inference_config.json")["decoding_resolved"])
        random.seed(42)
        torch.manual_seed(42)
        torch.cuda.manual_seed_all(42)
        quant = BitsAndBytesConfig(load_in_4bit=True, bnb_4bit_quant_type="nf4", bnb_4bit_use_double_quant=True, bnb_4bit_compute_dtype=torch.bfloat16)
        model = AutoModelForCausalLM.from_pretrained(str(snapshot), local_files_only=True, device_map={"": 0},
            quantization_config=quant, dtype=torch.bfloat16, low_cpu_mem_usage=True)
        model, _ = prepare_kbit_with_cpu_staged_large_layers(model, prepare_model_for_kbit_training, use_gradient_checkpointing=False)
        model.config.use_cache = True
        model = PeftModel.from_pretrained(model, str(resolve(parent["adapter"]["path"])), is_trainable=False, local_files_only=True)
        model.eval()
        if any(p.requires_grad for p in model.parameters()) or not model.is_loaded_in_4bit:
            raise RuntimeError("Expected frozen 4-bit inference only")
        before = {n: p._version for n, p in model.named_parameters()}
        samples = load_jsonl(out / "frozen_inputs.jsonl")
        if smoke:
            smoke_ids = set(read_json(out / "sample_ids.json")["smoke8"])
            samples = [s for s in samples if s["product_id"] in smoke_ids]
        report.update(gpu=torch.cuda.get_device_name(0), base_snapshot=str(snapshot), adapter_loaded=True, model_load_seconds=time.perf_counter() - started_run)
        for index, sample in enumerate(samples, 1):
            tasks = {}
            for task in TASK_TYPES:
                expected = sample["rag_tasks"][task]
                tick = time.perf_counter()
                inputs, meta = preflight(tokenizer, sample, task, cfg)
                preflight_seconds = time.perf_counter() - tick
                if meta != expected["preflight"]:
                    raise ValueError("Frozen tokenizer preflight changed")
                inputs = {k: v.to("cuda") for k, v in inputs.items()}
                random.seed(expected["seed"])
                torch.manual_seed(expected["seed"])
                torch.cuda.manual_seed_all(expected["seed"])
                torch.cuda.reset_peak_memory_stats()
                torch.cuda.synchronize()
                tick = time.perf_counter()
                with torch.no_grad(), torch.autocast("cuda", dtype=torch.bfloat16):
                    generated = model.generate(**inputs, generation_config=generation)
                torch.cuda.synchronize()
                latency = time.perf_counter() - tick
                ids = generated[0, inputs["input_ids"].shape[1]:].tolist()
                text = tokenizer.decode(ids, skip_special_tokens=True).strip()
                row = {"product_id": sample["product_id"], "variant": VARIANTS[1], "task_type": task,
                    "seed": expected["seed"], "prompt_sha256": expected["prompt_sha256"], "preflight": meta,
                    "preflight_seconds": preflight_seconds, "raw_text": text, "parsed": parse_task(task, text),
                    "generated_tokens": len(ids), "hit_max_new_tokens": len(ids) >= parent["decoding"]["max_new_tokens"],
                    "latency_seconds": latency, "memory": memory()}
                if not text:
                    raise RuntimeError("Empty output; stop without retries")
                tasks[task] = row
                append_jsonl(folder / "task_outputs.jsonl", row)
                del generated, inputs
            product = assemble(sample, tasks, VARIANTS[1])
            product["rag_selection_seconds"] = sample["rag_selection_seconds"]
            product["latency_seconds"]["three_task_plus_retrieval"] = product["latency_seconds"]["three_task_total"] + sample["rag_selection_seconds"]
            product["latency_seconds"]["three_task_plus_retrieval_and_preflight"] = product["latency_seconds"]["three_task_plus_retrieval"] + sum(t["preflight_seconds"] for t in tasks.values())
            append_jsonl(folder / "product_outputs.jsonl", product)
            print(json.dumps({"phase": "smoke" if smoke else "development24", "products_complete": index, "total": len(samples),
                              "elapsed_seconds": round(time.perf_counter() - started_run, 2)}, ensure_ascii=False), flush=True)
        if before != {n: p._version for n, p in model.named_parameters()} or any(p.grad is not None for p in model.parameters()):
            raise RuntimeError("Inference changed parameters or accumulated gradients")
        verify_protocol()
        report.update(status="passed", completed_at_utc=now(), products=len(samples), task_calls=len(samples) * 3,
                      parsed_tasks=sum(r["parsed"]["success"] for r in load_jsonl(folder / "task_outputs.jsonl")),
                      model_parameters_unchanged=True, wall_seconds=time.perf_counter() - started_run)
        atomic_json(folder / "run_report.json", report)
    except Exception as error:
        report.update(status="failed", error=str(error), failed_at_utc=now())
        atomic_json(folder / "run_report.json", report)
        raise


def paired_data():
    cfg, _, frozen = verify_protocol()
    out = resolve(OUTPUT)
    report = read_json(out / "lora_rag/run_report.json")
    if report["status"] != "passed" or report["products"] != 24 or report["task_calls"] != 72:
        raise ValueError("Complete generation required before review")
    samples = {s["product_id"]: s for s in load_jsonl(out / "frozen_inputs.jsonl")}
    versions = {}
    for variant, path in zip(VARIANTS, (out / "control_product_outputs.jsonl", out / "lora_rag/product_outputs.jsonl")):
        rows = load_jsonl(path)
        versions[variant] = {r["product_id"]: r for r in rows}
        if len(rows) != 24 or set(versions[variant]) != set(samples):
            raise ValueError("Missing/duplicate generated products")
        for row in rows:
            sample = samples[row["product_id"]]
            for task in TASK_TYPES:
                item = row["tasks"][task]
                expected = sample["tasks" if variant == VARIANTS[0] else "rag_tasks"][task]
                if any(item[k] != expected[k] for k in ("seed", "prompt_sha256", "preflight")):
                    raise ValueError("Output seed/input/preflight changed")
    independent = load_jsonl(out / "lora_rag/task_outputs.jsonl")
    if len(independent) != 72 or len({(t["product_id"], t["task_type"]) for t in independent}) != 72:
        raise ValueError("Independent task coverage mismatch")
    for task in independent:
        if task != versions[VARIANTS[1]][task["product_id"]]["tasks"][task["task_type"]]:
            raise ValueError("Independent task differs from assembled task")
    mapping = read_json(out / "blind_mapping.json")["pairs"]
    if mapping != mapping_for(list(samples.values()), cfg["pairing"]["blind_seed"]):
        raise ValueError("Anonymous mapping changed")
    return cfg, frozen, samples, versions, mapping


def grounding_context(sample, *, with_rag):
    context = copy.deepcopy(sample["rag_context"]) if with_rag else {"identity_facts": [], "mandatory_core_facts": [], "supplemental_facts": [], "negative_constraint_facts": []}
    # The validator must see facts present in ORIGINAL P7 core input as well as RAG.
    context["selected_facts"] = [{"fact_id": f"original-core:{sample['product_id']}:{field}", "canonical_field": field,
        "normalized_values": values, "product_id": sample["product_id"], "quality_status": "eligible"}
        for field, values in sample["used_attributes"].items()]
    if not with_rag:
        context["identity_facts"] = [{"fact_id": f"category:{field}", "canonical_field": field, "normalized_values": [sample[field]]}
                                     for field in ("category_l1", "category_l2")]
    return context


def percentile(values, q):
    values = sorted(values)
    pos = (len(values) - 1) * q
    return values[math.floor(pos)] + (values[math.ceil(pos)] - values[math.floor(pos)]) * (pos - math.floor(pos))


def automatic_summary(samples, versions):
    summaries, diagnostics = {}, {}
    denominator = sum(s["core_attribute_count"] for s in samples.values())
    for variant in VARIANTS:
        rows = list(versions[variant].values())
        tasks = [r["tasks"][t] for r in rows for t in TASK_TYPES]
        hints = {r["product_id"]: automatic_hints(samples[r["product_id"]], r) for r in rows}
        grounding = {}
        for r in rows:
            output = {"generated_title": r["tasks"]["title"]["raw_text"],
                      "selling_points": r["assembled"]["selling_points"] or [r["tasks"]["selling_points"]["raw_text"]],
                      "short_description": r["tasks"]["short_description"]["raw_text"]}
            grounding[r["product_id"]] = validate_generation_grounding(output, grounding_context(samples[r["product_id"]], with_rag=variant == VARIANTS[1]))
        diagnostics[variant] = {pid: {"automatic_hints": hints[pid], "grounding_REVIEW_ONLY": grounding[pid]} for pid in samples}
        matched = sum(h["literal_matched_count_NOT_HUMAN"] for h in hints.values())
        latency = {t: [r["latency_seconds"][t] for r in rows] for t in (*TASK_TYPES, "three_task_total")}
        if variant == VARIANTS[1]:
            latency["three_task_plus_retrieval"] = [r["latency_seconds"]["three_task_plus_retrieval"] for r in rows]
            latency["three_task_plus_retrieval_and_preflight"] = [r["latency_seconds"]["three_task_plus_retrieval_and_preflight"] for r in rows]
        expected = sum(t["preflight"].get("mandatory_core_expected", 0) for t in tasks)
        verified = sum(t["preflight"].get("mandatory_core_verified", 0) for t in tasks)
        summaries[variant] = {"products": len(rows), "task_parse_success": {t: sum(r["tasks"][t]["parsed"]["success"] for r in rows) for t in TASK_TYPES},
            "assembled_structure_success": sum(r["assembled"]["structure_success"] for r in rows),
            "literal_core_coverage_NOT_HUMAN": {"matched": matched, "denominator": denominator, "rate": matched / denominator},
            "latency_seconds": {k: {"mean": statistics.mean(v), "p95": percentile(v, .95), "max": max(v)} for k, v in latency.items()},
            "retrieval_selection_mean_seconds": statistics.mean(r.get("rag_selection_seconds", 0) for r in rows),
            "empty_tasks": sum(not t["raw_text"].strip() for t in tasks), "max_token_tasks": sum(t["hit_max_new_tokens"] for t in tasks),
            "claim_suspect_samples_NOT_FACTUAL_RATE": sum(bool(h["claim_terms_REVIEW_ONLY"]) for h in hints.values()),
            "number_suspect_samples_NOT_FACTUAL_RATE": sum(bool(h["number_terms_REVIEW_ONLY"]) for h in hints.values()),
            "grounding_pass_NOT_HUMAN": sum(g["status"] == "PASS" for g in grounding.values()),
            "peak_allocated_mib": max(t["memory"]["peak_allocated_mib"] for t in tasks),
            "peak_reserved_mib": max(t["memory"]["peak_reserved_mib"] for t in tasks),
            "input_tokens_mean": statistics.mean(t["preflight"]["input_token_count"] for t in tasks),
            "input_tokens_max": max(t["preflight"]["input_token_count"] for t in tasks),
            "mandatory_core_injection_per_task": {"expected": expected, "verified": verified, "rate": verified / expected if expected else None},
            "by_category": {c: {"products": sum(samples[r["product_id"]]["category_l2"] == c for r in rows),
                "literal_matched": sum(hints[r["product_id"]]["literal_matched_count_NOT_HUMAN"] for r in rows if samples[r["product_id"]]["category_l2"] == c),
                "core_denominator": sum(samples[r["product_id"]]["core_attribute_count"] for r in rows if samples[r["product_id"]]["category_l2"] == c)}
                for c in sorted({s["category_l2"] for s in samples.values()})}}
    return summaries, diagnostics


def review_fixed(sample, pair, versions):
    values = fixed_values(sample, pair, versions)
    for side in ("A", "B"):
        variant = pair[side]
        rag = variant == VARIANTS[1]
        values[f"{side}_input_facts"] = json.dumps({"retained_core": sample["used_attributes"],
            "rag_context": sample["rag_context"] if rag else None}, ensure_ascii=False)
        values[f"{side}_task_messages"] = json.dumps({t: sample["rag_tasks" if rag else "tasks"][t]["messages"] for t in TASK_TYPES}, ensure_ascii=False)
    return values


def integration_columns():
    return review_columns() + [(f"{side}_{key}", f"{side}：{label}") for side in ("A", "B") for key, label in
                              (("input_facts", "实际输入事实（原核心＋检索四区）"), ("task_messages", "真实三任务Prompt；仅需要时查看"))]


def build_review():
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    cfg, frozen, samples, versions, mapping = paired_data()
    out = resolve(OUTPUT)
    path = out / "lora_rag_development24_blind_review.xlsx"
    if path.exists() or (out / "automatic_summary.json").exists():
        raise FileExistsError("Never overwrite integration review/results")
    summaries, diagnostics = automatic_summary(samples, versions)
    wb = Workbook()
    guide = wb.active
    guide.title = "复核说明"
    instructions = ["本工作簿为已知非holdout的development24，8个二级品类各3件，共48份文案；对照不导入旧人工评分。",
                    "匿名比较两个冻结方案，不猜版本；A/B实际输入事实同一行，Prompt隐藏列按需查看。",
                    "先填黄色人工字段，再看隐藏自动线索；自动指标不等于人工事实错误率。",
                    "两版都保留原可靠核心属性，组合额外加入四区；P5和RAG多值门控差异不修改原分母，源争议单列。",
                    *GUIDE[2:18], "每行正式三指标、四项偏好、reviewer及review_confirmed=1全部完成后才允许汇总。",
                    "匿名映射veryHidden，非加密。不要查看带版本名的输出/报告；复核完成前不解盲。",
                    "不运行test/32条holdout，不用本轮结果继续调Prompt或LoRA。"]
    guide.append(["顺序", "复核口径与操作"])
    for i, instruction in enumerate(instructions, 1):
        guide.append([i, instruction])
        guide.row_dimensions[i + 1].height = 50
    guide.column_dimensions["B"].width = 130
    for row in guide:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    columns = integration_columns()
    ws = wb.create_sheet("盲评成对复核")
    ws.append([k for k, _ in columns])
    ws.append([label for _, label in columns])
    for pair in mapping:
        values = review_fixed(samples[pair["product_id"]], pair, versions)
        ws.append([values.get(k) for k, _ in columns])
    fixed_keys = set(review_fixed(samples[mapping[0]["product_id"]], mapping[0], versions))
    indices = {k: i for i, (k, _) in enumerate(columns, 1)}
    for row in ws:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if cell.row <= 2:
                cell.fill = PatternFill("solid", fgColor="244062")
                cell.font = Font(color="FFFFFF", bold=True)
            else:
                editable = columns[cell.column - 1][0] not in fixed_keys
                cell.protection = Protection(locked=not editable)
                if editable:
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
    for key, index in indices.items():
        letter = get_column_letter(index)
        ws.column_dimensions[letter].width = 62 if key.endswith((*TASK_TYPES, "_complete")) else 42 if any(v in key for v in ("attributes", "notes", "input_facts")) else 18
        if key.endswith("_task_messages"):
            ws.column_dimensions[letter].hidden = True
        ws.cell(1, index).comment = Comment(columns[index - 1][1], "Integration v1")
    for row in range(3, 27):
        ws.row_dimensions[row].height = 210
        ws.cell(row, indices["product_id"]).number_format = "@"
    for p in PREFERENCES:
        dv = DataValidation(type="list", formula1='"A更好,B更好,持平"', allow_blank=True)
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(indices[p])}3:{get_column_letter(indices[p])}26")
    for side in ("A", "B"):
        for field in (*METRICS, *DIAGNOSTICS):
            dv = DataValidation(type="whole", operator="greaterThanOrEqual" if field == "factual_error_count" else "between",
                formula1="0", formula2=None if field == "factual_error_count" else "$I3" if field == "matched_attribute_count" else "1", allow_blank=True)
            dv.showErrorMessage = True
            ws.add_data_validation(dv)
            letter = get_column_letter(indices[f"{side}_{field}"])
            dv.add(f"{letter}3:{letter}26")
    dv = DataValidation(type="list", formula1='"1"', allow_blank=True)
    ws.add_data_validation(dv)
    letter = get_column_letter(indices["review_confirmed"])
    dv.add(f"{letter}3:{letter}26")
    ws.freeze_panes = "F3"
    ws.auto_filter.ref = ws.dimensions
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    sealed = wb.create_sheet("匿名映射_汇总后解盲")
    sealed.append(["pair_id", "product_id", "A", "B"])
    for p in mapping:
        sealed.append([p[k] for k in ("pair_id", "product_id", "A", "B")])
    sealed.sheet_state = "veryHidden"
    hints_ws = wb.create_sheet("自动线索_非人工指标")
    hints_ws.append(["pair_id", "side", "product_id", "diagnostics_REVIEW_ONLY"])
    for pair in mapping:
        for side in ("A", "B"):
            hints_ws.append([pair["pair_id"], side, pair["product_id"], json.dumps(diagnostics[pair[side]][pair["product_id"]], ensure_ascii=False)])
    hints_ws.sheet_state = "hidden"
    wb.save(path)
    atomic_json(out / "automatic_summary.json", {"status": "awaiting_human_review_NOT_FINAL", "statistics": frozen["statistics"],
        "versions": summaries, "human_metrics": None, "human_preferences": None, "guards": cfg["guards"],
        "historical_control_latency": True, "old_single_call_rag_compared": False})
    atomic_json(out / "automatic_diagnostics.json", diagnostics)
    lines = ["# LoRA + RAG 独立 development24 验证", "", "自动生成完成，人工复核待完成。**不以自动 Grounding 判定事实改善。**", "",
             "同24件、原核心分母134；每版本三任务。对照复用已冻结P7输出，组合新增72调用，smoke另24调用不计分。", "",
             "| 自动辅助指标 | LoRA control | LoRA + RAG |", "|---|---:|---:|"]
    for title, key in (("三任务组装结构成功", "assembled_structure_success"), ("疑似评价/效果样本（非错误率）", "claim_suspect_samples_NOT_FACTUAL_RATE"),
                       ("Grounding PASS（非人审）", "grounding_pass_NOT_HUMAN"), ("峰值显存 MiB（allocated）", "peak_allocated_mib")):
        lines.append(f"| {title} | {summaries[VARIANTS[0]][key]} | {summaries[VARIANTS[1]][key]} |")
    for variant in VARIANTS:
        s = summaries[variant]
        lines += ["", f"## {variant}", "", f"字面覆盖（非人工）：{s['literal_core_coverage_NOT_HUMAN']}",
                  f"三任务解析成功数：{s['task_parse_success']}", f"三任务总耗时：{s['latency_seconds']['three_task_total']}",
                  f"Mandatory token 注入：{s['mandatory_core_injection_per_task']}"]
    lines += ["", "## 边界与下一步", "", "人工 matched_attribute_count / fluency_pass / factual_error_count 和四项匿名偏好均未预填，当前不能判断组合是否更优。",
              "人工工作簿：lora_rag_development24_blind_review.xlsx。只填写黄色列，完整确认后再汇总。",
              "旧对照 latency 是历史测量；组合增量包含中性四区事实组织及系统约束从可靠核心属性到可靠属性的必要适配，不是纯检索因果隔离。",
              "P7既有耳机类别多值输入与RAG门控存在1项差异，已记录protocol_manifest；原输入/分母不改，也不将该字段重复注入RAG区。",
              "无训练、无test读取、无holdout清单读取/生成；原RAG v2归档与frozen_not_executed状态不改。没有commit/push。"]
    (out / "phase_report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    verify_protocol()
    paths = [p for p in out.rglob("*") if p.is_file()]
    items = [{"path": str(p.relative_to(ROOT)).replace("\\", "/"), "sha256": digest(p)} for p in sorted(paths)]
    atomic_json(out / "review_manifest.json", {"status": "awaiting_human_review", "initial_workbook_sha256": digest(path),
        "files": [i for i in items if i["path"] != str(path.relative_to(ROOT)).replace("\\", "/")],
        "workbook": str(path.relative_to(ROOT)).replace("\\", "/"), "labels_prefilled": False, "products": 24, "documents": 48})
    items.append({"path": str((out / "review_manifest.json").relative_to(ROOT)).replace("\\", "/"), "sha256": digest(out / "review_manifest.json")})
    atomic_json(out / "automatic_phase_manifest.json", {"status": "generation_complete_human_pending", "files": items,
        "created_at_utc": now(), "guards": cfg["guards"], "code_and_prompt_freeze": "protocol_manifest.json"})
    (out / "checksums.sha256").write_text("".join(f"{i['sha256']}  {i['path']}\n" for i in items), encoding="utf-8")
    print(json.dumps({"status": "awaiting_human_review", "workbook": str(path), "versions": summaries}, ensure_ascii=False), flush=True)


def validate_review_rows(ws, samples, versions, mapping):
    columns = integration_columns()
    if ws.max_row != 26 or [c.value for c in ws[1]] != [k for k, _ in columns]:
        raise ValueError("Review coverage/columns changed")
    rows = []
    for index, pair in enumerate(mapping, 3):
        row = {k: ws.cell(index, j).value for j, (k, _) in enumerate(columns, 1)}
        for key, expected in review_fixed(samples[pair["product_id"]], pair, versions).items():
            if row[key] != expected:
                raise ValueError(f"Frozen review data changed: {pair['pair_id']} {key}")
        if str(row["review_confirmed"]).strip() != "1" or not str(row["reviewer"] or "").strip():
            raise ValueError("Incomplete human confirmation")
        if any(row[p] not in CHOICES for p in PREFERENCES):
            raise ValueError("Incomplete/invalid anonymous preferences")
        for side in ("A", "B"):
            official_metrics([{ "product_id": pair["product_id"], "core_attribute_count": row["core_attribute_count"],
                               **{m: row[f"{side}_{m}"] for m in METRICS}}])
            for d in DIAGNOSTICS:
                value = row[f"{side}_{d}"]
                if value is not None and str(value) not in ("0", "1"):
                    raise ValueError("Invalid diagnostic")
                if str(value) == "1" and not row[f"{side}_notes"]:
                    raise ValueError("Diagnosis needs evidence")
            if int(row[f"{side}_factual_error_count"]) > 0 and not row[f"{side}_notes"]:
                raise ValueError("Factual errors need evidence")
        rows.append(row)
    return rows


def summarize():
    from openpyxl import load_workbook
    _, _, samples, versions, mapping = paired_data()
    out = resolve(OUTPUT)
    for name in ("blind_summary.json", "human_metrics_unblinded.json", "human_review_final.csv"):
        if (out / name).exists():
            raise FileExistsError("Never overwrite finalized human metrics")
    manifest = read_json(out / "review_manifest.json")
    for item in manifest["files"]:
        if digest(resolve(item["path"])) != item["sha256"]:
            raise ValueError("Frozen review output changed")
    path = out / "lora_rag_development24_blind_review.xlsx"
    wb = load_workbook(path, data_only=False)
    sealed = wb["匿名映射_汇总后解盲"]
    if list(sealed.values)[1:] != [tuple(p[k] for k in ("pair_id", "product_id", "A", "B")) for p in mapping]:
        raise ValueError("Workbook mapping changed")
    rows = validate_review_rows(wb["盲评成对复核"], samples, versions, mapping)
    atomic_json(out / "blind_summary.json", {"status": "complete_anonymous_aggregate_before_unblinding",
        "preferences": {p: dict(Counter(r[p] for r in rows)) for p in PREFERENCES}, "workbook_sha256": digest(path)})
    judgments, prefs = {v: [] for v in VARIANTS}, {p: Counter() for p in PREFERENCES}
    for row, pair in zip(rows, mapping):
        for side in ("A", "B"):
            item = {"product_id": pair["product_id"], "variant": pair[side], "category_l2": samples[pair["product_id"]]["category_l2"],
                "core_attribute_count": row["core_attribute_count"], **{m: int(row[f"{side}_{m}"]) for m in METRICS},
                **{d: row[f"{side}_{d}"] for d in DIAGNOSTICS}, "reviewer": row["reviewer"], "review_notes": row[f"{side}_notes"] or ""}
            judgments[pair[side]].append(item)
        for p in PREFERENCES:
            prefs[p]["tie" if row[p] == "持平" else pair[row[p][0]]] += 1
    atomic_json(out / "human_metrics_unblinded.json", {"status": "completed_human_review", "versions": {v: official_metrics(r) for v, r in judgments.items()},
        "by_category": {v: {c: official_metrics([r for r in judgments[v] if r["category_l2"] == c]) for c in sorted({s["category_l2"] for s in samples.values()})} for v in VARIANTS},
        "preferences": {p: {k: prefs[p][k] for k in (*VARIANTS, "tie")} for p in PREFERENCES},
        "workbook_sha256": digest(path), "blind_summary_sha256": digest(out / "blind_summary.json"), "no_test_or_holdout": True})
    exported = [r for rows in judgments.values() for r in rows]
    with (out / "human_review_final.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(exported[0]))
        writer.writeheader()
        writer.writerows(exported)
    print("Human review complete. Stop; no test, holdout, training or prompt tuning.")
